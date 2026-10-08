"""Verify existing numerical evidence and restored maps before document installation."""
from pathlib import Path
import hashlib
import json
import re
import shutil
import zipfile
import xml.etree.ElementTree as ET

import numpy as np
import pandas as pd
import pymupdf
from docx import Document
from openpyxl import load_workbook

from .prepare import HERE, ROOT, STUDY, METRICS, sha

STAGING = HERE / 'staging'
ACTIVE = ROOT / 'publication/european_wheat_stb'
BEFORE = HERE / 'before/publication/european_wheat_stb'
REVISION = ROOT / 'analysis/paper_study/nature_food_revision_20261007'
MAIN_MAP = 'fig6_european_spatial_results'
SUPPLEMENT_MAPS = [
    'figS13_spatial_symptom_changes', 'figS14_spatial_yield_transfers',
    'figS15_archived_domain', 'figS16_archived_era5_baseline',
    'figS17_archived_climate_maps', 'figS18_archived_production_exposure',
    'figS19_full_grid_example_2001',
]


def dump(path, value):
    Path(path).write_text(json.dumps(value, indent=2, ensure_ascii=False) + '\n')


def old_package():
    with zipfile.ZipFile(BEFORE / 'Software_and_Evidence.zip') as archive:
        return json.loads(archive.read('PACKAGE_MANIFEST.json'))


def verify():
    receipt = json.loads((STAGING / 'document_build_receipt.json').read_text())
    for relative, digest in receipt['sources'].items():
        assert sha(ROOT / relative) == digest, relative
    changed_sources = {
        str((STUDY / 'manuscript' / name).relative_to(ROOT))
        for name in ['build.py', 'figures.py', 'sections.py']
    }
    original_manifest = old_package()
    for record in original_manifest:
        publication_output = record['path'].startswith((
            'publication/european_wheat_stb/',
            str((STUDY / 'manuscript/figures').relative_to(ROOT)) + '/'))
        if record['path'] not in changed_sources and not publication_output:
            assert sha(ROOT / record['path']) == record['sha256'], record['path']
    for record in json.loads((HERE / 'before/archive_manifest.json').read_text()):
        assert sha(HERE / 'before' / record['path']) == record['sha256'], record['path']

    assert receipt['main_figures'] == 6 and receipt['main_tables'] == 0
    assert receipt['supplementary_figures'] == 19 and receipt['supplementary_tables'] == 33
    assert receipt['reference_count'] == 27 and receipt['abstract_words'] <= 150
    body_files = ['introduction.txt', 'field_results.txt', 'baseline_results.txt',
                  'climate_results.txt', 'discussion_field_and_limits.txt']
    body_words = sum(len((STAGING / name).read_text().split()) for name in body_files)
    assert body_words <= 3000
    assert (STAGING / 'introduction.txt').read_text().splitlines()[0] != 'Introduction'
    text = (STAGING / 'Manuscript.txt').read_text()
    supplement = (STAGING / 'Supplementary_Information.txt').read_text()
    assert not re.search(r'\bn\.d\.|\[\[|APSIM Initiative|Python Crop Simulation Environment', text)
    assert not re.search(r'(?<![\d.])\d{1,2},\s+\d{3}(?![\d.])|\d(?=[a-z]{2})', text + supplement)
    assert 'Foulkes et al.’s Table 2 supplies' in text
    assert 'Table 2 entries are mixed-model' in supplement
    assert 'Reported Table 3 coefficients' in supplement
    assert 'Treatment-associated relative harvest yields are recovered' in text
    assert 'do not establish a general absolute yield-loss function' in text
    assert 'one harvest year does not constitute completed baseline' in text
    records = json.loads((STAGING / 'references.csl.json').read_text())
    assert len(records) == 27 and all(record.get('issued') for record in records)

    # The existing validation figures and every earlier numerical data sheet remain.
    for name in ['fig1_framework', 'fig2_flag_stage_validation',
                 'fig3_onset_and_detection', 'fig4_timing_and_yield_relevance',
                 'fig5_conditional_climate_changes']:
        assert sha(STAGING / 'figures' / f'{name}.png') == sha(BEFORE / 'figures' / f'{name}.png'), name
    german = pd.read_csv(STAGING / 'figures/fig2_german_stage_errors.csv')
    intervals = pd.read_csv(STAGING / 'figures/fig2_field_stage_intervals.csv')
    assert len(german) == 79011 and len(intervals) == 259
    assert set(german.BBCH) == {10, 31, 51, 85}
    assert set(intervals.event) == {31, 32, 33, 37, 39, 51, 65, 85}
    workbook = load_workbook(STAGING / 'Source_Data.xlsx', read_only=True, data_only=True)
    previous = load_workbook(BEFORE / 'Source_Data.xlsx', read_only=True, data_only=True)
    assert len(receipt['source_workbook']) == 54 and len(workbook.sheetnames) == 56
    assert len(previous.sheetnames) == 44
    for name in previous.sheetnames[:-2]:
        assert list(workbook[name].values) == list(previous[name].values), name
    for record in receipt['source_workbook']:
        source = ROOT / record['source_path']
        if record['source_path'].startswith('publication/european_wheat_stb/'):
            source = STAGING / Path(record['source_path']).relative_to('publication/european_wheat_stb')
        assert sha(source) == record['sha256'], record['source_path']
        assert sum(1 for _ in workbook[record['sheet']].values) - 1 == record['rows'], record['sheet']
    previous.close()
    workbook.close()

    map_receipt = json.loads((HERE / 'map_data_receipt.json').read_text())
    for relative, digest in map_receipt['source_sha256'].items():
        assert sha(ROOT / relative) == digest, relative
    for name, digest in map_receipt['table_sha256'].items():
        assert sha(HERE / name) == digest, name
    assert map_receipt['maps_interpolated'] is False
    assert map_receipt['current_maps_full_domain_census'] is False
    wheat = pd.read_csv(HERE / 'map_wheat_domain.csv')
    fields = pd.read_csv(HERE / 'map_field_locations.csv')
    assert len(wheat) == wheat.cell_id.nunique() == 14941
    assert len(fields) == 216
    draws = pd.read_csv(HERE / 'map_draw_values.csv')
    ensemble = pd.read_csv(HERE / 'map_ensemble_values.csv')
    expected = np.divide(draws.numerator, draws.denominator,
                         out=np.full(len(draws), np.nan), where=draws.denominator.to_numpy() > 0)
    np.testing.assert_allclose(expected, draws.mapped_value, rtol=0, atol=1e-12, equal_nan=True)
    assert draws.spatial_draw_id.nunique() == 64 and draws.cell_id.nunique() == 62
    assert set(draws.metric) == set(METRICS)
    assert set(draws.scenario) == {'ssp126', 'ssp245', 'ssp585'}
    keys = ['kind', 'scenario', 'period', 'metric', 'cell_id']
    unique = draws.drop_duplicates(keys + ['model'])
    computed = unique.groupby(keys, as_index=False).agg(
        value=('mapped_value', 'mean'), gcm_min=('mapped_value', 'min'),
        gcm_max=('mapped_value', 'max'), n_gcm=('mapped_value', 'count'))
    merged = ensemble.merge(computed, on=keys, validate='one_to_one', suffixes=('_map', '_check'))
    assert len(merged) == len(ensemble) and merged.n_gcm_map.eq(3).all()
    for name in ['value', 'gcm_min', 'gcm_max', 'n_gcm']:
        np.testing.assert_allclose(merged[name + '_map'], merged[name + '_check'], rtol=0, atol=1e-12)
    assert ensemble.groupby(['kind', 'scenario', 'period', 'metric']).size().eq(62).all()

    main_document = Document(STAGING / 'Manuscript.docx')
    supplementary_document = Document(STAGING / 'Supplementary_Information.docx')
    old_main_document = Document(BEFORE / 'Manuscript.docx')
    assert len(main_document.inline_shapes) == 6 and len(main_document.tables) == 0
    assert len(supplementary_document.inline_shapes) == 19 and len(supplementary_document.tables) == 33
    for original, moved in zip(old_main_document.tables, supplementary_document.tables[28:31]):
        assert [[cell.text for cell in row.cells] for row in original.rows] == [
            [cell.text for cell in row.cells] for row in moved.rows]
    with zipfile.ZipFile(STAGING / 'Manuscript.docx') as archive:
        assert hashlib.sha256(archive.read('word/media/image6.png')).hexdigest() == sha(
            STAGING / 'figures' / f'{MAIN_MAP}.png')
    archived = json.loads((HERE / 'archived_map_registry.json').read_text())
    with zipfile.ZipFile(STAGING / 'Supplementary_Information.docx') as archive:
        for record in archived:
            number = int(record['figure'].split('_')[0].replace('figS', ''))
            for extension in ['png', 'pdf']:
                assert sha(STAGING / 'supplementary_figures' / f"{record['figure']}.{extension}") == record[extension + '_sha256']
            assert hashlib.sha256(archive.read(f'word/media/image{number}.png')).hexdigest() == record['png_sha256']
    for figure in SUPPLEMENT_MAPS:
        for extension in ['png', 'pdf']:
            assert (STAGING / 'supplementary_figures' / f'{figure}.{extension}').is_file()
    for path in [*STAGING.glob('figures/*.svg'), *STAGING.glob('supplementary_figures/*.svg')]:
        ET.parse(path)

    full = REVISION / 'continental_replay/annual_outputs/nasa/ACCESS-CM2/ssp585/2001.parquet'
    annual = json.loads(full.with_suffix('.json').read_text())
    assert sha(full) == annual['output_sha256']
    data = pd.read_parquet(full)
    assert len(data) == 14941 and data.cell_id.nunique() == 14941
    assert int(data.valid_complete_season.sum()) == 14772
    assert annual['status_counts'] == dict(complete=14772, soft_dough_not_reached=160, missing_calendar=9)
    reconciliation = json.loads((REVISION / 'continental_replay/sample_to_full_grid_reconciliation.json').read_text())
    assert reconciliation['status'] == 'passed' and reconciliation['single_grid_season_only']
    assert reconciliation['numeric_metrics_reproduced'] == 58 and reconciliation['date_metrics_reproduced'] == 10

    pages = {}
    for name in ['Manuscript', 'Supplementary_Information']:
        with pymupdf.open(STAGING / f'{name}.pdf') as document:
            pages[name] = len(document)
            content = '\n'.join(page.get_text() for page in document)
            for number in ([6] if name == 'Manuscript' else range(13, 20)):
                assert f'Figure {"S" if name != "Manuscript" else ""}{number} |' in content
    yield_checks = json.loads((REVISION / 'yield_response/final_verification_receipt.json').read_text())
    assert yield_checks['status'] == 'complete' and yield_checks['targeted_tests'] == 13
    assert yield_checks['transferable_yield_response_demonstrated'] is False
    response_manifest = json.loads((REVISION / 'yield_response/artifact_manifest.json').read_text())
    for record in response_manifest['files']:
        assert sha(REVISION / 'yield_response' / record['file']) == record['sha256'], record['file']
    maps = dict(status='passed', main_map_figure=MAIN_MAP, supplementary_map_figures=SUPPLEMENT_MAPS,
                historical_design_regression='Earlier proposed two-map layout consolidated into one four-panel main figure.',
                all_original_map_images_unchanged=True, current_sample_points=62,
                completed_full_grid_example_year=2001, current_full_period_census_complete=False)
    dump(HERE / 'map_presence_after.json', maps)
    return dict(
        status='passed', scope='Restored spatial evidence, all prior validation, and completed crop-response diagnostics',
        main_figures=6, main_tables=0, main_body_words=body_words, abstract_words=receipt['abstract_words'],
        reference_count=27, supplementary_figures=19, supplementary_tables=33,
        source_workbook_sheets=56, prior_data_sheets_retained=42, pdf_pages=pages,
        numerical_sample_map_records=len(draws), unique_spatial_map_records=len(ensemble),
        archived_map_figures_retained=4, all_map_image_hashes_verified=True,
        current_full_grid_example=dict(year=2001, registered=14941, complete=14772),
        original_model_calibration_and_frozen_evidence_unchanged=True,
        original_publication_archive_unchanged=True, all_prior_validation_retained=True,
        all_original_main_tables_retained_in_supplement=True,
        new_yield_analysis_targeted_tests=13, original_repository_tests=305,
        original_repository_test_run_repeated=False, predictive_chain_validated=False,
        actual_grain_loss_estimates_identified=False, map_presence=maps,
    )


def install(audit):
    dump(HERE / 'preinstall_verification.json', audit)
    for path in STAGING.rglob('*'):
        if path.is_file():
            target = ACTIVE / path.relative_to(STAGING)
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, target)
    for path in (STAGING / 'figures').iterdir():
        if path.is_file():
            shutil.copy2(path, STUDY / 'manuscript/figures' / path.name)
    dump(HERE / 'publication_verification.json', audit)
    dump(ACTIVE / 'spatial_and_crop_response_verification.json', audit)
    dump(ACTIVE / 'validation_and_figures_verification.json', audit)
    paths = {ROOT / record['path'] for record in old_package()}
    safe_directories = [
        HERE, ACTIVE / 'figures', ACTIVE / 'supplementary_figures',
        REVISION / 'epidemic_evaluation', REVISION / 'yield_response',
        REVISION / 'basf_crop_response', REVISION / 'basf_yield_linkage',
        REVISION / 'disease_data_audit', REVISION / 'europe_forcing_audit',
        REVISION / 'continental_replay',
    ]
    for directory in safe_directories:
        for path in directory.rglob('*'):
            if not path.is_file() or any(part in ['before', 'staging', 'preview', '__pycache__']
                                         for part in path.relative_to(directory).parts):
                continue
            if path.suffix in {'.py', '.csv', '.json', '.parquet', '.txt', '.md', '.log', '.png', '.svg', '.pdf'}:
                paths.add(path)
    # Audit code and metadata link to original public sources; AHDB numeric extraction
    # and third-party protocol/document reproductions are not redistributed.
    audit_directory = REVISION / 'yield_data_audit'
    paths.update(path for path in audit_directory.glob('*.py'))
    paths.update(audit_directory / name for name in [
        'combined_audit_summary.json', 'final_data_verification.json',
        'nfts_extraction_summary.json', 'ahdb_extraction_summary.json',
        'source_usage_review.json',
    ] if (audit_directory / name).exists())
    paths.add(ACTIVE / 'spatial_and_crop_response_verification.json')
    paths.add(HERE / 'before/archive_manifest.json')
    assert not any('canopy_forecast' in str(path) for path in paths)
    assert not any(path.suffix == '.xlsx' and 'yield_data_audit' in str(path) for path in paths)
    reproduction = (BEFORE / 'REPRODUCIBILITY.txt').read_text()
    reproduction += (
        '\nSpatial and crop-response evidence revision (2026-10-07).\n'
        'The current article has six main figures, no main tables, 19 supplementary figures '
        'and 33 supplementary tables. Source_Data.xlsx has 54 numerical data sheets plus '
        'the source manifest and definitions. All previous 42 data sheets are retained.\n'
        'The four earlier map images remain unchanged in Figures S15–S18 and keep their '
        'seasonal-v1 model identity. Current multi-period spatial results contain 64 '
        'registered draws at 62 unique cells, shown without spatial interpolation. '
        'Figure S19 uses a completed full-grid replay for 2001 only. A complete current-model '
        'baseline/future census and validated European grain-loss predictions are absent.\n'
        'python -m analysis.paper_study.map_restoration_20261007.prepare regenerates map source tables '
        'from existing completed outputs. For a new complete document build, import '
        'analysis.paper_study.overwinter_leaf_model_20261006.manuscript.build, set build.DEST '
        'to a new directory and call build.main(). Source_Data manifests identify every CSV '
        'and checksum used by the figures and tables.\n'
        'Completed public-trial crop-response analyses and their holdout memberships are '
        'included as failed-transfer research results, outside model/. The Nordic source '
        'records retain attribution and source URLs. Public source HTML and AHDB numeric '
        'source workbooks are not bundled. Original raw climate caches are not bundled.\n'
        'The current numerical epidemic and crop-response errors limit inference to '
        'conditional scenarios. The original repository test run recorded 305 passing tests '
        'with seven existing warnings; 13 later crop-response tests and three spatial '
        'bookkeeping tests were recorded separately. This document installation does '
        'not claim a repeated complete model test run.\n'
    )
    (ACTIVE / 'REPRODUCIBILITY.txt').write_text(reproduction)
    manifest = [dict(path=str(path.relative_to(ROOT)), bytes=path.stat().st_size, sha256=sha(path))
                for path in sorted(paths)]
    package = ACTIVE / 'Software_and_Evidence.zip'
    with zipfile.ZipFile(package, 'w', zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for record in manifest:
            archive.write(ROOT / record['path'], record['path'])
        archive.writestr('PACKAGE_MANIFEST.json', json.dumps(manifest, indent=2) + '\n')
        archive.writestr('REPRODUCIBILITY.txt', reproduction)
    with zipfile.ZipFile(package) as archive:
        assert archive.testzip() is None and len(archive.namelist()) == len(manifest) + 2
        for record in manifest:
            assert hashlib.sha256(archive.read(record['path'])).hexdigest() == record['sha256']
    dump(ACTIVE / 'software_package_receipt.json', dict(
        status='complete', members=len(manifest) + 2, bundled_sources_and_evidence=manifest,
        package_sha256=sha(package), raw_climate_cache_bundled=False,
        publication_revision='Restored maps and completed crop-response diagnostics, 2026-10-07'))
    final = json.loads((BEFORE / 'final_verification.json').read_text())
    final.update(
        manuscript_pdf_pages=audit['pdf_pages']['Manuscript'],
        supplementary_pdf_pages=audit['pdf_pages']['Supplementary_Information'],
        main_figures=6, main_tables=0, supplementary_figures=19, supplementary_tables=33,
        workbook_sheets=56, software_package_members=len(manifest) + 2,
        main_body_words=audit['main_body_words'], abstract_words=audit['abstract_words'],
        verification_scope=audit['scope'], original_repository_tests=305,
        original_repository_test_run_repeated=False, completed_response_tests=13,
        predictive_chain_validated=False, current_full_period_european_census_complete=False,
        spatial_and_crop_response_verification=audit)
    final['review_receipts'] = [dict(path=str((HERE / 'publication_verification.json').relative_to(ROOT)),
                                      sha256=sha(HERE / 'publication_verification.json'))]
    final['artifact_sha256'] = {str(path.relative_to(ACTIVE)): sha(path)
                               for path in sorted(ACTIVE.rglob('*'))
                               if path.is_file() and path.name != 'final_verification.json'}
    dump(ACTIVE / 'final_verification.json', final)
    for name, digest in final['artifact_sha256'].items():
        assert sha(ACTIVE / name) == digest, name
    print(json.dumps({key: audit[key] for key in [
        'status', 'main_figures', 'main_tables', 'main_body_words', 'abstract_words',
        'supplementary_figures', 'supplementary_tables', 'source_workbook_sheets', 'pdf_pages']},
        indent=2))


if __name__ == '__main__':
    install(verify())
