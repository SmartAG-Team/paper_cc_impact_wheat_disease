"""Independent reporting and package checks; these do not validate biological prediction."""
from pathlib import Path
from zipfile import ZipFile
import argparse
import hashlib
import json
import re
import numpy as np
import pandas as pd
from docx import Document
import pymupdf
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]


def verify(output):
    from .verify_spatial import verify as verify_spatial
    output = Path(output).resolve()
    main = (output / 'Manuscript.txt').read_text()
    si = (output / 'Supplementary_Information.txt').read_text()
    abstract = (output / 'abstract.txt').read_text().strip()
    if len(abstract.split()) > 150 or abstract not in main or re.search(r'\[[1-9]\d*', abstract):
        raise AssertionError('Abstract length or cross-format consistency failed.')
    observations = pd.read_csv(ROOT / 'analysis/paper_study/nature_food_revision_20261007/epidemic_evaluation/observed_vs_frozen_assessments.csv')
    checks = {}
    for part in ['reused_BASF2019', 'reused_strict_Corteva']:
        x = observations[observations.partition.eq(part)].copy()
        x['sq'] = (x.frozen_damage_percent - x.observed_percent)**2
        groups = (x.groupby(['coordinate_year', 'field_id', 'leaf_index']).sq.mean()
                  .groupby(['coordinate_year', 'field_id']).mean().groupby('coordinate_year').mean())
        score = float(np.sqrt(groups.mean()))
        if f'{score:.2f}' not in main:
            raise AssertionError('Severity metric does not match source observations.')
        checks[part] = {'severity_RMSE_pp': score, 'assessments': len(x), 'groups': len(groups)}

    source = pd.read_csv(ROOT / 'analysis/paper_study/yield_evidence_20261007/foulkes2006_group_separated_predictions.csv')
    weights = 1/source.groupby('genetic_background').genetic_background.transform('size')
    weights /= weights.sum()
    error = source.relative_held_out_prediction - source.derived_yield_loss_fraction
    relative_rmse = float(100*np.sqrt(np.sum(weights*error**2)))
    if f'{relative_rmse:.2f}' not in main or len(source) != 8 or source.genetic_background.nunique() != 5:
        raise AssertionError('Crop-response population or score differs.')
    for name, following in [('S6a', 'S6b'), ('S6b', 'S6c')]:
        section = si.split(f'Table {name} |', 1)[1].split(f'Table {following} |', 1)[0]
        rows = [line for line in section.splitlines() if '\t' in line]
        if len(rows)-1 != 8:
            raise AssertionError(f'{name} omits contrasts used in the metrics.')
    checks['published_canopy_yield'] = {'contrasts': 8, 'backgrounds': 5, 'relative_RMSE_pp': relative_rmse}

    components = pd.read_csv(ROOT / 'analysis/paper_study/nature_food_fix_20261007/climate/decomposition/supported_forcing_decomposition_by_gcm.csv')
    means = components.groupby('component')['mean'].mean()
    net = float(means['weather_shapley'] + means['host_shapley'])
    ratio = float(-100*means['host_shapley']/means['weather_shapley'])
    derived = pd.read_csv(HERE / 'derived/weather_host_offset.csv')
    row = derived[derived.climate_model.eq('Three-model ensemble')].iloc[0]
    if abs(row.net_change-net) > 1e-12 or abs(row.host_offset_percent-ratio) > 1e-12:
        raise AssertionError('Contribution ratio does not reconcile independently.')
    if '+2.64' not in abstract or '71%' not in abstract or '70.6%' not in main:
        raise AssertionError('Rounded climate contribution summary differs.')
    checks['model_decomposition'] = {'net_change': net, 'host_offset_percent': ratio}

    captions = json.loads((output / 'figures/captions.json').read_text())
    if len(captions) != 5 or len(list((output / 'figures').glob('fig*.png'))) != 5:
        raise AssertionError('Main figure selection differs.')
    refs = re.findall(r'(?<!Supplementary )(?:Fig\.|Figure) (\d+)', main.split('Figure legends')[0])
    if list(dict.fromkeys(map(int, refs))) != [1,2,3,4,5]:
        raise AssertionError('Main figure citation order differs.')
    declared = set(re.findall(r'^Table (S\d+[a-z]?) \|', si, flags=re.M))
    for identifier in re.findall(r'(?:Table|Tables) (S\d+[a-z]?)', main):
        if identifier not in declared and not any(x.startswith(identifier) for x in declared):
            raise AssertionError('Unresolved supplementary table reference: '+identifier)
    doc = Document(output / 'Manuscript.docx')
    if len(doc.tables) != 1 or len(doc.inline_shapes) != 5:
        raise AssertionError('Editable manuscript lacks its main table or figures.')
    wordtext = '\n'.join(p.text for p in doc.paragraphs)
    if abstract not in wordtext or 'Gang Zhao' not in wordtext or 'no competing interests' not in wordtext:
        raise AssertionError('Word manuscript differs from its declared text.')
    abstract_index = next(i for i, paragraph in enumerate(doc.paragraphs) if paragraph.text == abstract)
    introduction = doc.paragraphs[abstract_index + 1]
    if not introduction.paragraph_format.page_break_before:
        raise AssertionError('Word Introduction is not separated from the abstract.')
    if any(run.font.superscript for run in doc.paragraphs[abstract_index].runs):
        raise AssertionError('Word abstract contains superscript citations.')
    pdf_counts = {}
    for name in ['Manuscript.pdf','Supplementary_Information.pdf','Supplementary_Model_Specification.pdf']:
        pdf = pymupdf.open(output/name)
        if any(len(page.get_text().strip()) < 5 for page in pdf):
            raise AssertionError('Blank or inaccessible PDF page: '+name)
        if name == 'Manuscript.pdf':
            first_page = ' '.join(pdf[0].get_text().split())
            second_page = ' '.join(pdf[1].get_text().split())
            if abstract not in first_page or introduction.text[:80] in first_page or introduction.text[:80] not in second_page:
                raise AssertionError('PDF abstract and Introduction are not clearly separated.')
        pdf_counts[name] = len(pdf)

    workbook = load_workbook(output/'Source_Data.xlsx', read_only=True)
    for values in list(workbook['Source_manifest'].values)[1:]:
        _, name, _, expected = values
        path = Path(name)
        if name.startswith('publication/european_wheat_stb/'):
            path = output/path.relative_to('publication/european_wheat_stb')
        else:path = ROOT/path
        if hashlib.sha256(path.read_bytes()).hexdigest() != expected:
            raise AssertionError('Source workbook provenance mismatch: '+name)
    archive_counts = None
    if (output/'Software_and_Evidence.zip').exists():
        with ZipFile(output/'Software_and_Evidence.zip') as archive:
            if archive.testzip() is not None:raise AssertionError('Archive integrity failed.')
            manifest = json.loads(archive.read('PACKAGE_MANIFEST.json'))
            for name, record in manifest['members'].items():
                content = archive.read(name)
                if hashlib.sha256(content).hexdigest() != record['sha256'] or len(content) != record['bytes']:
                    raise AssertionError('Archived member checksum mismatch: '+name)
            for local in [HERE/'sections.py',HERE/'build.py',HERE/'figures.py',ROOT/'analysis/paper_study/yield_evidence_20261007/tables.py']:
                name = local.relative_to(ROOT).as_posix()
                if archive.read(name) != local.read_bytes():raise AssertionError('Archived publication source is stale: '+name)
            archive_counts = len(manifest['members'])
    report = {'status': 'passed', 'scope': 'Reporting, source calculations and archive consistency; no claim of biological validation',
              'abstract_words': len(abstract.split()), 'main_figures': 5, 'main_tables': 1,
              'abstract_citation_free_and_separated': True,
              'supplementary_tables': len(declared), 'supplementary_figures': len(list((output/'supplementary_figures').glob('figS*.png'))),
              'independent_calculations': checks, 'independent_spatial_checks': verify_spatial(), 'PDF_pages': pdf_counts,
              'source_workbook_sheets': len(workbook.sheetnames), 'archived_members_checked': archive_counts}
    (output/'verification_receipt.json').write_text(json.dumps(report, indent=2)+'\n')
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, required=True)
    print(json.dumps(verify(parser.parse_args().output), indent=2))
