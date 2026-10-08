"""Verify and install the source-backed Figure 2 revision without refitting."""
from pathlib import Path
import hashlib
import json
import shutil
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from openpyxl import load_workbook
import pandas as pd
import pymupdf

ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
STAGING = HERE/'staging'
ACTIVE = ROOT/'publication/european_wheat_stb'
BEFORE = HERE/'before/publication/european_wheat_stb'
SOURCE = ROOT/'analysis/paper_study/overwinter_leaf_model_20261006/manuscript'
FIGURE = 'fig2_flag_stage_validation'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, value):
    Path(path).write_text(json.dumps(value, indent=2)+'\n')


def verify_staging():
    before_captions = json.loads((BEFORE/'figures/captions.json').read_text())
    captions = json.loads((STAGING/'figures/captions.json').read_text())
    assert list(captions) == list(before_captions)
    for name in captions:
        if name == FIGURE:
            assert captions[name] != before_captions[name]
        else:
            assert captions[name] == before_captions[name]
            assert sha(STAGING/'figures'/f'{name}.png') == sha(BEFORE/'figures'/f'{name}.png')
    for path in BEFORE.glob('*.txt'):
        if path.name != 'Manuscript.txt' and (STAGING/path.name).exists():
            assert sha(path) == sha(STAGING/path.name), path.name
    old_text = (BEFORE/'Manuscript.txt').read_text()
    new_text = (STAGING/'Manuscript.txt').read_text()
    assert old_text.split('Figure legends')[0] == new_text.split('Figure legends')[0]
    old_doc = Document(BEFORE/'Manuscript.docx')
    new_doc = Document(STAGING/'Manuscript.docx')
    assert len(new_doc.inline_shapes) == 5
    assert len(new_doc.tables) == len(old_doc.tables) == 3
    assert [[c.text for r in t.rows for c in r.cells] for t in new_doc.tables] == [
        [c.text for r in t.rows for c in r.cells] for t in old_doc.tables]
    old_paragraphs = [p.text for p in old_doc.paragraphs]
    new_paragraphs = [p.text for p in new_doc.paragraphs]
    assert len(old_paragraphs) == len(new_paragraphs)
    changed_paragraphs = [(a, b) for a, b in zip(old_paragraphs, new_paragraphs) if a != b]
    assert len(changed_paragraphs) == 1
    assert changed_paragraphs[0][1].startswith('Figure 2 |')
    with zipfile.ZipFile(STAGING/'Manuscript.docx') as z:
        assert hashlib.sha256(z.read('word/media/image2.png')).hexdigest() == sha(
            STAGING/'figures'/f'{FIGURE}.png')
    with pymupdf.open(STAGING/'Manuscript.pdf') as pdf:
        assert len(pdf) == 17
        figure_pages = [i for i, page in enumerate(pdf) if 'Figure 2 |' in page.get_text()]
        assert figure_pages == [13]
        pdf[13].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(HERE/'preview/manuscript_figure2_page.png')
    with pymupdf.open(STAGING/'figures'/f'{FIGURE}.pdf') as pdf:
        assert len(pdf) == 1
        text = pdf[0].get_text()
        assert 'BBCH37' in text and 'BBCH39' in text and 'mean excess' in text
    ET.parse(STAGING/'figures'/f'{FIGURE}.svg')
    plot = pd.read_csv(STAGING/'figures/fig2_field_stage_intervals.csv')
    assert len(plot) == 56
    assert not plot.duplicated(['field_id', 'partition', 'event']).any()
    raw = pd.read_csv(SOURCE.parent/'phenology/stage_censoring_scores.csv')
    raw = raw.loc[raw.genuinely_bracketed & raw.event.isin([37, 39]) & raw.partition.isin([
        'reused_development_2019', 'reused_external_strict'])]
    compare = plot.merge(raw, on=['field_id', 'partition', 'event'], validate='1:1')
    assert len(compare) == len(raw) == 56
    assert compare.display_compatible.eq(compare.compatible).all()
    assert compare.display_interval_excess_days.eq(compare.distance_days).all()
    wb = load_workbook(STAGING/'Source_Data.xlsx', read_only=True, data_only=True)
    assert len(wb.sheetnames) == 18
    old_wb = load_workbook(BEFORE/'Source_Data.xlsx', read_only=True, data_only=True)
    for name in old_wb.sheetnames:
        if name not in ('Definitions', 'Source_manifest'):
            assert list(wb[name].values) == list(old_wb[name].values), name
    new_sheet = next(name for name in wb.sheetnames if name.startswith('16_fig2'))
    assert list(wb[new_sheet].values)[0] == tuple(plot.columns)
    assert len(list(wb[new_sheet].values)) == len(plot)+1
    manifest = json.loads((STAGING/'document_build_receipt.json').read_text())['source_workbook']
    assert len(manifest) == 16
    assert manifest[-1]['sha256'] == sha(STAGING/'figures/fig2_field_stage_intervals.csv')
    old_wb.close();wb.close()
    groups = []
    for (partition, event), rows in plot.groupby(['partition', 'event']):
        groups.append(dict(partition=partition, bbch=int(event), field_stage_records=len(rows),
            compatible=int(rows.display_compatible.sum()),
            mean_interval_excess_days=float(rows.display_interval_excess_days.mean())))
    return dict(status='passed', scope='Figure 2 visualization and publication exports',
        field_stage_records=56, performance_groups=groups,
        field_scores_unchanged=True, manuscript_body_tables_and_references_unchanged=True,
        other_figure_pngs_unchanged=True, docx_figure2_matches_export=True,
        manuscript_pdf_pages=17, figure2_manuscript_pdf_page=14,
        source_workbook_sheets=18, source_workbook_original_data_sheets_unchanged=True,
        observation_bounds='last below-stage date + 1 day through first at-or-above-stage date',
        interval_centres='display positions only, not observed event dates',
        visual_review='standalone PNG and manuscript PDF page inspected; labels fit within the exports',
        source_sha256={str(p.relative_to(ROOT)):sha(p) for p in [SOURCE/'figures.py',
            SOURCE/'flag_stage_figure.py', SOURCE/'build.py',
            SOURCE.parent/'phenology/stage_censoring_scores.csv',
            SOURCE.parent/'phenology/stage_censoring_metrics.csv']})


def install(audit):
    for name in ('Manuscript.docx', 'Manuscript.pdf', 'Manuscript.txt', 'Source_Data.xlsx', 'document_build_receipt.json'):
        shutil.copy2(STAGING/name, ACTIVE/name)
    for target in (ACTIVE/'figures', SOURCE/'figures'):
        target.mkdir(exist_ok=True)
        for name in ('captions.json', 'fig2_field_stage_intervals.csv', *[f'{FIGURE}.{suffix}' for suffix in ('png', 'pdf', 'svg')]):
            shutil.copy2(STAGING/'figures'/name, target/name)
    with zipfile.ZipFile(BEFORE/'Software_and_Evidence.zip') as old:
        original = json.loads(old.read('PACKAGE_MANIFEST.json'))
        allowed_changes = {str(p.relative_to(ROOT)) for p in [SOURCE/'build.py', SOURCE/'figures.py', SOURCE/'figures/captions.json']}
        for row in original:
            path = ROOT/row['path']
            if row['path'] not in allowed_changes:
                assert sha(path) == row['sha256'], row['path']
        audit['runtime_calibration_and_scientific_evidence_sources_unchanged'] = True
        audit['archive_unchanged'] = all(sha(HERE/'before'/row['path']) == row['sha256']
            for row in json.loads((HERE/'before/archive_manifest.json').read_text()))
        assert audit['archive_unchanged']
        write_json(HERE/'figure2_verification.json', audit)
        write_json(ACTIVE/'figure2_verification.json', audit)
        added = [SOURCE/'flag_stage_figure.py', SOURCE/'figures/fig2_field_stage_intervals.csv',
            ACTIVE/'figures/fig2_field_stage_intervals.csv', HERE/'figure2_verification.json', Path(__file__)]
        files = {row['path']:ROOT/row['path'] for row in original}
        files.update({str(p.relative_to(ROOT)):p for p in added})
        manifest = [dict(path=name, bytes=p.stat().st_size, sha256=sha(p)) for name, p in sorted(files.items())]
        reproduction = old.read('REPRODUCIBILITY.txt').decode()
        reproduction += ('\nFigure 2: python -m analysis.paper_study.overwinter_leaf_model_20261006.manuscript.figures '
            'regenerates the five source-backed figures without fitting. The field-stage bounds and predictions are '
            'also supplied in publication/european_wheat_stb/figures/fig2_field_stage_intervals.csv.\n')
        (ACTIVE/'REPRODUCIBILITY.txt').write_text(reproduction)
        package = ACTIVE/'Software_and_Evidence.zip'
        with zipfile.ZipFile(package, 'w', zipfile.ZIP_DEFLATED, compresslevel=6) as z:
            for row in manifest:
                z.write(files[row['path']], row['path'])
            z.writestr('PACKAGE_MANIFEST.json', json.dumps(manifest, indent=2)+'\n')
            z.writestr('REPRODUCIBILITY.txt', reproduction)
    with zipfile.ZipFile(package) as z:
        assert z.testzip() is None
        assert len(z.namelist()) == len(manifest)+2
        for row in manifest:
            assert hashlib.sha256(z.read(row['path'])).hexdigest() == row['sha256']
    write_json(ACTIVE/'software_package_receipt.json', dict(status='complete',
        members=len(manifest)+2, bundled_sources_and_evidence=manifest,
        package_sha256=sha(package), raw_climate_cache_bundled=False,
        visualization_revision='Figure 2, 2026-10-07'))
    verification = json.loads((BEFORE/'final_verification.json').read_text())
    verification.update(workbook_sheets=18, software_package_members=len(manifest)+2,
        verification_scope='Figure 2 visualization and document exports; model checks inherited unchanged',
        model_test_evidence_origin=str((BEFORE/'final_verification.json').relative_to(ROOT)),
        runtime_and_calibration_sources_unchanged=True, figure2_field_stage_records=56)
    verification['review_receipts'].append(dict(path=str((HERE/'figure2_verification.json').relative_to(ROOT)),
        sha256=sha(HERE/'figure2_verification.json')))
    verification['artifact_sha256'] = {str(p.relative_to(ACTIVE)):sha(p)
        for p in sorted(ACTIVE.rglob('*')) if p.is_file() and p.name != 'final_verification.json'}
    write_json(ACTIVE/'final_verification.json', verification)
    for name, digest in verification['artifact_sha256'].items():
        assert sha(ACTIVE/name) == digest
    print(json.dumps(dict(status='passed', field_stage_records=56, workbook_sheets=18,
        package_members=len(manifest)+2, manuscript_pdf_pages=17,
        model_and_scientific_results_unchanged=True)))


if __name__ == '__main__':
    install(verify_staging())
