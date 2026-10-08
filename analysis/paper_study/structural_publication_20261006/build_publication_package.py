"""Render source-backed methods, results, tables and scientific figures."""

from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import shutil
import sys
import zipfile

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT=Path(__file__).resolve().parents[3]
sys.path.insert(0,str(ROOT))
from model.seasonal_septoria.structural import CanopyParameters, simulate_canopy

HERE=Path(__file__).parent
DEST=ROOT/'publication/structural_model_v6_20261006'
METRICS=HERE/'nested_summary/severity_metrics.csv'
LABELS={'original_seir':'Original compartment model','daily_or_canopy':'Daily exposure canopy',
    'duration_canopy':'Duration exposure canopy','fixed_structural_ensemble':'Canopy ensemble',
    'weighted_training_mean':'Training mean','leaf_rank_training_mean':'Leaf-rank mean',
    'weather_development_ridge':'Weather/development ridge'}
COLORS={'original_seir':'#777777','daily_or_canopy':'#84B3BA','duration_canopy':'#DAAD6A',
    'fixed_structural_ensemble':'#14696E','weighted_training_mean':'#C3C3C3',
    'leaf_rank_training_mean':'#A6A6A6','weather_development_ridge':'#9B6245'}


def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def row(frame,kind,model,upper=False):
    selected=frame.loc[frame.evaluation_kind.eq(kind)&frame.report_scope.eq('pooled')&frame.model.eq(model)
        &frame.endpoint.eq('final_numeric_assessment')&frame.leaf_scope.eq('upper_three' if upper else 'all_ordinal_leaves')]
    assert len(selected)==1
    return selected.iloc[0]


def save_figure(fig,name):
    fig.savefig(DEST/'figures'/f'{name}.pdf',bbox_inches='tight')
    fig.savefig(DEST/'figures'/f'{name}.png',dpi=300,bbox_inches='tight')
    plt.close(fig)


def figures(metrics,benchmarks,predictions,onsets):
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'axes.spines.top':False,
        'axes.spines.right':False,'pdf.fonttype':42,'ps.fonttype':42})
    fig,axes=plt.subplots(1,2,figsize=(8,3.1),layout='constrained')
    development=np.linspace(200,1050,500)
    thresholds={r['BBCH']:r['Cumulative_t_pp_v_GDD'] for r in json.loads((ROOT/'process_model/parameters/calibration.json').read_text())['thresholds']}
    flag=thresholds[31]+.3*(thresholds[51]-thresholds[31])
    for rank in (0,1,2,4):
        birth=max(thresholds[10],flag-rank*120)
        area=np.clip((development-birth)/100,0,1)
        axes[0].plot(development,area,label=f'F{rank+1}')
    axes[0].set(xlabel='Transferred development index (effective units)',ylabel='Normalized unfolding capacity',
        title='a  Independent emergence and expansion',ylim=(-.03,1.06))
    axes[0].legend(frameon=False,ncol=4,loc='lower right')
    days=70;t=np.full((1,days),18.);h=np.full_like(t,90.);rain=np.zeros_like(t)
    active=np.ones((1,days,1),bool);renewal=np.zeros_like(active,float)
    trajectory=simulate_canopy(t,h,rain,active,renewal,CanopyParameters(initiation=.045669,amplification=0,latent_days=20))
    for column,label,color in [(0,'Unaffected capacity','#8AA5BD'),(1,'Pending expression','#D3A359'),(2,'Expressed score','#14696E')]:
        axes[1].plot(np.arange(days+1),trajectory.state[0,:,0,column],label=label,color=color)
    axes[1].axvline(20,ls=':',color='black',lw=.8)
    axes[1].set(xlabel='Days under illustrative constant forcing',ylabel='Normalized effective state',
        title='b  Thermal queue and delayed expression',ylim=(-.03,1.06))
    axes[1].legend(frameon=False,fontsize=8)
    save_figure(fig,'figure1_canopy_structure')
    models=['weighted_training_mean','leaf_rank_training_mean','original_seir','weather_development_ridge','fixed_structural_ensemble']
    both=pd.concat([metrics,benchmarks.loc[~benchmarks.model.eq('fixed_structural_ensemble')]],ignore_index=True)
    fig,axes=plt.subplots(1,2,figsize=(8,3.6),layout='constrained')
    for axis,upper,title in zip(axes,[False,True],['a  Last numerical assessment, all leaves','b  Last numerical assessment, upper three']):
        for j,kind in enumerate(['spatial','forward']):
            offset=(j-.5)*.3
            values=[row(both,kind,model,upper).rmse for model in models]
            axis.plot(values,np.arange(len(models))+offset,'o' if kind=='spatial' else 's',
                color='#14696E' if kind=='spatial' else '#9B6245',label='Spatial' if kind=='spatial' else 'Forward year')
        axis.set(yticks=np.arange(len(models)),yticklabels=[LABELS[x] for x in models],
            xlabel='RMSE (percentage points)',title=title,xlim=(26,38))
        axis.invert_yaxis();axis.grid(axis='x',color='#E5E5E5',lw=.7)
    axes[1].legend(frameon=False,loc='lower right')
    save_figure(fig,'figure2_nested_severity')
    fig,axes=plt.subplots(1,2,figsize=(7.4,3.4),layout='constrained')
    for axis,kind,label in zip(axes,['spatial','forward'],['a  Spatial location folds','b  Forward year, novel locations']):
        frame=predictions.loc[predictions.evaluation_kind.eq(kind)&predictions.model.eq('fixed_structural_ensemble')]
        frame=frame.sort_values('date').groupby(['field_id','endpoint_series']).tail(1)
        axis.scatter(frame.value,frame.predicted_percent,s=10,alpha=.3,color='#14696E',edgecolors='none')
        axis.plot([0,100],[0,100],'--',color='#666666',lw=.8)
        axis.set(xlabel='Recorded INFECT percentage',ylabel='Predicted percentage',title=label,
            xlim=(-2,102),ylim=(-2,102),aspect='equal')
    save_figure(fig,'figure3_terminal_scores')
    fig,axes=plt.subplots(1,2,figsize=(7.6,3.1),layout='constrained')
    for axis,kind,label in zip(axes,['spatial','forward'],['a  Spatial','b  Forward year']):
        for model in ['original_seir','fixed_structural_ensemble']:
            frame=onsets.loc[onsets.evaluation_kind.eq(kind)&onsets.model.eq(model)&onsets.censoring.eq('interval')].sort_values('cutoff_percent')
            axis.plot(np.arange(3),100*frame.weighted_compatible,'o-',color=COLORS[model],label=LABELS[model])
            if model=='fixed_structural_ensemble':
                for i,r in enumerate(frame.itertuples()):axis.text(i,100*r.weighted_compatible+4,f'n={r.n}',ha='center',fontsize=8)
        axis.set(xticks=np.arange(3),xticklabels=['0.1','1','5'],xlabel='Recorded-score onset threshold (%)',
            ylabel='Interval-compatible predictions (%)',title=label,ylim=(0,85))
    axes[1].legend(frameon=False,fontsize=8,loc='upper right')
    save_figure(fig,'figure4_interval_onset')


def main():
    if DEST.exists():raise FileExistsError('Use a new publication package.')
    for folder,receipt in [(HERE/'nested_summary','verification.json'),(HERE/'nested_benchmarks','receipt.json'),(HERE/'full_refit','receipt.json')]:
        record=json.loads((folder/receipt).read_text())
        assert record['status'] in ('passed','complete')
        for name,digest in record.get('output_sha256',{}).items():assert sha(folder/name)==digest
    DEST.mkdir(parents=True);(DEST/'figures').mkdir();(DEST/'tables').mkdir()
    metrics=pd.read_csv(METRICS)
    benchmarks=pd.read_csv(HERE/'nested_benchmarks/severity_metrics.csv')
    differences=pd.read_csv(HERE/'nested_summary/paired_location_bootstrap.csv')
    onset=pd.read_csv(HERE/'nested_summary/onset_metrics.csv')
    predictions=pd.read_parquet(HERE/'nested_summary/all_out_of_fold_predictions.parquet')
    figures(metrics,benchmarks,predictions,onset)
    assumptions=[
        ['Observation','Source INFECT percentage; denominator unspecified','Effective score, without identified incidence or lesion-area interpretation'],
        ['Canopy initiation','Flag threshold C31+0.3(C51-C31); interval120','Transferred development index; fitted timing is not a measured phyllochron'],
        ['Expansion','100 transferred development units','New unaffected capacity dilutes pending and expressed fractions'],
        ['Expression delay','20 reference days at18°C in both deployment components','Thermal FIFO delay; selected effective waiting time'],
        ['Amplification','Fixed zero after initial candidate comparison','Data do not support an additional fitted amplification coefficient'],
        ['Weather response','Equal-weight daily-OR and inferred-duration operators','Averaging retains unresolved structural assumptions'],
        ['Phenology and sowing','Fixed transferred TPV; winter rainfed calendar scenario','Real sowing dates, cultivar and wheat type unavailable'],
        ['Senescence','No disease recovery fitted','Senescent tissue and green leaf area unobserved'],
        ['Prediction uncertainty','Component range is structural spread','It is not a calibrated prediction interval'],
    ]
    final=[]
    for kind in ['spatial','forward']:
        for model in ['original_seir','weighted_training_mean','leaf_rank_training_mean','weather_development_ridge','fixed_structural_ensemble']:
            r=row(metrics if model in ['original_seir','fixed_structural_ensemble'] else benchmarks,kind,model)
            final.append([kind.title(),LABELS[model],f'{r.rmse:.2f}',f'{r.mae:.2f}',f'{r.bias:+.2f}',f'{r.weighted_r2:.3f}'])
    pd.DataFrame(assumptions,columns=['Process','Specification','Interpretation']).to_csv(DEST/'tables/model_assumptions.csv',index=False)
    pd.DataFrame(final,columns=['Evaluation','Model','RMSE','MAE','Bias','Weighted_R2']).to_csv(DEST/'tables/nested_final_severity.csv',index=False)
    for source,name in [(METRICS,'all_nested_metrics.csv'),(HERE/'nested_summary/paired_location_bootstrap.csv','paired_model_differences.csv'),
        (HERE/'nested_benchmarks/paired_complete_location_history_differences.csv','paired_benchmark_differences.csv'),
        (HERE/'nested_summary/onset_metrics.csv','onset_by_censoring.csv'),(HERE/'nested_summary/high_score_metrics.csv','high_score_metrics.csv')]:
        shutil.copyfile(source,DEST/'tables'/name)
    sections=[]
    def heading(text):sections.append(('heading',text))
    def paragraph(text):sections.append(('paragraph',text))
    def table(headers,rows):sections.append(('table',(headers,rows)))
    heading('A development-indexed delayed canopy model for seasonal wheat Septoria scores')
    paragraph('Seasonal untreated wheat Septoria observations support a leaf-resolved effective scoring model with explicit emergence, unfolding dilution and delayed symptom expression. The revised structure improves onset compatibility substantially and yields a modest reduction in final all-leaf score error under retrospective location-grouped evaluation. Final-score errors remain approximately30 percentage points, and superiority to a weather/development ridge baseline is inconsistent. The source infection-percentage denominator, cultivar, true sowing date and genetic wheat type remain unidentified; consequently, the supported estimand is the recorded assessment score rather than an independently measured disease-area fraction.')
    heading('Observation basis and estimand')
    paragraph('The combined development archive contains2,305 numerical assessments from218 untreated field-seasons at127 rounded-coordinate locations, representing168 coordinate-years. BASF contributes584 assessments from73 field-seasons and Corteva contributes1,721 from145. Two modeled locations occur in both sources and retain a unified location identity. The primary endpoint is the last numerical assessment within each source ordinal-leaf series, giving906 endpoints. It is a leaf-specific terminal recorded assessment; different leaves within a field commonly have different last-assessment dates. It does not define a simultaneous final canopy average or an observed maturity endpoint.')
    paragraph('INFECT percentage records have an unspecified scoring denominator. Source multiplicities and repeated rows do not identify a binomial sample size. Normalized model states describe effective assessment capacity; their values cannot be interpreted as the proportion of infected plants, necrotic leaf area or pycnidial density. Image-based wheat STB research distinguishes lesion-area damage from pathogen reproduction and uses conditional sampling of diseased leaves (Karisto et al.,2018). These measurement domains require distinct observation operators. Independent pycnidial datasets therefore remain separate from the INFECT calibration.')
    heading('Development-indexed canopy and delayed score expression')
    paragraph('Daily mean and maximum temperature drive the transferred temperature–photoperiod–vernalization development index. Rainfed winter-wheat calendar scenarios supply sowing dates because trial sowing and cultivar information are absent. Development thresholds C10, C31, C51 and C85 retain their frozen phenology values. For ordinal rank k=0,...,6, the effective emergence threshold is Bk=max[C10,C31+f(C51−C31)−kI], where f controls flag timing and I independently controls spacing. The first leaf rank is the flag leaf. Availability uses the previous day’s accumulated development index, preventing within-day development from making new tissue available retrospectively.')
    paragraph('For an available adult leaf, normalized unfolding capacity is A(t)=min[1,max(0,(G(t)−Bk)/E)], where G is accumulated development and E is the expansion interval. Growth dilution is r(t)=max[0,A(t)−A(t−1)]/A(t). Pending and expressed fractions are multiplied by1−r(t), and the remaining capacity becomes unexpressed. This conservation rule represents the dilution of existing affected capacity by new unaffected growth; A is not measured leaf area. The juvenile reservoir retains its inherited300°C-day renewal assumption, but the selected zero-amplification model does not use it as a cross-leaf transmission source. Adult natural senescence is unobserved and is not fitted as disease recovery or recoded as a zero score. Canopy growth, healthy-area duration and cultivar physiology affect disease consequences, so predicted scores alone do not establish yield losses (van den Berg et al.,2017).')
    paragraph('Each available leaf has unexpressed S, pending Q and expressed D states with S+Q+D=1. Exposure consumes S through the effective hazard aW(t). Each newly pending cohort enters a FIFO queue and becomes expressed only when accumulated reference heat since entry reaches L: the increment is max[Tmean(t),0]dt/18. This deterministic thermal waiting rule removes the original Erlang delay’s near-zero expression tail. It does not identify a laboratory latent period or a dated infection event. A separate candidate hazard b·max(Tmean,0)D/18 permits further expressed-score amplification; b is fixed at zero in the retained procedure because the initial finite structural comparison did not favor fitting it. Infection-route identities and field inoculum loads are not inferred.')
    paragraph('The daily-OR exposure is W=exp[−0.5((Tmean−18)/8)²]·{1−(1−H)(1−R)}, with H=logistic[(RHmean−85)/5] and R=1−exp(−P/2). The second operator derives24 symmetric temperature bins around daily mean temperature, with amplitude Tmax−Tmean. Constant within-day vapour pressure is chosen numerically so that saturated, clipped bin humidities reproduce mean RH. Bins at RH≥90% contribute humid exposure. The assumed rain fraction is1−exp[−P/(24v)], with v estimated as the median daily precipitation/rain-hour ratio on positive-rain training days warmer than2°C. Expected rain/humidity overlap assumes temporal independence. Gaussian bin-temperature responses weight the resulting exposure. These synthetic duration indices do not reconstruct measured hourly weather or canopy leaf wetness.')
    paragraph('The deployed predictor averages the independently selected daily-OR and duration-proxy canopy scores with fixed weights0.5 and0.5. The weights were registered before reading pooled nested scores and were not estimated from outer outcomes. Each component carries its own frozen preprocessing and coefficients. Component spread quantifies only disagreement between these two structural alternatives. It does not include observation residuals, parameter estimation, cultivar effects or phenology uncertainty and is not a prediction interval.')
    table(['Process','Specification','Interpretation'],assumptions)
    heading('Calibration, selection and retrospective validation')
    paragraph('The initial finite structural experiment used only BASF2017–2018 calibration outcomes:48 combinations of flag-timing fraction, coupled or independent leaf spacing,20- or30-reference-day delay, and absent or fitted amplification. Further finite comparisons evaluated the duration operator and an identity versus nonlinear observation link. The nonlinear link failed to improve calibration final-leaf performance and remained excluded. Expansion intervals0,100 and200 units were examined with frozen coefficients; this is parameter-fixed sensitivity, not a separately fitted ablation.')
    paragraph('The nested retained procedure crosses flag fractions0.3,0.6 and0.9; independent intervals80,120 and160; and delays20 and30, with expansion100 and amplification zero. Three inner folds keep complete location histories together across years and both sources. Each candidate minimizes an equal mixture of hierarchical all-assessment and final-assessment squared error, using deterministic starts and declared coefficient bounds. Inner pooled final-leaf RMSE selects the canopy assumptions. The original compartment comparator independently selects10-,20- or30-day delay with its secondary coefficient fitted, on the same fields and loss weights. Daily initiation is bounded at0.1 and duration initiation at1.0; comparisons therefore assess complete fitted procedures rather than isolate one weather operator.')
    paragraph('Five outer spatial folds hold out all histories from a location. Four forward evaluations reserve2016,2017,2018 and2019, train only on preceding years, and additionally require test locations to be absent from all preceding training years in both sources. Forward evaluation contains1,005 assessments from110 field-seasons at93 novel locations, yielding449 final-leaf endpoints. The optimizer receives physically subset training fields; held-out prediction inputs have all disease values removed. No observed disease initialization, observed crop stage, location identity or future meteorology enters prediction. Outer fits and selections are frozen before outer scoring.')
    paragraph('Historical BASF2019 and Corteva outcomes had already informed preceding model development. The nested evaluation is therefore retrospective development evidence conditional on the investigated model family, rather than an untouched confirmatory test. Exact memberships, source hashes, candidate traces, frozen fits and independent score calculations provide an auditable separation of fitting, selection and scoring within each run; they cannot remove prior architectural exposure to those outcomes.')
    paragraph('Metrics give equal total weight to coordinate-years, equal weight to source field units within each coordinate-year, equal weight to leaf series within a field, and equal weight to dates within a series for the all-assessment endpoint. Final-assessment scoring retains one date per series. Paired percentile intervals use5,000 draws of complete location histories, preserving repeated years and cross-source dependence. They condition on frozen out-of-fold predictions and do not repeat selection or refitting. Simple weighted training means, leaf-rank means and a weather/development ridge model use the identical nested memberships and loss weights. The finite empirical benchmark plan was added after structural nested results were available and frozen before baseline fitting and scoring. Ridge penalties0.1,1,10,100 and1000 are selected by pooled inner validation final-leaf RMSE within each outer training partition; feature scaling is fitted on training subsets and predictions are clipped to the declared0–100 bounds. Its inherited7-,21- and60-day weather windows, original cohort timing, leaf rank and transferred-development features restrict the empirical comparison to that declared family. All methods use forcing through the assessment date, so this evaluation does not establish a specified advance-warning forecast horizon.')
    heading('Final recorded severity and comparator performance')
    spatial=row(metrics,'spatial','fixed_structural_ensemble');forward=row(metrics,'forward','fixed_structural_ensemble')
    paragraph(f'Spatial all-leaf final-score RMSE is{spatial.rmse:.2f} percentage points for the canopy ensemble, compared with31.24 for the independently refitted original compartment procedure. The paired difference is−1.01 points (95% location-bootstrap interval−1.93 to−0.14). MAE falls from26.14 to{spatial.mae:.2f}, and weighted R² increases from0.192 to{spatial.weighted_r2:.3f}. Forward all-leaf RMSE falls from30.72 to{forward.rmse:.2f}; its paired difference is−0.74 points (−2.54 to0.98). Forward bias increases from+5.41 to+7.72 points, so lower total error does not imply better mean calibration.')
    paragraph('The spatial training-mean and leaf-rank-mean RMSE values are35.78 and35.49 points. The weather/development ridge achieves29.50, compared with30.23 for the canopy ensemble; the canopy-minus-ridge difference is+0.74 (−0.76 to2.12). Forward ridge RMSE is30.84, compared with29.98 for the canopy ensemble; the difference is−0.86 (−3.35 to1.69). These comparisons support spatial predictive information beyond simple means, while consistent superiority to the declared ridge family is unsupported.')
    paragraph('Spatial upper-three final-score RMSE is31.27 for both the original procedure and the canopy ensemble, with a paired difference near zero (−0.80 to0.76). Forward upper-three RMSE decreases from30.97 to29.76, a difference of−1.22 (−2.43 to−0.03). The spatial all-leaf improvement therefore does not establish a meaningful gain for the upper-leaf endpoint most directly relevant to crop production. Among spatial final observations at or above80%, ensemble mean underprediction remains43.97 points and RMSE45.39 points; the corresponding forward values are35.31 and38.69. Severe epidemic magnitude remains inadequately predicted.')
    table(['Evaluation','Model','RMSE','MAE','Bias','R²'],final)
    heading('Onset, leaf gradients and meteorological duration checks')
    paragraph('The timing diagnostic compares the predicted first threshold crossing with an assessment bracket: the last recorded score below the threshold precedes the first score at or above it, giving an exclusive lower and inclusive upper bracket. Interpreting this bracket as the first onset interval assumes that no earlier transient crossing occurred and subsequently fell below the threshold. Intermittent sampling and growth dilution do not establish that persistence assumption. Left-, interval- and right-censored series remain distinct. Missing predicted crossings remain in compatibility denominators. Compatibility for broadly left-censored series is weak timing evidence and is not combined with two-sided assessment-bracket results.')
    paragraph('At the0.1% threshold, spatial assessment-bracket compatibility increases from7.97% to54.87% (171 series), and forward compatibility increases from6.54% to62.99% (73 series). At1%, spatial compatibility increases from11.56% to51.32% and forward compatibility from18.56% to50.38%; at5%, the respective changes are15.42% to43.26% and22.90% to44.56%. Original unweighted median spatial bracket distances are13,12 and9 days early at those thresholds; ensemble medians are zero. These are separate retrospective timing diagnostics. They do not identify exact infection dates or calibrated occurrence probabilities.')
    paragraph('In the original2017–2018-only frozen transfer experiment, same-date adjacent-leaf gradients provide an additional structural diagnostic unaffected by asynchronous terminal dates. The strict Corteva subset contains1,110 adjacent-leaf/date pairs from139 field-seasons. Its mean recorded lower-minus-upper gradient is16.88 points, compared with4.48 for the original predictor and11.08 for the fixed canopy ensemble. Gradient RMSE decreases from21.54 to18.44. This supports a more realistic between-leaf score progression, conditional on the same reused source evidence; residual gradients remain substantial.')
    paragraph('Atmospheric duration checks compare inferred indices with archived hourly-derived reanalysis aggregates of high humidity and positive rain. On26,463 strict Corteva forcing days, inferred humid hours have RMSE3.37h and R²0.783; rain hours have RMSE1.97h and R²0.806. These checks concern consistency with the same underlying reanalysis forcing, rather than independent field measurements of weather or leaf wetness, and use the original calibration-only duration fit. They cannot validate within-day rain/dew overlap or stationary future-climate rainfall intensity.')
    heading('Phenology evidence, numerical verification and deployment fit')
    paragraph('A separate phenology diagnostic fitted a single development multiplier to BASF2017–2018 stage intervals, retaining complete and informative single-ended constraints while excluding reversed, conflicting or chronologically inconsistent intervals. The selected grid optimum isq=0.9725; minimum-loss grid values span0.9500–0.9725 and the within-one-day objective grid range is0.8525–1.155. The finite grid has401 values at spacing0.0025. These are objective ranges on the declared grid, not continuous optimum bounds or confidence intervals. Only20 calibration fields provide two-sided heading constraints; flag-appearance and full-unfolding constraints are sparse. Forward overall stage loss does not improve. The primary disease predictor retainsq=1; these data do not identify measured leaf emergence or justify a universally superior phenology correction. Flag visibility and full unfolding are separate BBCH37 and39 events (Aarhus University and SEGES Innovation,2026).')
    paragraph('Quarter-day competing hazards preserve nonnegative states and normalized total capacity. Frozen component experiments have maximum conservation error below9×10⁻¹⁶. Halving the integration step changes individual archived assessment predictions by at most1.05 percentage points and by approximately0.04–0.06 points on average. Independent verification reproduces978 nested metric and prediction checks within1.1×10⁻¹⁴; additional benchmark verification reproduces384 metric and432 paired-comparison calculations. Invalid assessment coordinates and incomplete duration inputs fail explicitly. Prediction never reads target disease values.')
    paragraph('After retrospective evaluation, each component was selected again on three location-grouped folds of all available development data and refitted to all2,305 assessments. Both select f=0.3, I=120, E=100 and L=20, with amplification zero. Effective initiation is0.0456691 for daily-OR and0.0772091 for the duration proxy; the latter freezes rain intensity at0.460mmh−1. The deployment model identity iscanopy_score_ensemble_v6_20261006. All-data predictions are fit diagnostics and do not constitute a further validation set. The default runtime exactly reproduces all2,305 archived deployment ensemble predictions after removing observed disease values.')
    heading('Interpretation and limits of climate extrapolation')
    paragraph('The model supports a reproducible association between declared crop/weather scenarios and the archived INFECT score, with substantial improvement in delayed onset representation and modest spatial final-score skill. Weather–disease associations and model searches require explicit control of iterative selection and distinguish occurrence from conditional severity (Pietravalle et al.,2003). The remaining30-point final-score error, high-score underprediction and uncertain comparative advantage limit claims of precise absolute severity forecasting.')
    paragraph('Unknown scoring denominators, missing cultivar and sowing metadata, transferred phenology, assumed winter-wheat type, unmeasured expansion and senescence, unobserved field inoculum and synthetic wet-duration assumptions remain sources of structural uncertainty. The evidence does not identify pathogen transmission routes, whole-season infection probabilities, causal climate effects or yield loss. New independent seasons or a source with verified scoring protocols and cultivar/sowing observations are needed for confirmatory validation. Climate extrapolation additionally requires replacement simulations, checking of meteorological and phenological support, and propagation of these assumptions. The existing continental climate tables retain their original model identity; the structural evaluation does not validate their numerical conclusions under the deployed canopy model.')
    heading('References')
    paragraph('Karisto,P., Hund,A., Yu,K., Anderegg,J., Walter,A., Mascher,F., McDonald,B.A., and Mikaberidze,A.2018. Ranking quantitative resistance to Septoria tritici blotch in elite wheat cultivars using automated image analysis. Phytopathology108:568–581. https://doi.org/10.1094/PHYTO-04-17-0163-R')
    paragraph('van den Berg,F., Paveley,N.D., Bingham,I.J., and van den Bosch,F.2017. Physiological traits determining yield tolerance of wheat to foliar diseases. Phytopathology107:1468–1478. https://doi.org/10.1094/PHYTO-07-16-0283-R')
    paragraph('Pietravalle,S., Shaw,M.W., Parker,S.R., and van den Bosch,F.2003. Modeling of relationships between weather and Septoria tritici epidemics on winter wheat: a critical approach. Phytopathology93:1329–1339. https://doi.org/10.1094/PHYTO.2003.93.10.1329')
    paragraph('Aarhus University, Department of Agroecology, and SEGES Innovation.2026. Crop Protection Online: crop growth stages, BBCH scale. Accessed6October2026. https://plantevaernonline.dlbr.dk/cp/SeasonPlan/CropScale.asp?id=djf&language=en')
    captions=[('figure1_canopy_structure','Figure1. Declared canopy timing and delayed score states. The emergence/unfolding example uses deployment timing, with effective development units rather than measured leaf area. The single-leaf state example uses constant18°C,90% mean humidity, zero rain and no growth dilution; it is illustrative and is not a field validation.'),
        ('figure2_nested_severity','Figure2. Final numerical assessment RMSE under matched retrospective nested spatial and forward-year evaluations. All models use identical outer and inner memberships. Upper-three and all-leaf results remain separate. Forward test locations are absent from all preceding training years in both sources.'),
        ('figure3_terminal_scores','Figure3. Recorded and out-of-fold ensemble final-leaf scores. Points represent906 spatial and449 forward endpoints and are displayed without hierarchical weighting; tabled metrics apply equal coordinate-year, field and leaf weights. The diagonal marks score equality.'),
        ('figure4_interval_onset','Figure4. Compatibility with observed assessment brackets at three score thresholds. Only two-sided series are displayed, with equal coordinate-year/field/series weighting. Labels give raw series counts; left- and right-censored results are retained in the supplementary table. A first-onset interpretation assumes no earlier transient crossing. These are descriptive timing diagnostics, not exact infection-date or probability validation.')]
    document=Document();normal=document.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(7)
    for section in document.sections:
        section.top_margin=Inches(.65);section.bottom_margin=Inches(.65)
        section.left_margin=Inches(.65);section.right_margin=Inches(.65)
    font_path=Path(matplotlib.get_data_path())/'fonts/ttf/DejaVuSans.ttf'
    pdfmetrics.registerFont(TTFont('ModelSans',str(font_path)))
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle('ModelBody',fontName='ModelSans',fontSize=9,leading=12.5,spaceAfter=6.5))
    styles.add(ParagraphStyle('ModelHeading',fontName='ModelSans',fontSize=12,leading=16,spaceBefore=10,spaceAfter=7))
    styles.add(ParagraphStyle('ModelCell',fontName='ModelSans',fontSize=7,leading=10))
    story=[];plain=[]
    def typography(text):
        protected={}
        def protect(match):
            token='TOKEN'+chr(65+len(protected))+'TOKEN'
            protected[token]=match.group(0)
            return token
        text=re.sub(r'https?://\S+|canopy_score_ensemble_v6_20261006|\bC(?:10|31|51|85)\b',protect,text)
        text=re.sub(r'(?<=[A-Za-z])(?=\d)', ' ', text)
        text=re.sub(r'(?<=\d)(?=[A-Za-z])', ' ', text)
        text=re.sub(r'(?<=[,;])(?=[A-Za-z])',' ',text)
        text=re.sub(r',(?=\d{4}[). ])',', ',text)
        text=re.sub(r',(?=\d{1,2}[%°])',', ',text)
        text=re.sub(r'(?<=[A-Za-z]\.)(?=\d{4}\b)',' ',text)
        text=text.replace('isTOKEN','is TOKEN').replace('mmh−1','mm h⁻¹')
        for token,value in protected.items():text=text.replace(token,value)
        return text
    def pdftext(text):
        from xml.sax.saxutils import escape
        return escape(text)
    for kind,content in sections:
        if kind in ('heading','paragraph'):
            text=typography(content)
            text=text.replace('percentage30','percentage 30')
            if kind=='heading':document.add_heading(text,0 if not plain else 1)
            else:document.add_paragraph(text)
            story.append(Paragraph(pdftext(text),styles['ModelHeading' if kind=='heading' else 'ModelBody']))
            plain.append(text)
        else:
            headers,rows=content
            doc_table=document.add_table(rows=1,cols=len(headers));doc_table.style='Table Grid'
            for cell,label in zip(doc_table.rows[0].cells,headers):cell.text=label
            for values in rows:
                for cell,text in zip(doc_table.add_row().cells,values):cell.text=text
            grid=[[Paragraph(pdftext(typography(str(x))),styles['ModelCell']) for x in values] for values in [headers,*rows]]
            widths=[75,140,295] if len(headers)==3 else [60,170,65,65,65,65]
            pdf_table=Table(grid,colWidths=widths,repeatRows=1,hAlign='LEFT')
            pdf_table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E9F1F1')),('VALIGN',(0,0),(-1,-1),'TOP'),
                ('BOTTOMPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),5),('LINEBELOW',(0,0),(-1,0),.6,colors.HexColor('#14696E'))]))
            story.extend([Spacer(1,5),pdf_table,Spacer(1,9)])
            plain.append('\n'.join(['\t'.join(headers),*['\t'.join(values) for values in rows]]))
    for name,caption in captions:
        document.add_page_break();document.add_picture(str(DEST/'figures'/f'{name}.png'),width=Inches(7))
        document.add_paragraph(typography(caption))
        story.extend([PageBreak(),Image(str(DEST/'figures'/f'{name}.png'),width=510,height=510*plt.imread(DEST/'figures'/f'{name}.png').shape[0]/plt.imread(DEST/'figures'/f'{name}.png').shape[1]),
            Spacer(1,10),Paragraph(pdftext(typography(caption)),styles['ModelBody'])])
        plain.append(typography(caption))
    document.save(DEST/'Model_Methods_and_Validation.docx')
    SimpleDocTemplate(str(DEST/'Model_Methods_and_Validation.pdf'),pagesize=(612,792),leftMargin=51,rightMargin=51,topMargin=42,bottomMargin=42).build(story)
    (DEST/'Model_Methods_and_Validation.txt').write_text('\n\n'.join(plain)+'\n')
    shutil.copyfile(ROOT/'configurations/wheat_stb/legacy/publication_model.json',DEST/'publication_model.json')
    reproducibility='''Structural canopy model v6: reproducibility

Model identity: canopy_score_ensemble_v6_20261006.
Observation: infection_percent_unspecified_basis.
Python forecast adapter: calibration.seasonal_septoria.publication_model.predict_publication_model.
Required forcing: mean and maximum temperature in degreesCelsius, mean relative humidity in percent, and daily precipitation inmm. The duration component requires maximum temperature and its frozen rain-intensity preprocessing. The development index and thresholds must match the declared crop/calendar/phenology scenario. Target disease values are not used by prediction.

Frozen model: publication_model.json; deployment coefficients fitted to all2,305 available development assessments after retrospective nested evaluation. Those all-data fit outputs are not validation evidence. The two components have weights0.5 each. The returned structural range is not a prediction interval.

Archived statistical evidence: tables/all_nested_metrics.csv, paired_model_differences.csv, paired_benchmark_differences.csv, onset_by_censoring.csv and high_score_metrics.csv. The paired intervals resample complete location histories while holding fitted predictions fixed. No untouched confirmatory data are available.

Runtime use in the repository:
    from calibration.seasonal_septoria.publication_model import predict_publication_model
    result = predict_publication_model(data, accumulation, thresholds, return_daily=True)
    predicted = result['predicted_percent']

Field preprocessing example for archived data:
    from analysis.paper_study.structural_evaluation.run import load_inputs, DEFAULT_PATHS
    data, weather, accumulation, thresholds = load_inputs(DEFAULT_PATHS)
    data.targets['value'] = float('nan')
    result = predict_publication_model(data, accumulation, thresholds)

Finite evaluation commands (new output directories are required):
    .venv/bin/python -m analysis.paper_study.structural_evaluation.run --weather-operator daily_or --output <new_daily_archive>
    .venv/bin/python -m analysis.paper_study.structural_evaluation.run --weather-operator duration_proxy --output <new_duration_archive>

The supplement preserves original archive identities. Existing continental climate results use the original compartment model; replacement climate simulations are required before applying these new coefficients to manuscript climate conclusions. Source data remain in their original local locations with archived checksums; the software package does not distribute third-party trial or climate records.
'''
    (DEST/'Reproducibility.txt').write_text(reproducibility)
    source_paths=[*sorted((ROOT/'calibration').rglob('*.py')),*sorted((ROOT/'calibration').rglob('*.txt')),
        *sorted((ROOT/'configurations/wheat_stb/legacy').glob('*.json')),*sorted((ROOT/'model/seasonal_septoria').glob('*.py')),*sorted((ROOT/'process_model').glob('*.py')),
        *sorted((ROOT/'process_model/_support').glob('*.py')),ROOT/'process_model/parameters/calibration.json',
        ROOT/'configurations/wheat_stb/legacy/publication_model.json',ROOT/'pytest.ini',
        *sorted((ROOT/'tests').glob('test_structural*.py')),ROOT/'tests/test_publication_model.py',ROOT/'tests/test_wetness_proxy.py',
        *sorted((HERE).glob('*.py')),*sorted((ROOT/'analysis/paper_study/structural_evaluation').glob('*.py')),
        ROOT/'analysis/paper_study/calibrate_seasonal.py',ROOT/'analysis/paper_study/refine_seasonal_accuracy.py',
        ROOT/'analysis/paper_study/run_empirical_benchmarks.py',ROOT/'analysis/paper_study/run_structural_canopy.py',
        ROOT/'analysis/paper_study/run_weather_canopy.py',HERE/'ensemble_protocol_before_nested_scoring.json',
        HERE/'default_runtime_verification.json']
    for folder in ['nested_summary','nested_benchmarks','full_refit']:
        source_paths.extend(p for p in (HERE/folder).glob('*') if p.is_file() and p.suffix in ('.json','.csv','.py'))
    for folder in ['nested_canopy_daily_20261006','nested_canopy_duration_20261006']:
        archive=ROOT/'analysis/paper_study'/folder
        source_paths.extend([archive/'receipt.json',archive/'configuration_before_fitting.json',archive/'all_outer_fits_frozen_before_scoring.json',archive/'nested_membership_before_fitting.csv'])
        source_paths.extend(archive.glob('outer_folds/*/frozen_*_outer_fit.json'))
    lock=ROOT/'publication/european_wheat_stb/requirements.lock.txt'
    if lock.exists():shutil.copyfile(lock,DEST/'requirements.lock.txt')
    with zipfile.ZipFile(DEST/'Reproducibility.zip','w',zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(set(source_paths)):archive.write(path,str(path.relative_to(ROOT)))
        for path in sorted((DEST/'tables').glob('*')):archive.write(path,str(path.relative_to(DEST)))
        for path in [DEST/'Reproducibility.txt',DEST/'publication_model.json',DEST/'requirements.lock.txt']:
            if path.exists():archive.write(path,path.name)
    receipt=dict(status='complete',model_id='canopy_score_ensemble_v6_20261006',untouched_test=False,
        historical_climate_manuscript_updated=False,default_model_sha256=sha(ROOT/'configurations/wheat_stb/legacy/publication_model.json'),
        source_sha256={str(p.relative_to(ROOT)):sha(p) for p in sorted(set(source_paths))},
        output_sha256={str(p.relative_to(DEST)):sha(p) for p in sorted(DEST.rglob('*')) if p.is_file()})
    (DEST/'package_receipt.json').write_text(json.dumps(receipt,indent=2)+'\n')
    print(json.dumps(dict(status='complete',output=str(DEST),figures=len(captions),formal_words=len(' '.join(plain).split()))))


if __name__=='__main__':main()
