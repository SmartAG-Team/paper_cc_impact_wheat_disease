"""Independent arithmetic, artifact and source-consistency checks."""
from pathlib import Path
import argparse
import hashlib
import json
import re
from zipfile import ZipFile
import numpy as np
import pandas as pd
from docx import Document
from openpyxl import load_workbook
import pymupdf
from .data import ROOT,HERE,GRID,DERIVED,CANOPY,SEVERITY,ONSET,FREQUENCY,YIELD,pooled,sha


def pooled_arithmetic():
    onset=pd.read_csv(DERIVED/'pooled_symptom_onset_membership.csv')
    means=onset.groupby(['coordinate_year','field_id','leaf_index'])[['model','benchmark']].mean()
    means=means.groupby(['coordinate_year','field_id']).mean().groupby('coordinate_year').mean().mean()
    for key in ['model','benchmark']:
        np.testing.assert_allclose(means[key],pooled('symptom_onset',key).value,rtol=0,atol=1e-12)
    severity=pd.read_csv(DERIVED/'pooled_severity_membership.csv')
    scores={}
    for key in ['model','benchmark']:
        error=severity[key]-severity.observed
        d=severity[['coordinate_year','field_id','leaf_index']].assign(squared=error**2,bias=error)
        x=d.groupby(['coordinate_year','field_id','leaf_index'])[['squared','bias']].mean()
        x=x.groupby(['coordinate_year','field_id']).mean().groupby('coordinate_year').mean().mean()
        scores[key+'_rmse']=float(np.sqrt(x.squared));scores[key+'_bias']=float(x.bias)
        for metric in ['rmse','bias']:
            np.testing.assert_allclose(scores[key+'_'+metric],pooled('severity',key+'_'+metric).value,rtol=0,atol=1e-12)
    assert len(severity)==2033 and severity.field_id.nunique()==188
    assert len(onset)==136 and onset.field_id.nunique()==78
    return dict(symptom_onset_days=means.to_dict(),severity=scores,assessment_records=2033,onset_histories=136)


def domain_arithmetic(directory=GRID,period='2071-2100'):
    directory=Path(directory)
    registry=pd.read_csv(directory/'full_landuse_cell_registry.csv')
    rows=pd.read_csv(directory/'environmental_region_changes_by_gcm.csv')
    country=pd.read_csv(directory/'country_changes_by_gcm.csv')
    metrics=[CANOPY,SEVERITY,ONSET,FREQUENCY,YIELD]
    observed=[];direct=True
    start=int(period[:4])
    paths=[ROOT/f'analysis/paper_study/nature_food_revision_20261007/continental_replay/annual_outputs/nasa/{model}/ssp585/{year}.parquet'
        for model in ['ACCESS-CM2','MPI-ESM1-2-HR','MRI-ESM2-0'] for year in list(range(1991,2021))+list(range(start,start+30))]
    if not all(p.exists() for p in paths):direct=False
    for model in ['ACCESS-CM2','MPI-ESM1-2-HR','MRI-ESM2-0']:
        if direct:
            numerator=np.zeros((len(registry),len(metrics)));denominator=np.zeros_like(numerator)
            for i in range(30):
                base=ROOT/f'analysis/paper_study/nature_food_revision_20261007/continental_replay/annual_outputs/nasa/{model}/ssp585'
                a=pd.read_parquet(base/f'{1991+i}.parquet',columns=['cell_id',*metrics])
                b=pd.read_parquet(base/f'{start+i}.parquet',columns=['cell_id',*metrics])
                assert np.array_equal(a.cell_id,registry.cell_id) and np.array_equal(b.cell_id,registry.cell_id)
                av=a[metrics].to_numpy();bv=b[metrics].to_numpy();valid=np.isfinite(av)&np.isfinite(bv)
                numerator+=np.where(valid,bv-av,0);denominator+=valid
        else:
            pairs=pd.read_parquet(directory/'full_grid_paired_changes_by_gcm.parquet')
            pairs=pairs[pairs.model.eq(model)&pairs.scenario.eq('ssp585')&pairs.period.eq(period)]
            numerator=np.empty((len(registry),len(metrics)));denominator=np.empty_like(numerator)
            for j,metric in enumerate(metrics):
                q=pairs[pairs.metric.eq(metric)].set_index('cell_id').loc[registry.cell_id]
                denominator[:,j]=q.valid_year_pairs
                numerator[:,j]=np.where(q.valid_year_pairs>0,q.change*q.valid_year_pairs,0)
        for name,mask,reference in [
            ('Europe',np.ones(len(registry),bool),rows[rows.environment_region.eq('Europe')]),
            ('Boreal',registry.environment_region.eq('Boreal').to_numpy(),rows[rows.environment_region.eq('Boreal')]),
            ('Mediterranean',registry.environment_region.eq('Mediterranean').to_numpy(),rows[rows.environment_region.eq('Mediterranean')]),
            ('France',registry.dominant_source_country.eq('France').to_numpy(),country[country.country.eq('France')])]:
            w=registry.harvested_total_ha.to_numpy()*mask
            for j,metric in enumerate(metrics):
                denominator_sum=float(w@denominator[:,j])
                value=float(w@numerator[:,j]/denominator_sum) if denominator_sum else np.nan
                q=reference[reference.model.eq(model)&reference.scenario.eq('ssp585')&reference.period.eq(period)&reference.metric.eq(metric)].iloc[0]
                np.testing.assert_allclose(value,q.change,rtol=0,atol=1e-11,equal_nan=True)
                observed.append(dict(domain=name,model=model,metric=metric,change=value))
    return dict(status='passed',comparisons=len(observed),source='Annual simulation outputs' if direct else 'Bundled paired grid values and valid-year counts',period=period)


def exposure_and_timing_arithmetic():
    """Reconstruct headline exposure directly, without the aggregation routine."""
    from analysis.paper_study.food_security_exposure_20261009.exposure import read
    from .timing import read as read_timing
    directory=ROOT/'analysis/paper_study/food_security_exposure_20261009'
    summary=read()['domain_summary']
    production=pd.read_csv(directory/'production_grid_input.csv').set_index('cell_id')
    pairs=pd.read_parquet(GRID/'full_grid_paired_changes_by_gcm.parquet',filters=[('metric','==',CANOPY)])
    comparisons=0
    for (scenario,period),part in pairs.groupby(['scenario','period']):
        changes=part.pivot(index='cell_id',columns='model',values='change').loc[production.index]
        counts=part.pivot(index='cell_id',columns='model',values='valid_year_pairs').loc[production.index]
        p=production.production_total_tonnes
        for mode,threshold in [('primary_any_paired_year',1),('robust_all3_ge27',27)]:
            complete=changes.notna().all(axis=1)&counts.ge(threshold).all(axis=1)
            unanimous=complete&changes.gt(0).all(axis=1)
            ensemble=complete&changes.mean(axis=1).gt(0)
            expected=summary[summary.domain_type.eq('Europe')&summary.scenario.eq(scenario)&summary.period.eq(period)&summary.analysis_mode.eq(mode)].iloc[0]
            for actual,target in [(p.sum(),expected.baseline_production_tonnes),
                (p[unanimous].sum(),expected.all3positive_production_tonnes),
                (p[ensemble].sum(),expected.ensemble_increasing_production_tonnes),
                (p[~complete].sum(),expected.unavailable_production_tonnes)]:
                np.testing.assert_allclose(actual,target,rtol=1e-12,atol=1e-7);comparisons+=1
            estimates=[]
            for model in changes.columns:
                valid=changes[model].notna()&counts[model].ge(threshold)
                if threshold>1:valid &= complete
                w=p[valid]*counts.loc[valid,model]
                estimates.append(float(w@changes.loc[valid,model]/w.sum()))
            np.testing.assert_allclose([np.mean(estimates),min(estimates),max(estimates)],
                [expected.production_weighted_change,expected.gcm_min,expected.gcm_max],rtol=0,atol=1e-11)
            comparisons+=3
    timing=read_timing().set_index(['environment_region','event'])
    for col in ['reference','future','change']:
        x=timing[col].unstack('event')
        np.testing.assert_allclose(x.symptom_days-x.flowering_days,x.symptom_relative_flowering_days,rtol=0,atol=1e-12)
        np.testing.assert_allclose(x.soft_dough_days-x.flowering_days,x.grain_fill_elapsed_days,rtol=0,atol=1e-12)
    return dict(status='passed',independent_exposure_comparisons=comparisons,
        timing_population='Identical complete pairs for all event and interval means',
        production_interpretation='Fixed baseline production located in exposed cells; not estimated grain loss')


def annual_risk_arithmetic():
    from analysis.paper_study.climate_robustness_20261009.run import inputs,read,MODELS,SCENARIOS,WEIGHTS
    registry,raw=inputs()
    published=read('annual_distribution_by_gcm')
    count=0
    for r in published.itertuples():
        mi,si=MODELS.index(r.model),SCENARIOS.index(r.scenario)
        start=30 if r.period=='2031-2060' else 60
        common=np.isfinite(raw[:,si,:30]).all(axis=(0,1)) & np.isfinite(raw[:,si,start:start+30]).all(axis=(0,1))
        domain=np.ones(len(registry),bool) if r.environment_region=='Europe' else registry.environment_region.fillna('Unassigned').eq(r.environment_region).to_numpy()
        use=common & domain
        w=registry[WEIGHTS[r.weighting]].to_numpy()[use]
        historical=raw[mi,si,:30][:,use]@w/w.sum()
        future=raw[mi,si,start:start+30][:,use]@w/w.sum()
        threshold=np.quantile(historical,.9)
        direct=[historical.mean(),future.mean(),threshold,np.quantile(future,.9),
                np.mean(future>threshold),np.sort(future)[-3:].mean(),
                w.sum()/registry.loc[domain,WEIGHTS[r.weighting]].sum()]
        reported=[r.historical_mean,r.future_mean,r.historical_q90,r.future_q90,
                  r.future_exceedance_fraction,r.future_top3_mean,r.complete_support_fraction]
        np.testing.assert_allclose(direct,reported,atol=1e-10,rtol=0)
        count+=len(direct)
    sensitivity=read('pairing_sensitivity')
    support=read('support_ensemble').pivot(index=['scenario','period','weighting','environment_region'],columns='method',values='change')
    assert sensitivity.direction_consistent.all()
    assert not (support.complete_common_cells*support.positional_pairs<0).any()
    return dict(status='passed',independent_annual_statistics=count,
        complete_population='Same cells across both 30-year periods and all three climate models',
        scope='Conditional annual canopy damage; no grain-loss or adaptation validation')


def management_grain_arithmetic():
    from .management_evidence import NORDIC,MIXTURE,read,headline
    nordic=read(NORDIC,'paired_management_grain.csv')
    french=read(MIXTURE,'french_contrasts.csv')
    n,f,_=headline()
    assert len(nordic)==307 and nordic.trial_key.nunique()==263
    np.testing.assert_allclose(nordic.treated_yield_t_ha-nordic.control_yield_t_ha,nordic.gain_t_ha,atol=1e-9,rtol=0)
    np.testing.assert_allclose(nordic.groupby('trial_key').trial_weight.sum(),1,atol=1e-12,rtol=0)
    direct=nordic.groupby('trial_key').gain_t_ha.mean().mean()
    np.testing.assert_allclose(direct,n.mean_gain_t_ha,atol=1e-9,rtol=0)
    plots=read(MIXTURE,'french_plots.csv').set_index('plot_id')
    assert not plots.index.duplicated().any() and len(french)==195
    baseline=(plots.loc[french.control_1_id,'yield_t_ha'].to_numpy()+plots.loc[french.control_2_id,'yield_t_ha'].to_numpy())/2
    delta=plots.loc[french.plot_id,'yield_t_ha'].to_numpy()-baseline
    np.testing.assert_allclose(delta,french.delta_t_ha,atol=1e-9,rtol=0)
    np.testing.assert_allclose(delta.mean(),f.loc['raw_primary','mean_delta_t_ha'],atol=1e-9,rtol=0)
    assert int((delta<0).sum())==79 and french.environment_id.nunique()==1
    applicability=read(MIXTURE,'applicability.json')
    for source in ['france','swiss']:
        assert not applicability[source]['validated_climate_adaptation']
        assert not applicability[source]['stb_mediated_grain_loss']
    return dict(status='passed',nordic_contrasts=len(nordic),nordic_trial_identifiers=263,
        french_constituent_comparisons=len(french),french_environments=1,
        scope='Observed management-associated grain outcomes; source populations are not pooled')


def verify(output):
    output=Path(output).resolve();main=(output/'Manuscript.txt').read_text();si=(output/'Supplementary_Information.txt').read_text()
    abstract=(output/'abstract.txt').read_text().strip()
    assert len(abstract.split())<=150 and not re.search(r'\[[1-9]\d*',abstract)
    assert main.count(abstract)==1
    assert 'private during manuscript preparation' not in main
    assert 'https://github.com/SmartAG-Team/paper_cc_impact_wheat_disease' in main
    # Archive names occur in source references; scientific Results do not divide
    # their observations into corporate datasets.
    result=main.split('\n\nResults\n\n',1)[1].split('\n\nDiscussion\n\n',1)[0]
    assert not re.search(r'\b(?:nan|inf)\b',result,flags=re.I)
    assert not re.search(r'BASF|Corteva|sampling weights|draw identities|area-proportional draws',result)
    assert not result.startswith('European wheat area')
    captions=json.loads((output/'figures/captions.json').read_text())
    assert len(captions)==5 and len(list((output/'figures').glob('*.png')))==5
    for stem,caption in captions.items():
        assert len(caption.split())<=220 and all((output/'figures'/f'{stem}.{ext}').exists() for ext in ['png','pdf','svg'])
    assert 'harvested area and field' not in list(captions.values())[0].lower()
    order=[int(n) for n in re.findall(r'(?<!Supplementary )Fig\. (\d+)',main.split('Figure legends')[0])]
    assert list(dict.fromkeys(order))==[1,2,3,4,5]
    declared=set(re.findall(r'^Table (S\d+[a-z]?) \|',si,flags=re.M))
    for identifier in re.findall(r'(?:Table|Tables) (S\d+[a-z]?)',main):
        assert identifier in declared or any(x.startswith(identifier) for x in declared),identifier
    doc=Document(output/'Manuscript.docx');assert len(doc.inline_shapes)==5 and len(doc.tables)==1
    word='\n'.join(p.text for p in doc.paragraphs);assert abstract in word
    i=next(i for i,p in enumerate(doc.paragraphs) if p.text==abstract)
    assert doc.paragraphs[i+1].paragraph_format.page_break_before
    pages={}
    for name in ['Manuscript.pdf','Supplementary_Information.pdf','Supplementary_Model_Specification.pdf']:
        pdf=pymupdf.open(output/name);pages[name]=len(pdf)
        assert all(len(p.get_text().strip())>5 for p in pdf)
        if name=='Manuscript.pdf':assert abstract in ' '.join(pdf[0].get_text().split())
    workbook=load_workbook(output/'Source_Data.xlsx',read_only=True)
    for sheet,name,count,expected in list(workbook['Source_manifest'].values)[1:]:
        path=output/Path(name).relative_to('publication/european_wheat_stb') if name.startswith('publication/european_wheat_stb/') else ROOT/name
        assert sha(path)==expected,name
    archive_count=None
    if (output/'Software_and_Evidence.zip').exists():
        with ZipFile(output/'Software_and_Evidence.zip') as z:
            assert z.testzip() is None
            manifest=json.loads(z.read('PACKAGE_MANIFEST.json'))
            for name,record in manifest['members'].items():
                content=z.read(name);assert hashlib.sha256(content).hexdigest()==record['sha256'] and len(content)==record['bytes'],name
            archive_count=len(manifest['members'])
    arithmetic=pooled_arithmetic();domains=domain_arithmetic();exposure=exposure_and_timing_arithmetic()
    annual=annual_risk_arithmetic()
    management=management_grain_arithmetic()
    report=dict(status='passed',scope='Numerical reporting, provenance and outcome-specific evaluation; end-to-end grain-loss predictive validity is not established',
        abstract_words=len(abstract.split()),main_figures=5,main_tables=1,Results_focus='Predictive support, canopy damage, production exposure, annual risk, crop–disease timing and observed management responses',
        pooled_evaluation=arithmetic,full_grid_domain_check=domains,PDF_pages=pages,
        production_exposure_and_timing=exposure,
        annual_canopy_risk=annual,
        observed_management_grain=management,
        source_workbook_sheets=len(workbook.sheetnames),archived_members_checked=archive_count)
    (output/'verification_receipt.json').write_text(json.dumps(report,indent=2)+'\n')
    return report


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,required=True)
    print(json.dumps(verify(parser.parse_args().output),indent=2))
