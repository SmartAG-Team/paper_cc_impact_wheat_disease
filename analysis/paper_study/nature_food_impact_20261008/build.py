"""Build the climate-impact manuscript and retain reproducible numerical evidence."""
from pathlib import Path
import argparse
import hashlib
import json
import re
import shutil
import zipfile
import numpy as np
import pandas as pd
from openpyxl import Workbook,load_workbook
from . import data,figures,tables,sections
from analysis.paper_study.nature_food_submission_20261008 import build as base

BASE_CURATE=base._curated_figures
BASE_PACKAGE=base._package
MAIN_ORDER=['abstract.txt','introduction.txt','climate_results.txt','field_results.txt','baseline_results.txt',
            'discussion_field_and_limits.txt','methods.txt','availability.txt']


class FigureCaptions(dict):
    def items(self):
        return iter(sorted(super().items(),key=lambda kv:int(re.match(r'figS(\d+)',kv[0]).group(1))))

    def values(self):
        return (value for _,value in self.items())


def prepare():
    data.prepare_pooled()
    from analysis.paper_study.nature_food_submission_20261008.prepare import prepare as prepare_previous
    prepare_previous()


def curate(captions,folder):
    for old,new in [('figS13_spatial_symptom_changes','figS12_spatial_symptom_changes'),
                    ('figS14_spatial_yield_transfers','figS13_spatial_yield_transfers')]:
        source=folder/f'{old}.csv.gz'
        if source.exists():source.rename(folder/f'{new}.csv.gz')
    captions=BASE_CURATE(captions,folder)
    previous='figS14_full_grid_example_2001';current='figS14_historical_disease_patterns'
    if previous in captions:
        captions[current]=captions.pop(previous)
        for suffix in ['png','pdf','svg','csv','csv.gz']:
            source=folder/f'{previous}.{suffix}'
            if source.exists():source.rename(folder/f'{current}.{suffix}')
    captions['figS15_climate_robustness']=figures.sensitivity_figure(folder)
    figures.contribution_figure(folder,'figS16_weather_host_decomposition',full=True)
    captions['figS16_weather_host_decomposition']=(
        'Figure S16 | Weather and wheat-development contributions and offsets. '
        '(a) Weather, crop-development and net HAD contributions with the full interaction shown separately. '
        'The interaction is already allocated equally between the two Shapley contributions and is not added again to the net. '
        'Gray segments span climate-model means; colored whiskers show one spatial Monte Carlo standard error. '
        '(b) Model-specific offsets and their ensemble-contribution ratio. The same 5,731 supported paired seasons underlie all HAD terms. '
        'The sampled analysis is separate from the full-grid estimates and does not measure intervention benefits.')
    captions['figS20_pooled_field_evaluation']=figures.pooled_validation_figure(folder)
    captions['figS21_study_domain']=figures.study_context(folder)
    # Full-grid source tables are archived separately; obsolete sampled source
    # CSVs from the predecessor figure selection do not enter this publication.
    return FigureCaptions(captions)


def source_workbook(destination,canonical='publication/european_wheat_stb'):
    sources=[*sorted(data.DERIVED.glob('*.csv'))]
    sources += [data.GRID/name for name in ['environmental_region_changes.csv','environmental_region_changes_by_gcm.csv',
        'country_changes.csv','country_changes_by_gcm.csv','canopy_change_sign_area_shares.csv',
        'annual_area_weighted_means.csv','annual_status_and_area_coverage.csv','full_landuse_cell_registry.csv']]
    original=data.ROOT/'analysis/paper_study/yield_evidence_20261007'
    sources += [original/name for name in ['parker2004_table4_yield_loss_matrix.csv','parker2004_table5_yield_slopes.csv',
        'foulkes2006_group_separated_predictions.csv','published_yield_response_metrics.csv','published_yield_domain_compatibility.csv']]
    sources += [data.ROOT/'analysis/paper_study/nature_food_revision_20261007/yield_response/model_comparison_metrics.csv',
                data.ROOT/'analysis/paper_study/nature_food_revision_20261007/basf_crop_response/year_transfer_metrics.csv']
    sources += [data.ROOT/'analysis/paper_study/nature_food_fix_20261007/climate/headline_robustness.csv',
                data.ROOT/'analysis/paper_study/nature_food_fix_20261007/climate/decomposition/supported_forcing_ensemble_decomposition.csv',
                data.ROOT/'analysis/paper_study/nature_food_fix_20261007/climate/decomposition/supported_forcing_decomposition_by_gcm.csv',
                destination/'Main_Table1.csv']
    sources += sorted((destination/'figures').glob('*.csv'))+sorted((destination/'figures').glob('*.csv.gz'))
    sources += sorted((destination/'supplementary_figures').glob('*.csv'))+sorted((destination/'supplementary_figures').glob('*.csv.gz'))
    workbook=Workbook(write_only=True);manifest=[]
    for i,path in enumerate(sources,1):
        frame=pd.read_csv(path)
        name=f'{i:02d}_'+re.sub(r'[^A-Za-z0-9]+','_',path.name.removesuffix('.csv.gz').removesuffix('.csv'))
        name=name[:31];sheet=workbook.create_sheet(name);sheet.append(list(frame.columns))
        for row in frame.itertuples(index=False,name=None):
            sheet.append([None if pd.isna(v) else v.item() if isinstance(v,np.generic) else v for v in row])
        source_path=(Path(canonical)/path.relative_to(destination)).as_posix() if path.is_relative_to(destination) else str(path.relative_to(data.ROOT))
        manifest.append([name,source_path,len(frame),data.sha(path)])
    sheet=workbook.create_sheet('Source_manifest');sheet.append(['Sheet','Source','Rows','SHA256'])
    for row in manifest:sheet.append(row)
    sheet=workbook.create_sheet('Definitions');sheet.append(['Quantity','Definition'])
    for row in [
      ['Full-grid population','14,941 fixed all-wheat cells; 14,932 imposed winter-wheat rainfed crop calendars.'],
      ['Full-grid paired change','Fixed harvested-area-weighted change over valid paired seasons, followed by three-model averaging.'],
      ['Normalized HAD loss','HAD loss divided by maximum reference upper-three-leaf LAI; units are days. Numerical values are unchanged when reference LAI equals one.'],
      ['Country group','Dominant SPAM source-country label of a quarter-degree cell; study-domain subset only.'],
      ['Symptom frequency','Fraction of complete simulated grid-seasons with flag-leaf symptoms before soft dough.'],
      ['Relative HAD loss','HAD loss divided by reference HAD; percentage of assumed canopy function lost during grain filling, distinct from lesion percentage.'],
      ['Estimated disease-related yield change','Minus the published coefficient times the change in HAD loss; kg ha-1 per unit maximum reference upper-canopy LAI.'],
      ['Climate-model range','Deterministic minimum and maximum across three climate-model domain estimates; not a confidence interval.'],
      ['Sampled diagnostics','Weather-crop decomposition and structural sensitivity retain 64 draws at 62 cells.'],
      ['Pooled field scores','Equal coordinate-year, field, source-leaf and assessment weight; calibration excluded.'],
      ['Pooled intervals','20,000 paired coordinate-year bootstrap resamples; source conventions retained.']]:sheet.append(row)
    workbook.save(destination/'Source_Data.xlsx')


def package(destination,evidence_archive=None):
    shutil.copy2(data.HERE/'Reproducibility.txt',destination/'REPRODUCIBILITY.txt')
    BASE_PACKAGE(destination,evidence_archive)
    target=destination/'Software_and_Evidence.zip';temporary=destination/'Software_and_Evidence.extended.tmp'
    additional={}
    for directory in [data.DERIVED,data.GRID]:
        for path in directory.iterdir():
            if path.is_file():additional[path.relative_to(data.ROOT).as_posix()]=path
    specification=data.ROOT/'analysis/paper_study/nature_food_fix_20261007/specification'
    for path in specification.rglob('*'):
        if path.is_file() and 'node_modules' not in path.parts and '__pycache__' not in path.parts:
            additional[path.relative_to(data.ROOT).as_posix()]=path
    additional[(data.HERE/'Reproducibility.txt').relative_to(data.ROOT).as_posix()]=data.HERE/'Reproducibility.txt'
    for relative in ['analysis/paper_study/full_grid_climate_20261008/run_configuration.json',
                     'analysis/paper_study/nature_food_revision_20261007/continental_replay/configuration_before_results.json']:
        path=data.ROOT/relative
        if path.exists():additional[relative]=path
    members={}
    with zipfile.ZipFile(target) as source,zipfile.ZipFile(temporary,'w',zipfile.ZIP_DEFLATED,compresslevel=4) as output:
        def write(name,content):
            output.writestr(name,content)
            members[name]={'sha256':hashlib.sha256(content).hexdigest(),'bytes':len(content)}
        for record in source.infolist():
            if record.is_dir() or record.filename=='PACKAGE_MANIFEST.json' or record.filename in additional:continue
            if record.filename=='REPRODUCIBILITY.txt':write(record.filename,(data.HERE/'Reproducibility.txt').read_bytes())
            elif record.filename=='publication/european_wheat_stb/REPRODUCIBILITY.txt':
                write(record.filename,(destination/'REPRODUCIBILITY.txt').read_bytes())
            else:write(record.filename,source.read(record))
        for name,path in additional.items():write(name,path.read_bytes())
        manifest={'package':'European wheat full-grid climate-impact manuscript',
            'entry_point':'analysis/paper_study/nature_food_impact_20261008/build.py',
            'publication_command':'python -m analysis.paper_study.nature_food_impact_20261008.build --output publication/european_wheat_stb_regenerated --documents-only',
            'member_count_excluding_manifest':len(members),'checksums_exclude_this_manifest':True,'members':members}
        output.writestr('PACKAGE_MANIFEST.json',json.dumps(manifest,indent=2)+'\n')
    temporary.replace(target)


def build(destination,documents_only=False,evidence_archive=None):
    base.sections=sections.sections;base.TITLE=sections.TITLE
    base.draw_figures=figures.main;base.main_tables=tables.main_tables
    base.validation_figures=figures.validation_supplementary;base.yield_figures=figures.yield_supplementary
    base.map_figures=figures.map_supplementary;base.supplementary_phenology=figures.supplementary_phenology
    base.environmental_comparison=figures.agreement_figure;base._curated_figures=curate
    base._curated_tables=tables.supplementary_tables;base._source_workbook=source_workbook
    base.prepare=prepare;base._package=package
    base.supplementary_framework=figures.supplementary_framework
    # The legacy table collector is retained as an immutable callable; replacing
    # it globally would create recursion in the selected published-data tables.
    tables.legacy_tables=LEGACY_TABLES
    receipt=base.build(Path(destination),documents_only,evidence_archive)
    path=Path(destination)/'REPRODUCIBILITY.txt'
    shutil.copy2(data.HERE/'Reproducibility.txt',path)
    return receipt


LEGACY_TABLES=base._curated_tables

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--documents-only',action='store_true');parser.add_argument('--evidence-archive',type=Path)
    args=parser.parse_args()
    print(json.dumps(build(args.output,args.documents_only,args.evidence_archive),indent=2))
