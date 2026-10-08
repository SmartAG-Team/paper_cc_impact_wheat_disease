"""Formal manuscript, supplement, PDF and source-data workbook."""
from pathlib import Path
import html,json,re,sys
import pandas as pd
import numpy as np
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak,Image,Table,TableStyle,KeepTogether
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from matplotlib import font_manager
from html.parser import HTMLParser

ROOT=Path(__file__).resolve().parents[3];sys.path.insert(0,str(ROOT))
from analysis.paper_study.run_regional import sha
from analysis.paper_study.manuscript.author_date_citations import plain
DEST=ROOT/'publication/european_wheat_stb';DATA=ROOT/'data/paper_study/publication'
TITLE='Crop development modulates simulated Septoria pressure on European wheat under climate change'
CITATION_PATTERN=r'\[([1-9]\d*(?:(?:,\s*|[–-])[1-9]\d*)*)\]'


def paragraphs(name):return (DEST/name).read_text().strip().split('\n\n')


def citation_numbers(text):
    for match in re.finditer(CITATION_PATTERN,text):
        for token in match[1].split(','):
            ends=re.split('[–-]',token.strip())
            yield from range(int(ends[0]),int(ends[-1])+1)


def article():
    discussion=paragraphs('discussion_field_and_limits.txt')
    body=paragraphs('abstract.txt')+paragraphs('introduction.txt')+paragraphs('field_results.txt')
    body+=paragraphs('baseline_results.txt')+paragraphs('climate_results.txt')
    body+=discussion[:2]+paragraphs('climate_discussion.txt')+discussion[2:]
    body+=paragraphs('methods.txt')+paragraphs('availability.txt')
    refs={int(p.split('.',1)[0]):p.split('.',1)[1].strip() for p in paragraphs('references.txt') if re.match(r'^\d+\.',p)}
    order=[]
    for n in citation_numbers('\n\n'.join(body)):
        assert n in refs
        if n not in order:order.append(n)
    assert set(order)==set(refs)
    rendered=json.loads((DEST/'citation_rendered.json').read_text())
    assert rendered['style_id']=='http://www.zotero.org/styles/apa'
    mapping={old:plain(rendered['citations'][f'[{old}]']) if f'[{old}]' in rendered['citations'] else f'ref-{old:02d}' for old in order}
    def renumber(text):
        def replace(match):
            return plain(rendered['citations'][match[0]])
        return re.sub(CITATION_PATTERN,replace,text)
    body=[renumber(p) for p in body]
    references=[plain(entry) for entry in rendered['bibliography_html']]
    assert len(references)==len(refs)==18
    return body,references,mapping


def configure_doc(document):
    section=document.sections[0]
    section.page_width=Inches(8.27);section.page_height=Inches(11.69)
    section.top_margin=section.bottom_margin=Inches(.75)
    section.left_margin=section.right_margin=Inches(.78)
    style=document.styles['Normal'];style.font.name='Times New Roman';style.font.size=Pt(11)
    style.paragraph_format.line_spacing=1.15;style.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',16),('Heading 1',13),('Heading 2',11)]:
        document.styles[name].font.name='Times New Roman';document.styles[name].font.size=Pt(size)
        document.styles[name].font.color.rgb=None
    footer=section.footer.paragraphs[0];footer.alignment=2
    run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.append(field)
    document.core_properties.title=TITLE;document.core_properties.author='Gang Zhao'


def is_heading(text):return len(text)<90 and not text.endswith(('.',']')) and '\n' not in text


def add_prose_doc(document,paragraphs):
    for text in paragraphs:
        if is_heading(text):
            document.add_heading(text,level=1 if text in ['Introduction','Results','Discussion','Methods','Data availability','Code availability','Supplementary Information','Supplementary Methods'] else 2)
        else:document.add_paragraph(text)


class ReferenceRuns(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True);self.runs=[];self.stack=[]
    def handle_starttag(self,tag,attrs):self.stack.append(tag)
    def handle_endtag(self,tag):
        if tag in self.stack:
            i=len(self.stack)-1-self.stack[::-1].index(tag);del self.stack[i:]
    def handle_data(self,data):
        if data.strip():self.runs.append((data,'i' in self.stack or 'em' in self.stack))


def add_reference_doc(document,value):
    paragraph=document.add_paragraph();paragraph.paragraph_format.left_indent=Inches(.3)
    paragraph.paragraph_format.first_line_indent=Inches(-.3)
    parser=ReferenceRuns();parser.feed(value)
    for text,italic in parser.runs:
        run=paragraph.add_run(text);run.italic=italic


def reference_inline(value):
    value=re.sub(r'</?div[^>]*>','',value).strip()
    value=re.sub(r'</?span[^>]*>','',value)
    return value


def table_contents():
    one=pd.read_csv(DATA/'table1_field_severity.csv')
    first=[['Evaluation','Leaf scope','Predictor','n','RMSE','MAE','Bias']]
    for r in one.itertuples():first.append([r.evaluation,r.leaf_scope,r.model,str(r.n),f'{r.rmse:.2f}',f'{r.mae:.2f}',f'{r.bias:.2f}'])
    two=pd.read_csv(DATA/'table2_climate_contrasts.csv');second=[['SSP','Period','Area-weighted change','Production-weighted change','Disease weather','Host development']]
    for r in two.itertuples():
        second.append([{'ssp126':'1-2.6','ssp245':'2-4.5','ssp585':'5-8.5'}[r.scenario],
            '2031–2060' if 'mid' in r.period else '2071–2100',
            f'{r.area_change:.2f} [{r.area_low:.2f}, {r.area_high:.2f}]',
            f'{r.production_change:.2f} [{r.production_low:.2f}, {r.production_high:.2f}]',
            f'{r.weather_contribution:.2f}',f'{r.host_contribution:.2f}'])
    captions=[
        'Table 1 | Final recorded severity under frozen prediction models. Errors and bias are percentage points. Final numeric endpoints are selected separately for each field–source-leaf series. Equal coordinate-year, field and selected-leaf weights are reconstructed after restricting the leaf scope. n counts source leaf endpoints rather than independent fields. The external subset excludes locations shared with the original development archive. The empirical comparator families were introduced after initial external scoring; their fitting and tuning use calibration data only.',
        'Table 2 | Conditional European climate contrasts. Values are percentage-point changes from each climate model–SSP-specific 1991–2020 reference, averaged over three models. Brackets are 2.5th–97.5th percentiles of 100 joint calibration-parameter draws, with structure and latent delay fixed. Area and production changes have exact full-grid point anchors. Weather and host columns are spatial-sample point contributions on common paired seasons; their sum is the sampled net and can differ from the full-grid change. Sampling errors and individual climate-model ranges accompany the source data.']
    return [(first,captions[0]),(second,captions[1])]


def add_native_table(document,rows,caption):
    document.add_paragraph(caption)
    table=document.add_table(rows=1,cols=len(rows[0]));table.style='Table Grid'
    for j,value in enumerate(rows[0]):table.rows[0].cells[j].text=value
    for row in rows[1:]:
        cells=table.add_row().cells
        for j,value in enumerate(row):cells[j].text=str(value)
    for i,row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.
                for run in p.runs:run.font.size=Pt(9);run.bold=i==0


def font_styles():
    for name,weight,style in [('PaperSerif','normal','normal'),('PaperSerifBold','bold','normal'),
            ('PaperSerifItalic','normal','italic'),('PaperSerifBoldItalic','bold','italic')]:
        path=font_manager.findfont(font_manager.FontProperties(family='DejaVu Serif',weight=weight,style=style))
        pdfmetrics.registerFont(TTFont(name,path))
    pdfmetrics.registerFontFamily('PaperSerif',normal='PaperSerif',bold='PaperSerifBold',
        italic='PaperSerifItalic',boldItalic='PaperSerifBoldItalic')
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle('PaperBody',fontName='PaperSerif',fontSize=10,leading=13.5,spaceAfter=7,splitLongWords=True))
    styles.add(ParagraphStyle('PaperTitle',fontName='PaperSerifBold',fontSize=15,leading=19,spaceAfter=16))
    styles.add(ParagraphStyle('PaperHeading',fontName='PaperSerifBold',fontSize=12,leading=15,spaceBefore=12,spaceAfter=7,keepWithNext=True))
    styles.add(ParagraphStyle('PaperCaption',fontName='PaperSerif',fontSize=8.5,leading=11,spaceAfter=7))
    styles.add(ParagraphStyle('PaperTable',fontName='PaperSerif',fontSize=7.2,leading=9))
    styles.add(ParagraphStyle('PaperReference',parent=styles['PaperBody'],leftIndent=16,firstLineIndent=-16))
    return styles


def pdf_prose(flow,texts,styles):
    for text in texts:flow.append(Paragraph(html.escape(text),styles['PaperHeading' if is_heading(text) else 'PaperBody']))


def pdf_tables(flow,tables,styles):
    for rows,caption in tables:
        flow.append(PageBreak());flow.append(Paragraph(html.escape(caption),styles['PaperCaption']))
        data=[[Paragraph(html.escape(str(value)),styles['PaperTable']) for value in row] for row in rows]
        widths=([95,66,83,24,38,38,38] if len(rows[0])==7 else [32,66,105,105,58,72])
        table=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT')
        table.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eef1f3')),
            ('LINEBELOW',(0,0),(-1,0),.5,colors.grey),('LINEBELOW',(0,-1),(-1,-1),.5,colors.grey),
            ('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
        flow.append(table)


def add_figures(document,flow,captions,styles):
    from PIL import Image as PILImage
    for name,caption in captions.items():
        document.add_page_break();document.add_picture(str(DEST/'figures'/f'{name}.png'),width=Inches(6.45));document.add_paragraph(caption)
        flow.append(PageBreak());path=DEST/'figures'/f'{name}.png'
        with PILImage.open(path) as im:w,h=im.size
        width=450;height=width*h/w
        if height>535:width*=535/height;height=535
        flow.append(Image(str(path),width=width,height=height));flow.append(Spacer(1,10));flow.append(Paragraph(html.escape(caption),styles['PaperCaption']))


def footer(canvas,doc):
    canvas.setFont('PaperSerif',8);canvas.drawRightString(A4[0]-45,30,str(doc.page))


def workbook():
    book=Workbook(write_only=True);sources=[]
    csvs=sorted(DATA.glob('*.csv'))
    extras=[ROOT/'analysis/paper_study/seasonal_statistics/calibration_parameter_intervals.csv',
        ROOT/'analysis/paper_study/seasonal_statistics/secondary_beta_training_profile.csv',
        ROOT/'analysis/paper_study/field_sensitivity/final_severity_sowing_sensitivity.csv',
        ROOT/'analysis/paper_study/field_sensitivity/independent_stage_transfer_intervals.csv',
        ROOT/'analysis/paper_study/external_statistics/onset_censoring_statistics.csv',
        ROOT/'data/paper_study/calendar_sensitivity_reporting/mirca_mixed_calendar_periods.csv',
        ROOT/'data/paper_study/calendar_sensitivity_reporting/sampled_sensitivity_periods.csv',
        ROOT/'data/paper_study/production_exposure_reporting/europe_period_production_exposure.csv',
        ROOT/'data/paper_study/production_exposure_reporting/conditional_production_projection_bands.csv']
    for i,path in enumerate(csvs+extras):
        name=(f'{i+1:02d}_'+path.stem)[:31];sheet=book.create_sheet(name);sheet.freeze_panes='A2'
        frame=pd.read_csv(path);header=[]
        for name in frame.columns:
            cell=WriteOnlyCell(sheet,value=name);cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='34495E');header.append(cell)
        sheet.append(header)
        for row in frame.itertuples(index=False,name=None):
            sheet.append([None if pd.isna(value) else value.item() if isinstance(value,np.generic) else value for value in row])
        sources.append(dict(sheet=sheet.title,source_path=str(path.relative_to(ROOT)),rows=len(frame),columns=len(frame.columns),sha256=sha(path)))
    sheet=book.create_sheet('Source_manifest')
    sheet.append(['Sheet','Source path','Rows','Columns','SHA256'])
    for row in sources:sheet.append(list(row.values()))
    sheet=book.create_sheet('Definitions')
    for row in [
        ['Definition','Meaning'],['Severity','Normalized untreated disease percentage proxy; source denominator incompletely documented'],
        ['Final field endpoint','Last numeric assessment separately for each source leaf series'],
        ['Final regional endpoint','Upper-three mean at predicted soft dough; not observed harvest'],
        ['Reference production','SPAM2020 metric tonnes associated with source agricultural statistics; not future yield'],
        ['Exposed production','Reference tonnes in simulated threshold-exceeding cell-years; not tonnes lost'],
        ['Parameter intervals','Conditional2.5–97.5% bootstrap percentiles; fixed structural assumptions; no observation residual uncertainty'],
        ['Spatial MCSE','Sampling uncertainty from64 registered area-proportional draws'],
        ['Climate range','Three point-parameter climate models; not a probabilistic confidence interval'],
        ['True primary infection date','Unobserved; internal external-exposure diagnostic is not validated as that date']]:sheet.append(row)
    book.save(DEST/'Source_Data.xlsx');return sources


def main():
    body,references,mapping=article();tables=table_contents();styles=font_styles()
    text=TITLE+'\n\nAbstract\n\n'+'\n\n'.join(body)+'\n\nReferences\n\n'+'\n\n'.join(references)+'\n'
    (DEST/'Manuscript.txt').write_text(text)
    document=Document();configure_doc(document);document.add_heading(TITLE,0);document.add_heading('Abstract',1)
    add_prose_doc(document,body);document.add_heading('References',1)
    reference_html=json.loads((DEST/'citation_rendered.json').read_text())['bibliography_html']
    for reference in reference_html:add_reference_doc(document,reference)
    flow=[Paragraph(html.escape(TITLE),styles['PaperTitle']),Paragraph('Abstract',styles['PaperHeading'])]
    pdf_prose(flow,body,styles);flow.append(Paragraph('References',styles['PaperHeading']))
    for reference in reference_html:flow.append(Paragraph(reference_inline(reference),styles['PaperReference']))
    for rows,caption in tables:document.add_page_break();add_native_table(document,rows,caption)
    pdf_tables(flow,tables,styles)
    captions={}
    for name in ['field_and_baseline_captions.json','climate_and_production_captions.json']:
        captions.update(json.loads((DEST/'figures'/name).read_text()))
    for rows,caption in tables:
        text+='\n'+caption+'\n\n'+'\n'.join('\t'.join(str(value) for value in row) for row in rows)+'\n'
    text+='\nFigure legends\n'+''.join('\n'+caption+'\n' for caption in captions.values())
    (DEST/'Manuscript.txt').write_text(text)
    add_figures(document,flow,captions,styles)
    document.save(DEST/'Manuscript.docx')
    SimpleDocTemplate(str(DEST/'Manuscript.pdf'),pagesize=A4,rightMargin=45,leftMargin=45,topMargin=45,bottomMargin=45,
        title=TITLE,author='Gang Zhao').build(flow,onFirstPage=footer,onLaterPages=footer)
    supplement=Document();configure_doc(supplement);add_prose_doc(supplement,paragraphs('supplementary_methods.txt'))
    sf=[];pdf_prose(sf,paragraphs('supplementary_methods.txt'),styles)
    sc=json.loads((DEST/'figures/supplementary_captions.json').read_text());add_figures(supplement,sf,sc,styles)
    supplement.save(DEST/'Supplementary_Information.docx')
    SimpleDocTemplate(str(DEST/'Supplementary_Information.pdf'),pagesize=A4,rightMargin=45,leftMargin=45,topMargin=45,bottomMargin=45).build(sf,onFirstPage=footer,onLaterPages=footer)
    sources=workbook()
    receipt=dict(status='generated',title=TITLE,main_figures=len(captions),main_tables=len(tables),supplementary_figures=len(sc),
        abstract_words=len(body[0].split()),methods_words=len((DEST/'methods.txt').read_text().split()),reference_count=len(references),
        citation_style='APA 7th edition (author–date)',citation_style_id='http://www.zotero.org/styles/apa',
        reference_id_mapping=mapping,references_generated_from_csl=True,zotero_word_plugin_fields=False,
        source_workbook=sources,source_code_sha256=sha(Path(__file__)),
        pdf_is_direct_layout_of_same_manuscript_text=True,docx_is_native_word_text_and_tables=True,
        authorship_and_funding_not_invented=True)
    (DEST/'document_build_receipt.json').write_text(json.dumps(receipt,indent=2)+'\n')
    print(json.dumps({k:v for k,v in receipt.items() if k!='source_workbook'}),flush=True)


if __name__=='__main__':main()
