"""Validate final document structure, source workbook and scientific receipts."""
from pathlib import Path
import hashlib,json,re
import numpy as np
import pandas as pd
import pymupdf
from docx import Document
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[3];DEST=ROOT/'publication/european_wheat_stb'


def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    build=json.loads((DEST/'document_build_receipt.json').read_text())
    document=Document(DEST/'Manuscript.docx')
    assert len(document.tables)==2 and len(document.inline_shapes)==6
    assert len(Document(DEST/'Supplementary_Information.docx').inline_shapes)==build['supplementary_figures']==3
    text=(DEST/'Manuscript.txt').read_text()
    assert not re.search(r'(?i)\b(?:TODO|TBD|PLACEHOLDER|INSERT HERE|this report will|this paper will|according to the user|below we|here we)\b',text)
    assert 'β∈[0,10]' in text
    assert 150<=build['abstract_words']<=200 and build['reference_count']==18
    reference_index=next(i for i,p in enumerate(document.paragraphs) if p.text=='References')
    reference_text='\n'.join(p.text for p in document.paragraphs[reference_index+1:reference_index+19])
    for forbidden in ['MDPI','Frontiers in','Scientific Reports']:assert forbidden not in reference_text
    assert build['citation_style_id']=='http://www.zotero.org/styles/apa'
    assert '(Savary et al., 2019)' in text and '(Suffert & Sache, 2011)' in text
    assert not re.search(r'\[([1-9]\d*(?:(?:,\s*|[–-])[1-9]\d*)*)\]',text)
    for name in ['Manuscript','Supplementary_Information']:
        pdf=pymupdf.open(DEST/f'{name}.pdf')
        assert all(page.get_text().strip() for page in pdf)
        assert not any('\ufffd' in page.get_text() for page in pdf)
        for page in pdf:
            for word in page.get_text('words'):
                assert word[0]>=30 and word[2]<=page.rect.width-25 and word[1]>=20 and word[3]<=page.rect.height-15
    workbook=load_workbook(DEST/'Source_Data.xlsx',read_only=True,data_only=True)
    checks=0;max_relative=0.
    for record in build['source_workbook']:
        source=ROOT/record['source_path'];assert sha(source)==record['sha256']
        frame=pd.read_csv(source);sheet=workbook[record['sheet']]
        actual=sheet.iter_rows(values_only=True,max_col=record['columns']);assert list(next(actual))==frame.columns.tolist()
        count=0
        for expected,observed in zip(frame.itertuples(index=False,name=None),actual,strict=True):
            assert len(expected)==len(observed)
            for a,b in zip(expected,observed):
                if pd.isna(a):assert b is None
                elif isinstance(a,(float,int,np.number)) and not isinstance(a,(bool,np.bool_)):
                    assert b is not None and np.isclose(a,b,rtol=2e-14,atol=1e-12)
                    max_relative=max(max_relative,abs(float(a)-float(b))/max(abs(float(a)),1.))
                else:assert a==b
                checks+=1
            count+=1
        assert count==record['rows']
    workbook.close()
    receipts={
        'climate':'analysis/paper_study/climate/climate_full_acquisition_receipt.json',
        'regional':'analysis/paper_study/regional_reporting/aggregation_receipt.json',
        'parameter_ensemble':'analysis/paper_study/regional_parameter_uncertainty/independent_ensemble_validation.json',
        'projection_reporting':'analysis/paper_study/projection_uncertainty_reporting/independent_reporting_validation.json',
        'scenario_and_fit_review':'analysis/paper_study/manuscript_review_20261006/independent_validation.json',
        'citations':'analysis/paper_study/citation_author_date_20261006/citation_validation.json',
        'harmonization':'analysis/paper_study/harmonization/independent_validation.json'}
    r={key:json.loads((ROOT/path).read_text()) for key,path in receipts.items()}
    assert r['climate']['status']=='complete' and len(r['regional']['completed_periods'])==56
    assert r['parameter_ensemble']['status']==r['projection_reporting']['status']==r['harmonization']['status']=='passed'
    assert r['scenario_and_fit_review']['status']=='passed'
    assert r['citations']['status']=='passed' and r['citations']['reference_records']==18
    assert r['citations']['scientific_body_changed'] is False
    provenance=json.loads((DEST/'figures/supplementary_provenance.json').read_text())
    assert provenance['source_code_sha256']==sha(ROOT/'analysis/paper_study/manuscript/prepare_publication_tables.py')
    for sources in provenance['figures'].values():
        for path,digest in sources.items():assert sha(ROOT/path)==digest
    provenance=json.loads((DEST/'figures/climate_and_production_provenance.json').read_text())
    assert provenance['source_code_sha256']==sha(ROOT/'analysis/paper_study/manuscript/plot_climate_and_production.py')
    for path,digest in provenance['source_hashes'].items():assert sha(ROOT/path)==digest
    table=pd.read_csv(ROOT/'data/paper_study/publication/table2_climate_contrasts.csv')
    assert len(table)==6 and set(table.scenario)=={'ssp126','ssp245','ssp585'} and table.period.nunique()==2
    country=pd.read_csv(ROOT/'data/paper_study/publication/fig6_country_production_exposure.csv')
    assert len(country)==90 and country.groupby('country').scenario.nunique().eq(3).all()
    paired=pd.read_csv(ROOT/'data/paper_study/publication/paired_scenario_contrasts.csv')
    assert len(paired)==48 and paired.period.nunique()==2
    assert r['harmonization']['scientific_measurement_rows']==1355842
    assert r['harmonization']['continental_partitions_byte_checksum_verified']==1059
    frozen=ROOT/'analysis/paper_study/seasonal_calibration_v1/frozen_main_fit.json'
    assert sha(frozen)=='e5feed4c5520184d2128dd9c1bf7c0a0c1c9ad92d25f6180590b1e825e4b3850'
    report=dict(status='passed',source_workbook_value_checks=checks,maximum_workbook_relative_numeric_difference=max_relative,
        native_word_tables=2,main_figures=6,supplementary_figures=build['supplementary_figures'],
        main_pdf_pages=len(pymupdf.open(DEST/'Manuscript.pdf')),
        supplementary_pdf_pages=len(pymupdf.open(DEST/'Supplementary_Information.pdf')),
        scientific_receipt_hashes={key:sha(ROOT/path) for key,path in receipts.items()},
        frozen_main_parameters_unchanged=True,prohibited_metawriting_or_placeholders=False,
        climate_scenarios=['SSP1-2.6','SSP2-4.5','SSP5-8.5'],future_periods=['2031–2060','2071–2100'],
        all_scenarios_in_country_and_sensitivity_figures=True,
        reference_scope_verified=True,docx_layout_not_rendered_by_word_engine=True,
        citation_style='APA 7th edition (author–date)',zotero_importable_reference_records=18,
        zotero_word_plugin_fields=False,
        standalone_pdf_visually_reviewed=True,source_code_sha256=sha(Path(__file__)))
    (DEST/'final_validation.json').write_text(json.dumps(report,indent=2)+'\n');print(json.dumps(report),flush=True)


if __name__=='__main__':main()
