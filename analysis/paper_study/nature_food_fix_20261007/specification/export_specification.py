"""Export one verified equation source to native Word mathematics and PDF.

Setup:
  uv pip install --python .venv/bin/python mathml2omml==0.0.2 resvg-py==0.5.0
  npm install --prefix <specification>/.conversion_tools mathjax-full@3.2.2
Run from the repository root:
  .venv/bin/python -m analysis.paper_study.nature_food_fix_20261007.specification.export_specification
"""
from pathlib import Path
from copy import deepcopy
import hashlib
import html
import json
import re
import subprocess
import zipfile

import pymupdf as fitz
import mathml2omml
import resvg_py
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle

DEST = Path(__file__).resolve().parent
SOURCE = DEST/'Supplementary_Model_Specification.md'
ASSETS = DEST/'rendered_equations'
MATH_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
WORD_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
WIDTH = A4[0]-2*51
EQUATION_WIDTH = WIDTH-36
FS = 11.


def equation_layout(tex, label):
    """Wrap independent top-level clauses; retain every mathematical token."""
    if '\\begin{aligned}' in tex:
        return tex
    if label == 'S23':
        return '\\begin{aligned}&' + tex.replace('\n+\\mathbf1', '\\\\\n&+\\mathbf1') + '\\end{aligned}'
    clauses=tex.split('\\qquad')
    if len(clauses)>1:
        return '\\begin{aligned}&'+'\\\\\n&'.join(x.strip() for x in clauses)+'\\end{aligned}'
    return tex


def parse_source():
    blocks=[];requests=[]
    for raw in SOURCE.read_text().strip().split('\n\n'):
        if not raw.strip():continue
        if raw.startswith('\\['):
            body=raw.removeprefix('\\[').removesuffix('\\]').strip()
            label=re.search(r'\\tag\{(S\d+)\}',body).group(1)
            tex=re.sub(r'\\tag\{S\d+\}', '',body).strip()
            idx=len(requests);requests.append(dict(id=idx,tex=equation_layout(tex,label),display=True,label=label))
            blocks.append(dict(kind='equation',id=idx,label=label));continue
        if raw.startswith('#'):
            blocks.append(dict(kind='heading',text=raw.lstrip('#').strip(),level=2 if raw.startswith('###') else 1));continue
        pieces=[];cursor=0
        for found in re.finditer(r'\\\((.*?)\\\)',raw,flags=re.S):
            pieces.append(('text',raw[cursor:found.start()]))
            idx=len(requests);requests.append(dict(id=idx,tex=found.group(1),display=False))
            pieces.append(('math',idx));cursor=found.end()
        pieces.append(('text',raw[cursor:]));blocks.append(dict(kind='paragraph',pieces=pieces))
    labels=[x['label'] for x in blocks if x['kind']=='equation']
    assert labels==[f'S{x}' for x in range(1,30)],labels
    return blocks,requests


def omml(mml, font_size):
    content=mathml2omml.convert(mml)
    # mathml2omml 0.0.2 closes accent properties with the parent tag. Correct
    # those two documented serializer branches without changing their payload.
    content=re.sub(r'(<m:groupChrPr>.*?<m:pos m:val="(?:top|bot)"/>)</m:groupChr>',
        r'\1</m:groupChrPr>',content,flags=re.S)
    wrapper=etree.fromstring((f'<root xmlns:m="{MATH_NS}" xmlns:w="{WORD_NS}">{content}</root>').encode())
    element=wrapper[0]
    for run in element.xpath('.//m:r',namespaces={'m':MATH_NS}):
        rp=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts')
        fonts.set(qn('w:ascii'),'Cambria Math');fonts.set(qn('w:hAnsi'),'Cambria Math');rp.append(fonts)
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(round(font_size*2)));rp.append(sz)
        run.insert(1 if len(run) and run[0].tag==qn('m:rPr') else 0,rp)
    return element


def image_math(record):
    svg=record['svg'];root=etree.fromstring(svg.encode())
    vb=[float(x) for x in root.attrib['viewBox'].split()]
    size=FS if not record['display'] else 11.5
    width,height=vb[2]*size/1000,vb[3]*size/1000
    if record['display'] and width>EQUATION_WIDTH:
        scale=EQUATION_WIDTH/width;width*=scale;height*=scale;size*=scale
    scalept=size/1000
    depth=max(0,(vb[1]+vb[3])*scalept)
    # Explicit physical pixel dimensions make the SVG independent of CSS ex conversion.
    root.set('width',str(width*8));root.set('height',str(height*8))
    svg_path=ASSETS/f"math_{record['id']:03d}.svg"
    svg_path.write_bytes(etree.tostring(root))
    png_path=svg_path.with_suffix('.png')
    png_path.write_bytes(resvg_py.svg_to_bytes(svg_string=svg_path.read_text()))
    record.update(width_pt=width,height_pt=height,depth_pt=depth,font_size_pt=size,png=str(png_path))


def word_text(paragraph,value):
    text=value.replace('`','').replace('\n',' ')
    for piece in re.split(r'(⁻[¹²]|[¹²])',text):
        if not piece:continue
        if piece in ('⁻¹','⁻²','¹','²'):
            run=paragraph.add_run({'⁻¹':'-1','⁻²':'-2','¹':'1','²':'2'}[piece]);run.font.superscript=True
        else:paragraph.add_run(piece)


def docx_output(blocks, maths):
    doc=Document();section=doc.sections[0]
    section.page_width=Pt(A4[0]);section.page_height=Pt(A4[1]);section.left_margin=Pt(51);section.right_margin=Pt(51)
    section.top_margin=Pt(49);section.bottom_margin=Pt(48)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(FS)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.14
    normal.paragraph_format.widow_control=True
    for name,size in [('Title',16),('Heading 1',14),('Heading 2',12)]:
        doc.styles[name].font.name='Times New Roman';doc.styles[name].font.size=Pt(size)
        doc.styles[name].font.color.rgb=None
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Supplementary Methods | ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    for block_index,block in enumerate(blocks):
        if block['kind']=='heading':
            p=doc.add_paragraph(block['text'],'Title' if block['level']==1 else 'Heading 2')
            p.paragraph_format.keep_with_next=True;continue
        if block['kind']=='equation':
            record=maths[block['id']];table=doc.add_table(rows=1,cols=2);table.alignment=WD_TABLE_ALIGNMENT.CENTER
            table.autofit=False;table.columns[0].width=Pt(EQUATION_WIDTH);table.columns[1].width=Pt(36)
            row=table.rows[0];cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
            for cell in row.cells:cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=row.cells[0].paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            op=OxmlElement('m:oMathPara');props=OxmlElement('m:oMathParaPr');jc=OxmlElement('m:jc');jc.set(qn('m:val'),'center')
            props.append(jc);op.append(props);op.append(omml(record['mml'],record['font_size_pt']));p._p.append(op)
            number=row.cells[1].paragraphs[0];number.alignment=WD_ALIGN_PARAGRAPH.RIGHT;number.add_run('('+block['label']+')')
            continue
        p=doc.add_paragraph()
        if block_index+1<len(blocks) and blocks[block_index+1]['kind']=='equation':
            p.paragraph_format.keep_with_next=True
        for kind,value in block['pieces']:
            if kind=='text':word_text(p,value)
            else:p._p.append(omml(maths[value]['mml'],FS))
    path=DEST/'Supplementary_Model_Specification.docx';doc.save(path)
    return path


def pdf_output(blocks, maths):
    fontroot=Path('/System/Library/Fonts/Supplemental')
    for name,file in [('TimesDoc','Times New Roman.ttf'),('TimesDocBold','Times New Roman Bold.ttf'),('TimesDocItalic','Times New Roman Italic.ttf')]:
        pdfmetrics.registerFont(TTFont(name,str(fontroot/file)))
    body=ParagraphStyle('Body',fontName='TimesDoc',fontSize=FS,leading=14,spaceAfter=8,allowWidows=0,allowOrphans=0)
    leadin=ParagraphStyle('Equation lead-in',parent=body,keepWithNext=True)
    title=ParagraphStyle('Title',fontName='TimesDocBold',fontSize=16,leading=20,spaceAfter=14,keepWithNext=True)
    heading=ParagraphStyle('Heading',fontName='TimesDocBold',fontSize=12,leading=15,spaceBefore=10,spaceAfter=7,keepWithNext=True)
    num=ParagraphStyle('Equation number',fontName='TimesDoc',fontSize=FS,alignment=2)
    story=[]
    for block_index,block in enumerate(blocks):
        if block['kind']=='heading':story.append(Paragraph(html.escape(block['text']),title if block['level']==1 else heading));continue
        if block['kind']=='equation':
            record=maths[block['id']];image=Image(record['png'],width=record['width_pt'],height=record['height_pt'])
            table=Table([[image,Paragraph('('+block['label']+')',num)]],colWidths=[EQUATION_WIDTH,36])
            table.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ALIGN',(0,0),(0,0),'CENTER'),
                ('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),0),
                ('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
            story.append(table);continue
        parts=[]
        for kind,value in block['pieces']:
            if kind=='text':
                text=html.escape(value.replace('`','').replace('\n',' '))
                # Paragraph's image-fragment tokenizer discards ordinary edge
                # spaces; protected spaces retain the source's inline spacing.
                if text.startswith(' '):text='&#160;'+text[1:]
                if text.endswith(' '):text=text[:-1]+'&#160;'
                for source,replacement in [('⁻¹','<super>-1</super>'),('⁻²','<super>-2</super>'),('¹','<super>1</super>'),('²','<super>2</super>')]:
                    text=text.replace(source,replacement)
                parts.append(text)
            else:
                r=maths[value]
                parts.append(f'<img src="{r["png"]}" width="{r["width_pt"]:.4f}" height="{r["height_pt"]:.4f}" valign="{-r["depth_pt"]:.4f}"/>')
        style=leadin if block_index+1<len(blocks) and blocks[block_index+1]['kind']=='equation' else body
        story.append(Paragraph(''.join(parts),style))
    path=DEST/'Supplementary_Model_Specification.pdf'
    def page_footer(canvas,document):
        canvas.setFont('TimesDoc',9);canvas.setFillColor(colors.HexColor('#555555'))
        canvas.drawCentredString(A4[0]/2,27,f'Supplementary Methods | {document.page}')
    document=SimpleDocTemplate(str(path),pagesize=A4,leftMargin=51,rightMargin=51,topMargin=49,bottomMargin=48,
        title='Supplementary Methods: mathematical specification',author='Gang Zhao',pageCompression=1)
    document.build(story,onFirstPage=page_footer,onLaterPages=page_footer)
    return path


def verify(docx_path,pdf_path,maths):
    with zipfile.ZipFile(docx_path) as z:
        content=z.read('word/document.xml');xml=etree.fromstring(content)
        ns={'m':MATH_NS,'w':WORD_NS}
        equations=xml.xpath('//m:oMathPara',namespaces=ns)
        native=xml.xpath('//m:oMath',namespaces=ns)
        assert len(equations)==29
        text=' '.join(xml.xpath('//w:t/text()|//m:t/text()',namespaces=ns))
        assert not re.search(r'\\(?:\[|\]|\(|\)|begin|end|frac|tag|mathbf|mathcal|operatorname)',text)
        assert not xml.xpath('//m:t[contains(text(),"&")]',namespaces=ns)
        assert '⁻' not in text
    pdf=fitz.open(pdf_path);pdftext='\n'.join(page.get_text() for page in pdf)
    assert not re.search(r'\\(?:\[|\]|\(|\)|begin|end|frac|tag|mathbf|mathcal|operatorname)',pdftext)
    labels=re.findall(r'\(S\d+\)',pdftext);assert labels==[f'(S{x})' for x in range(1,30)],labels
    bounds=[]
    for i,page in enumerate(pdf):
        for item in page.get_text('dict')['blocks']:
            rect=fitz.Rect(item['bbox'])
            if rect.x0<35 or rect.x1>page.rect.width-35 or rect.y0<25 or rect.y1>page.rect.height-18:
                bounds.append(dict(page=i+1,bbox=list(rect)))
        page.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(DEST/f'preview_page_{i+1:02d}.png')
    assert not bounds,bounds
    report=dict(status='complete',source_sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        numbered_display_equations=29,native_word_math_objects=len(native),inline_math_objects=len(native)-29,
        pdf_pages=len(pdf),raw_latex_commands_in_visible_docx=False,raw_latex_commands_in_pdf_text=False,
        missing_equation_numbers=[],content_bounds_violations=bounds,
        superscript_units='Native Word superscript runs and PDF superscript text; unsupported U+207B glyph avoided.',
        word_equations='Native Office Math (OMML), editable in Microsoft Word.',
        pdf_equations='MathJax SVG rasterized at 576 dpi; shared mathematical source with native Word export.',
        minimum_display_equation_font_pt=min(x['font_size_pt'] for x in maths.values() if x['display']),
        output_sha256={x.name:hashlib.sha256(x.read_bytes()).hexdigest() for x in [docx_path,pdf_path]},
        conversion_dependencies={'mathjax-full':'3.2.2','mathml2omml':'0.0.2','resvg-py':'0.5.0'},
        conversion_script_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        math_renderer_sha256=hashlib.sha256((DEST/'render_math.cjs').read_bytes()).hexdigest(),
        limitation='DOCX native mathematics has structural checks; Microsoft Word layout was not rendered in this environment.')
    (DEST/'rendering_checks.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))


def main():
    ASSETS.mkdir(exist_ok=True);blocks,requests=parse_source()
    request=ASSETS/'math_requests.json';output=ASSETS/'math_rendered.json'
    request.write_text(json.dumps(requests))
    subprocess.run(['node',str(DEST/'render_math.cjs'),str(request),str(output)],check=True)
    maths={x['id']:x for x in json.loads(output.read_text())}
    for record in maths.values():image_math(record)
    docx_path=docx_output(blocks,maths);pdf_path=pdf_output(blocks,maths)
    verify(docx_path,pdf_path,maths)


if __name__=='__main__':
    main()
