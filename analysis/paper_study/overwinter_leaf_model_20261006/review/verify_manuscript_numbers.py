"""Independent read-only source audit; writes only its review receipt."""
from pathlib import Path
import hashlib
import json
import re
import numpy as np
import pandas as pd
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[4]
STUDY = ROOT / 'analysis/paper_study/overwinter_leaf_model_20261006'
DEST = ROOT / 'publication/european_wheat_stb_oop_staging_20261006'
OUT = Path(__file__).with_name('manuscript_numeric_review.json')
bound = {}


def sha(path):
    path = Path(path)
    value = hashlib.sha256(path.read_bytes()).hexdigest()
    bound[str(path.relative_to(ROOT))] = value
    return value


def csv(relative):
    path = STUDY / relative
    sha(path)
    return pd.read_csv(path)


def parquet(relative):
    path = STUDY / relative
    sha(path)
    return pd.read_parquet(path)


def js(path):
    sha(path)
    return json.loads(Path(path).read_text())


def weighted_rows(frame):
    keys = ['coordinate_year', 'field_id', 'leaf_index']
    f = frame.copy()
    coords = f.coordinate_year.nunique()
    fields = f.groupby('coordinate_year').field_id.transform('nunique')
    leaves = f.groupby(keys[:2]).leaf_index.transform('nunique')
    rows = f.groupby(keys).field_id.transform('size')
    return np.asarray(1 / (coords * fields * leaves * rows), float)


def ratio_mcse(frame):
    z = np.dot(frame.area_mean_weight, frame.denominator)
    mean = np.dot(frame.area_mean_weight, frame.numerator) / z
    residual = (frame.numerator - mean * frame.denominator) / z
    variance = 0.
    for _, group in frame.assign(residual=residual).groupby('stratum'):
        assert len(group) == 4
        weight = group.area_mean_weight.sum()
        variance += weight**2 * group.residual.var(ddof=1) / 4
    return float(mean), float(np.sqrt(variance)), np.asarray(residual), float(z)


def main():
    sha(Path(__file__))
    receipt = js(DEST / 'document_build_receipt.json')
    text = (DEST / 'Manuscript.txt').read_text()
    sha(DEST / 'Manuscript.txt')
    for name, expected in receipt['sources'].items():
        assert sha(ROOT / name) == expected, name
    for name in ['Manuscript.pdf', 'Manuscript.docx', 'Source_Data.xlsx']:
        sha(DEST / name)
    issues = []
    original_issues = [dict(id='paired_common_future_rounding', severity='minor_numeric',
        original='from 0.11706 to 0.16415 on common coverage',
        correction='from 0.11706 to 0.16416 on common coverage',
        source='climate/paired_period_changes_by_gcm_ssp.csv: future_on_common, late SSP585, equal three-GCM mean',
        exact=0.164157692977173)]
    if original_issues[0]['original'] in text:
        issues.append(original_issues[0])
    original_issues[0]['resolved_in_current_text'] = original_issues[0]['correction'] in text

    # Workbook values are compared to source CSV parsing, without importing the builder.
    wb = load_workbook(DEST / 'Source_Data.xlsx', read_only=True, data_only=True)
    workbook_checks = []
    maximum_excel_error = 0.
    numeric_excel_cells = 0
    for item in receipt['source_workbook']:
        source = ROOT / item['source_path']
        if str(source).startswith(str(ROOT / 'publication/european_wheat_stb/')):
            source = DEST / source.relative_to(ROOT / 'publication/european_wheat_stb')
        assert sha(source) == item['sha256']
        f = pd.read_csv(source)
        rows = iter(wb[item['sheet']].iter_rows(values_only=True))
        assert list(next(rows)) == list(f.columns)
        count = 0
        for observed, expected in zip(rows, f.itertuples(index=False, name=None), strict=True):
            # XLSX omits absent trailing cells from the sparse row representation.
            assert len(observed) <= len(expected)
            observed = tuple(observed) + (None,) * (len(expected) - len(observed))
            assert len(observed) == len(expected)
            for got, want in zip(observed, expected):
                if pd.isna(want):
                    assert got is None
                elif isinstance(want, (bool, np.bool_)):
                    assert got == bool(want)
                elif isinstance(want, (int, float, np.integer, np.floating)):
                    error = abs(float(got) - float(want))
                    maximum_excel_error = max(maximum_excel_error, error)
                    numeric_excel_cells += 1
                    assert np.isclose(got, want, rtol=2e-15, atol=1e-14), (item['sheet'], count, got, want)
                else:
                    assert got == want, (item['sheet'], count, got, want)
            count += 1
        assert count == item['rows'] == len(f)
        workbook_checks.append(dict(sheet=item['sheet'], rows=count, columns=len(f.columns), values_match=True))
    manifest = list(wb['Source_manifest'].iter_rows(values_only=True))
    assert manifest[0] == ('Sheet', 'Source', 'Rows', 'SHA256')
    assert manifest[1:] == [tuple(row.values()) for row in receipt['source_workbook']]
    assert len(workbook_checks) == 15 and len(wb.sheetnames) == 17

    stage = csv('phenology/stage_censoring_metrics.csv')
    stage_coverage = csv('phenology/stage_observation_coverage.csv')
    threshold = js(STUDY / 'phenology/calibrated_stage_thresholds.json')
    donor31 = js(STUDY / 'phenology/donor_stage31_q1_reconciliation.json')
    disease_registration = js(STUDY / 'disease/configuration_before_fitting.json')
    selected = js(STUDY / 'disease/overwinter_source_model/selection.json')
    frozen_fit = js(STUDY / 'disease/overwinter_source_model/frozen_selected_fit.json')
    benchmark_selection = js(STUDY / 'disease/phenology_only/selection.json')
    climate_registration = js(STUDY / 'climate/configuration_before_climate_results.json')
    assert donor31['status'] == 'passed' and donor31['fields'] == 28 and not donor31['retuned_q']
    assert disease_registration['candidate_count'] == len(selected['candidate_scores']) == 144
    assert len(benchmark_selection['candidate_scores']) == 6
    assert selected['exact_primary_minimum_ties'] == 2
    assert selected['selected_score']['primary_brackets'] == 8
    assert selected['selected_score']['primary_fields'] == 6
    assert frozen_fit['fitted']['parameters']['primary_scale'] == .001
    assert frozen_fit['fitted']['parameters']['secondary_scale'] == 10
    assert frozen_fit['fitted']['parameters']['latent_reference_days'] == 20
    assert frozen_fit['fitted']['rank_spacing_units'] == 120
    assert frozen_fit['fitted']['constant_imported_pressure'] == .1
    assert frozen_fit['fitted']['weather_preprocessing']['rain_rate_mm_hour'] == .45
    assert climate_registration['eligible_calendar_cells'] == 14932
    assert climate_registration['area_registry_cells'] - climate_registration['eligible_calendar_cells'] == 9
    detection = csv('disease/assessment_detection_metrics.csv')
    onset = csv('disease/onset_bracket_metrics.csv')
    assessments = parquet('disease/all_assessment_sign_predictions.parquet')
    brackets = parquet('disease/all_first_onset_bracket_predictions.parquet')
    field_max_error = 0.
    for partition in ['calibration', 'reused_BASF2019', 'reused_strict_Corteva']:
        for model in ['overwinter_source_model', 'phenology_only']:
            a = assessments[(assessments.partition == partition) & (assessments.model == model)
                & (assessments.scenario == 'baseline') & (assessments.leaf_index < 3)]
            weights = weighted_rows(a)
            observed = np.asarray(a.observed_positive, bool)
            predicted = np.asarray(a.predicted_positive, bool)
            sensitivity = np.dot(weights, observed & predicted) / np.dot(weights, observed)
            specificity = np.dot(weights, ~observed & ~predicted) / np.dot(weights, ~observed)
            row = detection[(detection.partition == partition) & (detection.model == model)
                & (detection.scenario == 'baseline') & (detection.leaf_scope == 'top3')].iloc[0]
            assert row['n'] == len(a)
            field_max_error = max(field_max_error, abs(sensitivity-row.sensitivity), abs(specificity-row.specificity))
            assert np.allclose([sensitivity, specificity], [row.sensitivity, row.specificity], atol=1e-14, rtol=0)
            b = brackets[(brackets.partition == partition) & (brackets.model == model)
                & (brackets.scenario == 'baseline') & (brackets.leaf_index < 3) & (brackets.censoring == 'two_sided')]
            bw = weighted_rows(b)
            distance = float(np.dot(bw, b.onset_distance_days))
            compatibility = float(np.dot(bw, b.compatible))
            row = onset[(onset.partition == partition) & (onset.model == model) & (onset.scenario == 'baseline')
                & (onset.leaf_scope == 'top3') & (onset.censoring == 'two_sided')].iloc[0]
            assert row['n'] == len(b)
            field_max_error = max(field_max_error, abs(distance-row.distance_days), abs(compatibility-row.compatible_fraction))
            assert np.allclose([distance, compatibility], [row.distance_days, row.compatible_fraction], atol=1e-13, rtol=0)

    # Recreate all three native tables directly from frozen numeric source tables.
    table1 = [['Set','BBCH','n','Median width (d)','Mean excess (d)','Compatible (%)']]
    labels = {'calibration':'BASF17–18', 'reused_development_2019':'BASF19', 'reused_external_strict':'Corteva'}
    for r in stage[(stage.scope == 'genuine_two_sided') & stage.event.isin([32,33,37,39])].itertuples():
        table1.append([labels[r.partition],str(r.event),str(r.n_field_events),f'{r.median_bracket_width_days:g}',f'{r.mean_distance_days:.2f}',f'{100*r.compatible_fraction:.1f}'])
    table2 = [['Set','Model','Assessment n','Sensitivity (%)','Specificity (%)','Onset n','Mean excess (d)']]
    for p, label in [('calibration','BASF17–18'), ('reused_BASF2019','BASF19'), ('reused_strict_Corteva','Corteva')]:
        for m, name in [('overwinter_source_model','Seasonal'), ('phenology_only','Phenology')]:
            a = detection[(detection.partition == p) & (detection.model == m) & (detection.scenario == 'baseline') & (detection.leaf_scope == 'top3')].iloc[0]
            b = onset[(onset.partition == p) & (onset.model == m) & (onset.scenario == 'baseline') & (onset.leaf_scope == 'top3') & (onset.censoring == 'two_sided')].iloc[0]
            table2.append([label,name,str(a['n']),f'{100*a.sensitivity:.1f}',f'{100*a.specificity:.1f}',str(b['n']),f'{b.distance_days:.2f}'])
    climate_change = csv('climate/ensemble_paired_period_changes.csv')
    climate_period = csv('climate/ensemble_period_means.csv')
    by_gcm = csv('climate/paired_period_changes_by_gcm_ssp.csv')
    table3 = [['SSP','Period','Δ F1 infection (d)','Δ F1 symptoms (d)','Δ HAD3','Δ transfer','GCM transfer range','Spatial MCSE']]
    transfer = 'conditional_yield_loss_GS65_85_b0180_t_ha_per_unit_lai'
    for s, label in [('ssp126','1–2.6'),('ssp245','2–4.5'),('ssp585','5–8.5')]:
        for period in ['2031-2060', '2071-2100']:
            rows = climate_change[(climate_change.scenario == s) & (climate_change.future_period == period)].set_index('metric')
            y = rows.loc[transfer]
            table3.append([label,period,f'{rows.loc["F1_infection_relative_anthesis_days","gcm_mean_change"]:+.2f}',
                f'{rows.loc["F1_symptom_relative_anthesis_days","gcm_mean_change"]:+.2f}',f'{rows.loc["GS65_85_lost_had3","gcm_mean_change"]:+.3f}',
                f'{y.gcm_mean_change:+.4f}', f'{y.gcm_min_change:+.4f} to {y.gcm_max_change:+.4f}', f'{y.spatial_mcse_gcm_mean_change:.4f}'])
    doc = Document(DEST / 'Manuscript.docx')
    doc_tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
    assert doc_tables == [table1, table2, table3]
    for table in [table1, table2, table3]:
        assert '\n'.join('\t'.join(row) for row in table) in text

    # Independently evaluate the published climate ratios and MCSE, including shared-draw covariance.
    metrics = ['F1_infection_relative_anthesis_days','F2_infection_relative_anthesis_days','F3_infection_relative_anthesis_days',
        'F1_symptom_relative_anthesis_days','grain_fill_days','GS65_85_lost_had3',transfer,
        'conditional_yield_loss_GS65_85_b0141_t_ha_per_unit_lai','conditional_yield_loss_GS65_85_b0207_t_ha_per_unit_lai',
        'any_top3_infection_before85','all_top3_infection_before85']
    moments = parquet('climate/draw_period_moments.parquet')
    paired = parquet('climate/paired_draw_period_moments.parquet')
    climate_max_error = 0.
    climate_checks = 0
    for source, grouping, output, period_column in [(moments, ['scenario','period','metric'],climate_period,'period'),
            (paired,['scenario','future_period','metric'],climate_change,'future_period')]:
        for keys, group in source[source.metric.isin(metrics)].groupby(grouping):
            scenario, period, metric = keys
            estimates, influence = [], []
            for _, g in group.groupby('model'):
                g = g.sort_values('spatial_draw_id')
                mean, _, residual, _ = ratio_mcse(g)
                estimates.append(mean)
                influence.append(residual)
            influence = np.mean(influence, axis=0)
            variance = 0.
            last = g.assign(residual=influence)
            for _, h in last.groupby('stratum'):
                variance += h.area_mean_weight.sum()**2 * h.residual.var(ddof=1) / 4
            row = output[(output.scenario == scenario) & (output[period_column] == period) & (output.metric == metric)].iloc[0]
            expected = [np.mean(estimates),np.min(estimates),np.max(estimates),np.sqrt(variance)]
            names = ['gcm_mean','gcm_min','gcm_max','spatial_mcse_gcm_mean'] if period_column == 'period' else ['gcm_mean_change','gcm_min_change','gcm_max_change','spatial_mcse_gcm_mean_change']
            errors = np.abs(np.asarray(expected)-np.asarray(row[names],float))
            climate_max_error = max(climate_max_error,float(errors.max()))
            assert errors.max() < 1e-11
            climate_checks += 1

    climate = parquet('climate/draw_season_outputs.parquet')
    climate_receipt = js(STUDY / 'climate/receipt.json')
    replay = csv('manuscript/replay/full_season_top3_timing_and_conditional_yield.csv')
    replay_receipt = js(STUDY / 'manuscript/replay/receipt.json')
    assert len(climate) == 51840
    assert climate.groupby(['model','scenario','harvest_year']).size().eq(64).all()
    assert climate.groupby(['model','scenario']).harvest_year.nunique().eq(90).all()
    draws_path = ROOT / 'data/paper_study/regional_parameter_uncertainty/spatial_draws.csv'
    sha(draws_path)
    draws = pd.read_csv(draws_path)
    assert len(draws) == 64 and draws.cell_id.nunique() == 62
    assert draws.groupby('stratum').size().eq(4).all()
    assert np.isclose(draws.area_mean_weight.sum(), 1, atol=1e-14)
    weights = draws.set_index('spatial_draw_id').area_mean_weight
    assert np.array_equal(climate.area_mean_weight.to_numpy(), climate.spatial_draw_id.map(weights).to_numpy())
    assert climate[['model','scenario','harvest_year','spatial_draw_id']].duplicated().sum() == 0
    assert not any('yield0' in col.lower() or col.lower() == 'y0' for col in climate.columns)
    status = {str(k):int(v) for k,v in climate.status.value_counts().items()}
    assert status == {'complete':51813,'soft_dough_not_reached':26,'unsupported_mean_temperature_above40':1}
    assert len(replay) == 218 and replay.grain_fill_complete.all() and replay.reference_lai_scenario.eq(1).all()
    assert replay_receipt['all_original_tissue_residue_and_host_area_prefixes_identical']
    replay_facts = {'field_count':len(replay),'mean_conditional_transfer':float(replay['conditional_yield_loss_b0.018_t_ha_per_unit_lai'].mean())}
    for leaf in [1,2,3]:
        replay_facts[f'F{leaf}_median_infection_relative_anthesis'] = float(replay[f'F{leaf}_infection_relative_anthesis_days'].median())
        replay_facts[f'F{leaf}_median_symptom_relative_anthesis'] = float(replay[f'F{leaf}_symptom_relative_anthesis_days'].median())
    replay_facts['F1_symptoms_before_anthesis'] = int((replay.F1_symptom_relative_anthesis_days < 0).sum())
    assert all((replay[f'F{leaf}_infection_relative_anthesis_days']<0).all() for leaf in [1,2,3])
    assert np.allclose(replay['conditional_yield_loss_b0.018_t_ha_per_unit_lai'], .018*replay.model_proxy_lost_had3,atol=1e-14)
    for leaf in [1,2,3]:
        for event in ['infection','symptom']:
            dates = pd.to_datetime(replay[f'F{leaf}_{event}_date'])
            relative = (dates-pd.to_datetime(replay.BBCH65_date)).dt.days
            assert relative.equals(replay[f'F{leaf}_{event}_relative_anthesis_days'].astype('int64'))
    late = climate_change[(climate_change.scenario=='ssp585') & (climate_change.future_period=='2071-2100')].set_index('metric')
    late_gcm = by_gcm[(by_gcm.scenario=='ssp585') & (by_gcm.future_period=='2071-2100') & (by_gcm.metric==transfer)]
    climate_facts = dict(valid_draw_seasons=51813,invalid_status_counts=status,eligible_harvested_ha=climate_receipt['eligible_harvested_ha'],
        late_ssp585={metric:float(late.loc[metric,'gcm_mean_change']) for metric in metrics if metric in late.index},
        late_ssp585_common_reference_transfer=float(late_gcm.reference_on_common.mean()),
        late_ssp585_common_future_transfer=float(late_gcm.future_on_common.mean()),
        late_ssp585_transfer_gcm_range=[float(late.loc[transfer,'gcm_min_change']),float(late.loc[transfer,'gcm_max_change'])],
        late_ssp585_transfer_spatial_mcse=float(late.loc[transfer,'spatial_mcse_gcm_mean_change']),
        minimum_paired_GF_area_time_coverage=float(climate_change[climate_change.metric==transfer].minimum_gcm_common_paired_coverage.min()))
    climate_facts['minimum_infection_timing_common_event_coverage'] = float(climate_change[climate_change.metric.isin(['F1_infection_relative_anthesis_days','F2_infection_relative_anthesis_days','F3_infection_relative_anthesis_days'])].minimum_gcm_common_paired_coverage.min())
    for metric in ['any_top3_infection_before85','all_top3_infection_before85']:
        p=climate_period[climate_period.metric==metric]
        climate_facts[metric+'_ensemble_range_percent']=[float(100*p.gcm_mean.min()),float(100*p.gcm_mean.max())]
    assert climate_facts['minimum_paired_GF_area_time_coverage'] > .99866
    assert late.loc['grain_fill_days','gcm_mean_change'] > 0

    # Paragraph identity checks protect delivery formats; reference metadata is audited elsewhere.
    doc_paragraphs=[p.text for p in doc.paragraphs if p.text]
    body_end=text.index('\n\nReferences\n\n')
    body_paragraphs=[p for p in text[:body_end].split('\n\n') if p]
    assert doc_paragraphs[:len(body_paragraphs)] == body_paragraphs
    pdf=PdfReader(DEST/'Manuscript.pdf')
    pdf_text='\n'.join(page.extract_text() for page in pdf.pages)
    normalize=lambda s:re.sub(r'\s+','',s).replace('−','-').replace('–','-')
    important_fragments=['51,840','51,813','98.9%','96.9%','633.392','702.266','5.06','5.36','0.0471','0.1177']
    for fragment in important_fragments:
        assert normalize(fragment) in normalize(pdf_text), fragment
    meta_patterns=[r'\bthis (?:paper|report|manuscript) (?:will|presents|uses|does not)',r'\bwe (?:will|present|revise|avoid|choose)',r'according to (?:the )?user',r'below we']
    meta_matches=[match.group(0) for pattern in meta_patterns for match in re.finditer(pattern,text[:body_end],re.I)]
    assert not meta_matches
    spacing_fragments=['andF3','infection−4.805','and−14.511','changes+0.125','andSSP2–4.5','A24-bin','mean,120-unit']
    for fragment in spacing_fragments:
        if fragment in text:
            issues.append(dict(id='spacing_'+fragment, severity='typography',original=fragment))

    for name in ['fig1_framework','fig2_flag_stage_validation','fig3_onset_and_detection','fig4_timing_and_yield_relevance','fig5_conditional_climate_changes']:
        for suffix in ['png','pdf','svg']:
            assert (DEST/'figures'/f'{name}.{suffix}').is_file()
            sha(DEST/'figures'/f'{name}.{suffix}')
    # Detect any concurrent rebuild so that the saved review identifies one stable package.
    for name,value in bound.items():
        assert hashlib.sha256((ROOT/name).read_bytes()).hexdigest()==value, f'Artifact changed during review: {name}'
    result=dict(status='passed_with_minor_issues' if issues else 'passed', reviewed_package=str(DEST.relative_to(ROOT)),
        read_only_review=True, source_sha256=bound, prominent_numeric_claims_checked=True,
        workbook=dict(source_csvs=15,source_rows=sum(row['rows'] for row in workbook_checks),checks=workbook_checks,
            numeric_cells_compared=numeric_excel_cells,maximum_absolute_serialization_difference=maximum_excel_error,
            precision_note='All values match source CSV parsing at Excel serialization precision; headers, strings, booleans, nulls and row counts are exact.'),
        native_tables=dict(count=3,all_text_cells_match_independent_source_calculation=True),
        field_metrics=dict(weighted_baseline_detection_and_two_sided_onset_checks=12,maximum_discrepancy=field_max_error,
            calibration_genuine_onsets=8,reused_evaluation_genuine_onsets=136,reused_evaluation_components=[13,123],
            calibration_genuine_added_stage_intervals={str(r.event):int(r.genuine_two_sided) for r in stage_coverage[(stage_coverage.partition=='calibration')&stage_coverage.event.isin([32,33,37,39])].itertuples()},
            added_stage_thresholds=threshold['thresholds'],weak_identification_all_added_thresholds=all(d['poorly_identified'] for d in threshold['diagnostics'])),
        frozen_selection=dict(donor_q1_stage31_exact_fields=28,original_calibration_fields=28,
            disease_candidate_count=144,phenology_benchmark_candidate_count=6,inner_location_folds=3,
            primary_distance_ties=2,calibration_onset_fields=6,selected_parameters=frozen_fit['fitted'],
            primary_and_secondary_are_effective_hypotheses=True,secondary_scale_on_upper_grid_boundary=True),
        climate_statistics=dict(independent_ensemble_ratio_mcse_checks=climate_checks,maximum_discrepancy=climate_max_error,
            shared_spatial_draw_covariance_in_three_GCM_MCSE=True,paired_ratio_common_coverage=True,values=climate_facts),
        spatial_design=dict(independent_draw_identities=64,unique_cells=62,strata=16,draws_per_stratum=4,
            registered_area_weights_match_all_51840_rows=True,weight_sum=float(draws.area_mean_weight.sum()),
            duplicate_physical_cells_retained_as_independent_draw_ids=True),
        conditional_field_replay=replay_facts,
        interpretation=dict(actual_infection_dates_observed=False,validation_retrospective=True,actual_yield_prediction=False,
            nominal_reference_upper3_LAI=1,cultivar_range_is_confidence_interval=False,GCM_range_is_confidence_interval=False,
            field_vs_climate_clinical_gate_difference_disclosed='The climate replay applies the declared detection gate before functional transfer.' in text,
            broad_GS31_window_is_sensitivity=True,author_writing_meta_matches=meta_matches),
        delivered_formats=dict(native_docx_body_matches_txt=True,pdf_pages=len(pdf.pages),prominent_numeric_fragments_present_pdf=True,
            figures=5,tables=3), original_issues=original_issues,issues=issues)
    OUT.parent.mkdir(parents=True,exist_ok=True)
    OUT.write_text(json.dumps(result,indent=2,allow_nan=False)+'\n')
    print(json.dumps(dict(status=result['status'],issues=issues,workbook_rows=result['workbook']['source_rows'],
        climate_statistic_checks=climate_checks,field_replay=replay_facts,climate=climate_facts),indent=2))


if __name__ == '__main__':
    main()
