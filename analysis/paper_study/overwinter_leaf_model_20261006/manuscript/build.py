"""Native manuscript, source workbook and versioned research package."""
from pathlib import Path
import hashlib,html,json,re,shutil,subprocess,zipfile
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from reportlab.platypus import SimpleDocTemplate,Paragraph,PageBreak,Image,Table,TableStyle,Spacer,KeepTogether
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from PIL import Image as PILImage

from analysis.paper_study.manuscript import build_documents as layout
from analysis.paper_study.manuscript.author_date_citations import plain
from .sections import sections,TITLE,HERE,STUDY
from .figures import main as draw_figures
from .figures import export
from analysis.paper_study.full_validation_20261007.evidence import HERE as VALIDATION
from analysis.paper_study.full_validation_20261007.figures import supplementary as supplementary_figures
from analysis.paper_study.full_validation_20261007.tables import supplementary_tables
from analysis.paper_study.yield_evidence_20261007.published_response import HERE as YIELD_EVIDENCE
from analysis.paper_study.yield_evidence_20261007.figures import supplementary as yield_figures
from analysis.paper_study.yield_evidence_20261007.tables import supplementary_tables as yield_tables
from analysis.paper_study.map_restoration_20261007.prepare import HERE as MAPS
from analysis.paper_study.map_restoration_20261007.figures import supplementary as map_figures
from analysis.paper_study.map_restoration_20261007.tables import supplemental_tables as crop_validation_tables

ROOT=STUDY.parents[2]
DEST=ROOT/'publication/european_wheat_stb_oop_staging_20261006'
TOKEN=r'\[\[([A-Za-z0-9_|]+)\]\]'


def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def polish(text):
    text=re.sub(r'(?<=[a-z])(?=\d)',' ',text)
    text=re.sub(r'(?<=[a-z])(?=[A-Z])',' ',text)
    text=re.sub(r'(?<=[A-Za-z])(?=[+−]\d)',' ',text)
    text=re.sub(r'\bA(?=\d)','A ',text)
    text=re.sub(r'\b(MAE|RMSE|SED|HAD|R²)(?=\d)',r'\1 ',text)
    text=re.sub(r'\b(BBCH|BASF|SSP|GS|LAI)(?=\d)',r'\1 ',text)
    text=re.sub(r'(?<=\d)(?=[a-z]{2})',' ',text)
    text=re.sub(r';(?=\S)','; ',text)
    text=re.sub(r'(?<!\d),(?=\S)',', ',text)
    text=re.sub(r',(?!\d{3}\b(?!\.\d))(?=\S)',', ',text)
    return text


def citations(raw):
    records=json.loads((HERE/'references/references.csl.json').read_text())
    lookup={row['id']:row for row in records}
    tokens=sorted(set(re.findall(TOKEN,'\n'.join(text for paragraphs in raw.values() for text in paragraphs))))
    ids=sorted({key for token in tokens for key in token.split('|')})
    citation_input=dict(items=[lookup[key] for key in ids],clusters={token:token.split('|') for token in tokens})
    (DEST/'citation_input.json').write_text(json.dumps(citation_input,indent=2,ensure_ascii=False)+'\n')
    subprocess.run(['node',str(ROOT/'analysis/paper_study/manuscript/render_csl.js'),str(DEST/'citation_input.json'),str(DEST/'citation_rendered.json')],check=True,capture_output=True,text=True)
    rendered=json.loads((DEST/'citation_rendered.json').read_text())
    converted={name:[polish(re.sub(TOKEN,lambda match:plain(rendered['citations'][match[1]]),text)) for text in paragraphs]
        for name,paragraphs in raw.items()}
    (DEST/'references.csl.json').write_text(json.dumps(citation_input['items'],indent=2,ensure_ascii=False)+'\n')
    (DEST/'references_author_date.txt').write_text('References\n\n'+'\n\n'.join(plain(row) for row in rendered['bibliography_html'])+'\n')
    shutil.copy2(HERE/'references/references.bib',DEST/'references.bib')
    return converted,rendered


def tables():
    stage=pd.read_csv(VALIDATION/'field_stage_metrics_all.csv')
    stage=stage.loc[stage.scope.eq('genuine_two_sided') & stage.forcing_scope.eq('full_season')
        & stage.partition.isin(['reused_development_2019','reused_external_strict'])]
    names={'calibration':'BASF17–18','reused_development_2019':'BASF19','reused_external_strict':'Corteva'}
    first=[['Set','BBCH','n','Median width (d)','Mean excess (d)','Compatible (%)']]
    for row in stage.itertuples():first.append([names[row.partition],str(row.event),str(row.n_field_events),
        '—' if pd.isna(row.median_bracket_width_days) else f'{row.median_bracket_width_days:g}',
        '—' if pd.isna(row.mean_distance_days) else f'{row.mean_distance_days:.2f}',
        '—' if pd.isna(row.compatible_fraction) else f'{100*row.compatible_fraction:.1f}'])
    detection=pd.read_csv(STUDY/'disease/assessment_detection_metrics.csv')
    onset=pd.read_csv(STUDY/'disease/onset_bracket_metrics.csv')
    second=[['Set','Model','Scope','Assessment n','Sensitivity (%)','Specificity (%)','Onset n','Mean excess (d)']]
    for partition,label in [('calibration','BASF17–18'),('reused_BASF2019','BASF19'),('reused_strict_Corteva','Corteva')]:
        for model,mname in [('overwinter_source_model','Seasonal'),('phenology_only','Phenology')]:
            for scope,sname in [('top3','F1–F3'),('all_ordinal_leaves','All ranks')]:
                a=detection.loc[detection.partition.eq(partition)&detection.model.eq(model)&detection.scenario.eq('baseline')&detection.leaf_scope.eq(scope)].iloc[0]
                b=onset.loc[onset.partition.eq(partition)&onset.model.eq(model)&onset.scenario.eq('baseline')&onset.leaf_scope.eq(scope)&onset.censoring.eq('two_sided')].iloc[0]
                second.append([label,mname,sname,str(a['n']),f'{a.sensitivity*100:.1f}',f'{a.specificity*100:.1f}',str(b['n']),f'{b.distance_days:.2f}'])
    climate=pd.read_csv(STUDY/'climate/ensemble_paired_period_changes.csv')
    third=[['SSP','Period','Δ F1 infection (d)','Δ F1 symptoms (d)','Δ HAD3','Δ transfer','GCM transfer range','Spatial MCSE']]
    for scenario,slabel in [('ssp126','1–2.6'),('ssp245','2–4.5'),('ssp585','5–8.5')]:
        for period in ['2031-2060','2071-2100']:
            part=climate.loc[climate.scenario.eq(scenario)&climate.future_period.eq(period)].set_index('metric')
            y=part.loc['conditional_yield_loss_GS65_85_b0180_t_ha_per_unit_lai']
            third.append([slabel,period,f"{part.loc['F1_infection_relative_anthesis_days','gcm_mean_change']:+.2f}",
                f"{part.loc['F1_symptom_relative_anthesis_days','gcm_mean_change']:+.2f}",
                f"{part.loc['GS65_85_lost_had3','gcm_mean_change']:+.3f}",f'{y.gcm_mean_change:+.4f}',
                f'{y.gcm_min_change:+.4f} to {y.gcm_max_change:+.4f}',f'{y.spatial_mcse_gcm_mean_change:.4f}'])
    return [(first,'Table1 | All retained stages in genuine two-sided field validation. n counts field-stage events; zero identifies absent interval evidence. Widths and excess are calendar days. Full-season forcing extends the observation horizon without changing existing threshold crossings. Calibration, one-sided constraints and overlapping full-source transfer are retained in Tables S2–S3. These are reused records, not untouched tests.',[65,42,26,96,96,100]),
        (second,'Table2 | Disease assessment detection and first-symptom intervals for F1–F3 and all ordinal ranks under frozen parameters. Sensitivity/specificity and onset excess use equal coordinate-year, field and leaf hierarchy. The eight top-three calibration onset histories cannot identify ecological source and transmission parameters. The phenology benchmark is independently selected on the same calibration membership. The two leaf scopes overlap and are not pooled.',[57,55,49,54,67,67,43,64]),
        (third,'Table3 | Paired three-climate-model mean changes from each SSP-specific1991–2020 reference. HAD3 is GLAI-days per unit nominal upper-three-leaf LAI. Conditional transfer is t ha⁻¹ per unit nominal LAI at b=0.018, using a model functional-loss proxy. GCM range and spatial Monte Carlo standard error are separate uncertainty summaries; neither includes ecological-parameter, host-transfer or observation uncertainty.',[33,58,58,58,43,48,83,47])]


def workbook():
    paths=[STUDY/'phenology'/name for name in ['stage_censoring_metrics.csv','stage_censoring_scores.csv','stage_observation_coverage.csv','modeled_final_leaf_ranks_at_stages.csv','ordered_threshold_joint_profiles.csv']]
    paths += [STUDY/'disease'/name for name in ['assessment_detection_metrics.csv','onset_bracket_metrics.csv','observed_window_occurrence_metrics.csv']]
    paths += [HERE/'replay/full_season_top3_timing_and_conditional_yield.csv',DEST/'figures/fig4_specified_yield_scenarios.csv']
    paths += [STUDY/'climate'/name for name in ['ensemble_period_means.csv','ensemble_paired_period_changes.csv','period_means_by_gcm_ssp.csv','paired_period_changes_by_gcm_ssp.csv','coverage_by_period_and_status.csv']]
    paths += [DEST/'figures/fig2_field_stage_intervals.csv']
    paths += [DEST/'figures/fig2_german_stage_errors.csv']
    paths += [VALIDATION/name for name in ['german_stage_metrics_all.csv','field_stage_constraints_all.csv',
        'field_stage_scores_all.csv','field_stage_metrics_all.csv','field_stage_coverage_all.csv',
        'disease_assessment_detection_metrics_baseline.csv','disease_onset_bracket_metrics_baseline.csv',
        'disease_observed_window_occurrence_metrics_baseline.csv','archived_seasonal_model_results.csv',
        'french_transfer_v1.csv','french_transfer_v2.csv','validation_source_inventory.csv']]
    paths += [YIELD_EVIDENCE/name for name in ['parker2004_table4_yield_loss_matrix.csv',
        'parker2004_table4_yield_loss_long.csv','parker2004_table5_yield_slopes.csv',
        'parker2004_reported_fit_correlations.csv','stb_independent/foulkes2006_table2_paired_predicted_means.csv',
        'stb_independent/foulkes2006_table3_coefficients.csv','stb_independent/foulkes2006_table4_siteyear_slope_deviations.csv',
        'foulkes2006_group_separated_predictions.csv','published_yield_response_metrics.csv',
        'published_yield_domain_compatibility.csv','bryson1997_reported_yield_fits.csv','subbarao1989_reported_yield_fits.csv']]
    paths += [DEST/'supplementary_figures/figS6_leaf_onset_distances.csv']
    paths += [MAPS/name for name in ['map_wheat_domain.csv','map_field_locations.csv','map_draw_values.csv','map_ensemble_values.csv','map_metric_coverage.csv']]
    paths += [ROOT/'data/paper_study/publication'/name for name in ['fig1_wheat_production.csv',
        'fig3_era5_cell_summary.csv','fig4_climate_change_cells.csv','fig6_country_production_exposure.csv']]
    revision=ROOT/'analysis/paper_study/nature_food_revision_20261007'
    paths += [revision/'epidemic_evaluation/severity_trajectory_metrics.csv',
        revision/'yield_response/model_comparison_metrics.csv',revision/'basf_crop_response/year_transfer_metrics.csv']
    wb=Workbook(write_only=True);manifest=[]
    for i,path in enumerate(paths,1):
        frame=pd.read_csv(path);sheet=wb.create_sheet((f'{i:02d}_'+path.stem)[:31]);sheet.freeze_panes='A2';sheet.append(list(frame.columns))
        for row in frame.itertuples(index=False,name=None):sheet.append([None if pd.isna(v) else v.item() if isinstance(v,np.generic) else v for v in row])
        source_path=(Path('publication/european_wheat_stb')/path.relative_to(DEST)
            if path.is_relative_to(DEST) else path.relative_to(ROOT))
        manifest.append(dict(sheet=sheet.title,source_path=str(source_path),rows=len(frame),sha256=sha(path)))
    sheet=wb.create_sheet('Source_manifest');sheet.append(['Sheet','Source','Rows','SHA256'])
    for row in manifest:sheet.append(list(row.values()))
    sheet=wb.create_sheet('Definitions')
    for row in [('Metric','Definition'),('Stage/onset interval excess','Distance outside observed interval; not exact-date MAE'),
        ('Figure 2 observed stage window','Daily admissible dates: last below-stage assessment + 1 day through first at-or-above-stage assessment'),
        ('Figure 2 interval centre','Horizontal display position only; not an observed event date'),
        ('German event error','Exact predicted minus observed date, with station/sowing-cycle/BBCH identity retained'),
        ('German box whiskers','Fifth and ninety-fifth distribution percentiles; not confidence intervals'),
        ('Full Corteva transfer','Includes two fields at locations shared with BASF; overlaps the strict subset; not pooled'),
        ('Archived French transfer','Conditioned on initial pycnidial observations; earlier model versions; not current-model onset validation'),
        ('Field forcing scope','Observed-horizon and full-season scores retained separately; existing threshold dates reconcile exactly'),
        ('Effective infection','Model total affected-fraction threshold event; actual infection unobserved'),
        ('Assessment specificity','Observed-zero dates; not independently identified disease-free complete seasons'),
        ('Model functional loss','Simulated visible affected tissue used as a conditional proxy; actual green-area loss unmeasured'),
        ('Reference upper3LAI','Nominal mature LAI1 with equal mature F1/F2/F3 capacity'),
        ('Conditional transfer','b×HADloss; t/ha per unit nominal LAI; not actual field yield prediction'),
        ('Cultivar range','0.0141–0.0207 t/ha per GLAI-day; not confidence interval'),
        ('Published paired means','Foulkes Table 2 mixed-model genotype-by-fungicide means; shared estimation and unbalanced site coverage'),
        ('Yield-response group holdout','Five genetic backgrounds; each prediction excludes its own background from coefficient estimation'),
        ('Yield-response weights','Equal total weight for each genetic background, divided among its observed lines'),
        ('Relative crop response','Protected-minus-affected HAD or yield divided by corresponding protected mean'),
        ('Leaf-scope compatibility','Top-three STB, top-five paired means, canopy yellow-rust and top-four relative leaf-rust domains remain separate'),
        ('Reported rust R2','Original within-study fit summary; not independent disease or yield prediction accuracy'),
        ('Missing published contrasts','Unreported cultivar/site-year combinations remain missing; negative contrasts are retained'),
        ('Current spatial maps','62 unique sampled cells; no interpolation of uncomputed grid cells'),
        ('Full-grid example','Completed2001 crop replay only; not a completed30-year historical/future census'),
        ('Archived spatial maps','Earlier seasonal-v1 values retain original definitions; not current-model validation'),
        ('Public trial response','Signed1−untreated/treated grain yield; management-associated response, not disease-free absolute loss'),
        ('Severity-window integral','Measured disease percentage-days between common observations; not HAD or absolute green area'),
        ('Spatial MCSE','Sampling uncertainty for registered64 area-proportional draws'),
        ('GCM range','Three deterministic model estimates; not probabilistic confidence interval')]:sheet.append(row)
    wb.save(DEST/'Source_Data.xlsx');return manifest


def render(body,reference_html,captions,table_data):
    layout.TITLE=TITLE
    doc=Document();layout.configure_doc(doc);doc.add_heading(TITLE,0);doc.add_heading('Abstract',1)
    layout.add_prose_doc(doc,body);doc.add_heading('References',1)
    for ref in reference_html:layout.add_reference_doc(doc,ref)
    styles=layout.font_styles();flow=[Paragraph(html.escape(TITLE),styles['PaperTitle']),Paragraph('Abstract',styles['PaperHeading'])]
    layout.pdf_prose(flow,body,styles);flow.append(Paragraph('References',styles['PaperHeading']))
    for ref in reference_html:flow.append(Paragraph(layout.reference_inline(ref),styles['PaperReference']))
    for rows,caption,widths in table_data:
        doc.add_page_break();layout.add_native_table(doc,rows,polish(caption))
        flow += [PageBreak(),Paragraph(html.escape(polish(caption)),styles['PaperCaption'])]
        cells=[[Paragraph(html.escape(str(x)),styles['PaperTable']) for x in row] for row in rows]
        table=Table(cells,colWidths=widths,repeatRows=1)
        table.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eef1f3')),
            ('LINEBELOW',(0,0),(-1,0),.5,colors.grey),('LINEBELOW',(0,-1),(-1,-1),.5,colors.grey),
            ('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
        flow.append(table)
    for name,caption in captions.items():
        path=DEST/'figures'/f'{name}.png'
        doc.add_page_break();doc.add_picture(str(path),width=Inches(6.45));doc.add_paragraph(polish(caption))
        with PILImage.open(path) as im:w,h=im.size
        width=460;height=width*h/w
        if height>520:width*=520/height;height=520
        flow += [PageBreak(),Image(str(path),width,height),Spacer(1,10),Paragraph(html.escape(polish(caption)),styles['PaperCaption'])]
    doc.save(DEST/'Manuscript.docx')
    SimpleDocTemplate(str(DEST/'Manuscript.pdf'),pagesize=A4,rightMargin=45,leftMargin=45,topMargin=45,bottomMargin=45,
        title=TITLE,author='Gang Zhao').build(flow,onFirstPage=layout.footer,onLaterPages=layout.footer)


def render_supplement(paragraphs,table_data,captions):
    doc=Document();layout.configure_doc(doc);layout.add_prose_doc(doc,paragraphs)
    flow=[];styles=layout.font_styles();layout.pdf_prose(flow,paragraphs,styles)
    for rows,caption,widths in table_data:
        doc.add_paragraph();layout.add_native_table(doc,rows,polish(caption))
        doc.paragraphs[-1].paragraph_format.keep_with_next=True
        caption_flow=Paragraph(html.escape(polish(caption)),styles['PaperCaption'])
        cells=[[Paragraph(html.escape(str(value)),styles['PaperTable']) for value in row] for row in rows]
        table=Table(cells,colWidths=widths,repeatRows=1)
        table.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eef1f3')),
            ('LINEBELOW',(0,0),(-1,0),.5,colors.grey),('LINEBELOW',(0,-1),(-1,-1),.5,colors.grey),
            ('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3),
            ('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
        flow += [Spacer(1,12),KeepTogether([caption_flow,table])]
    for name,caption in captions.items():
        path=DEST/'supplementary_figures'/f'{name}.png'
        doc.add_page_break();doc.add_picture(str(path),width=Inches(6.45));doc.add_paragraph(polish(caption))
        with PILImage.open(path) as image:w,h=image.size
        width=460;height=width*h/w
        if height>510:width*=510/height;height=510
        flow += [PageBreak(),Image(str(path),width,height),Spacer(1,9),Paragraph(html.escape(polish(caption)),styles['PaperCaption'])]
    doc.save(DEST/'Supplementary_Information.docx')
    SimpleDocTemplate(str(DEST/'Supplementary_Information.pdf'),pagesize=A4,rightMargin=45,leftMargin=45,
        topMargin=45,bottomMargin=45,title='Supplementary Information',author='Gang Zhao').build(flow,
        onFirstPage=layout.footer,onLaterPages=layout.footer)


def main():
    DEST.mkdir(parents=True,exist_ok=True)
    raw=sections();converted,rendered=citations(raw)
    for name,paras in converted.items():(DEST/name).write_text('\n\n'.join(paras)+'\n')
    order=['abstract.txt','introduction.txt','field_results.txt','baseline_results.txt','climate_results.txt','discussion_field_and_limits.txt','methods.txt','availability.txt']
    body=[text for name in order for text in converted[name]]
    captions=draw_figures(DEST/'figures');moved_tables=tables();table_data=[]
    render(body,rendered['bibliography_html'],captions,table_data)
    text=TITLE+'\n\nAbstract\n\n'+'\n\n'.join(body)+'\n\nReferences\n\n'+'\n\n'.join(plain(r) for r in rendered['bibliography_html'])
    for rows,caption,_ in table_data:text+='\n\n'+polish(caption)+'\n\n'+'\n'.join('\t'.join(map(str,row)) for row in rows)
    text+='\n\nFigure legends\n\n'+'\n\n'.join(polish(caption) for caption in captions.values())+'\n'
    (DEST/'Manuscript.txt').write_text(text)
    supplement_captions=supplementary_figures(DEST/'supplementary_figures',export)
    supplement_captions.update(yield_figures(DEST/'supplementary_figures',export))
    supplement_captions.update(map_figures(DEST/'supplementary_figures',export))
    (DEST/'supplementary_figures/supplementary_captions.json').write_text(json.dumps(supplement_captions,indent=2)+'\n')
    migrated=[(rows,caption.replace(f'Table{index}',f'Table S11{letter}'),widths)
        for index,letter,(rows,caption,widths) in zip([1,2,3],'abc',moved_tables)]
    supplement_tables=supplementary_tables()+yield_tables()+migrated+crop_validation_tables()
    render_supplement(converted['supplementary_methods.txt'],supplement_tables,supplement_captions)
    supplement_text='\n\n'.join(converted['supplementary_methods.txt'])
    for rows,caption,_ in supplement_tables:supplement_text+='\n\n'+polish(caption)+'\n\n'+'\n'.join('\t'.join(map(str,row)) for row in rows)
    supplement_text+='\n\n'+'\n\n'.join(polish(caption) for caption in supplement_captions.values())+'\n'
    (DEST/'Supplementary_Information.txt').write_text(supplement_text)
    manifest=workbook()
    shutil.copy2(ROOT/'WHEAT_STB_ENGINE.md',DEST/'WHEAT_STB_ENGINE.md')
    wheel=ROOT/'publication/european_wheat_stb/wheat_stb_model-0.1.0-py3-none-any.whl'
    if wheel.resolve() != (DEST/wheel.name).resolve():shutil.copy2(wheel,DEST/wheel.name)
    receipt=dict(status='generated',title=TITLE,main_figures=len(captions),main_tables=len(table_data),reference_count=len(rendered['bibliography_html']),
        abstract_words=len(converted['abstract.txt'][0].split()),source_workbook=manifest,
        runtime_only_model_folder=True,calibration_outside_model=True,untouched_test=False,
        direct_infection_dates_observed=False,actual_yield_predictions=False,
        supplementary_figures=len(supplement_captions),supplementary_tables=len(supplement_tables),
        sources={str(p.relative_to(ROOT)):sha(p) for p in [Path(__file__),HERE/'sections.py',HERE/'figures.py',
            HERE/'references/references.csl.json',HERE/'references/references.bib',
            *[VALIDATION/name for name in ['figures.py','workflow_figure.py','tables.py','evidence.py','evidence_receipt.json']],
            *[YIELD_EVIDENCE/name for name in ['text.py','figures.py','tables.py','published_response.py','published_response_receipt.json']],
            *[MAPS/name for name in ['prepare.py','figures.py','text.py','tables.py','map_data_receipt.json']],
            ROOT/'analysis/paper_study/nature_food_revision_20261007/epidemic_evaluation/evaluation_receipt.json',
            ROOT/'analysis/paper_study/nature_food_revision_20261007/yield_response/final_verification_receipt.json',
            ROOT/'calibration/yield_transfer.py',ROOT/'tests/test_published_yield_transfer.py',
            *[(STUDY/group/'receipt.json') for group in ('phenology','disease','climate')]]},
        native_docx_text_and_tables=True,pdf_uses_same_manuscript_body=True)
    (DEST/'document_build_receipt.json').write_text(json.dumps(receipt,indent=2)+'\n')
    print(json.dumps({key:value for key,value in receipt.items() if key not in ('source_workbook','sources')}))


if __name__=='__main__':main()
