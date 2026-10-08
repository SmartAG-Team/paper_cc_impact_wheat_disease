"""Full-season event replay and literature-based upper-canopy yield scenarios."""

from datetime import datetime,timezone
from pathlib import Path
from types import SimpleNamespace
import hashlib,json,sys,shutil,zipfile
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches,Pt
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Image,PageBreak
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT=Path(__file__).resolve().parents[3]
sys.path.insert(0,str(ROOT))
from analysis.paper_study.structural_evaluation.run import load_inputs,DEFAULT_PATHS
from model.seasonal_septoria.regional import tpv_accumulation
from calibration.seasonal_septoria.infection_priority import predict_infection_priority, DEFAULT_MODEL
from model.seasonal_septoria.yield_relevance import upper_leaf_had_loss,transfer_had_yield_loss
from model.seasonal_septoria.structural import canopy_host, simulate_canopy, CanopyParameters
from model.seasonal_septoria.wetness import duration_exposure

HERE=Path(__file__).parent
DEST=ROOT/'publication/infection_and_yield_priority_20261006'


def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def extend_forcing(data,weather):
    weather=weather.copy();weather['date']=pd.to_datetime(weather.date)
    groups={name:part.set_index('date').sort_index() for name,part in weather.groupby('location_id')}
    records=[];metadata=data.metadata.copy();lengths=[]
    for meta in metadata.itertuples():
        sow=pd.Timestamp(meta.sowing_date)
        end=max(pd.Timestamp(int(meta.season_year),9,30),pd.Timestamp(meta.last_forcing_date))
        group=groups[meta.site_id];end=min(end,group.index.max())
        dates=pd.date_range(sow,end)
        matched=group.reindex(dates)
        cols=['tmean_c','tmax_c','rh_mean_pct','precipitation_mm']
        missing=np.flatnonzero(~np.isfinite(matched[cols].to_numpy(float)).all(axis=1))
        if len(missing):
            matched=matched.iloc[:int(missing[0])]
            dates=dates[:int(missing[0])]
        if not len(dates) or dates[-1]<pd.Timestamp(meta.last_forcing_date):
            raise ValueError('A complete archived extended crop weather window is required.')
        records.append(matched[cols].to_numpy(float));lengths.append(len(dates))
    arrays=np.zeros((len(records),max(lengths),4))
    for i,values in enumerate(records):arrays[i,:len(values)]=values
    metadata['forcing_days']=lengths
    metadata['last_forcing_date']=[pd.Timestamp(s)+pd.Timedelta(days=n-1) for s,n in zip(metadata.sowing_date,lengths)]
    forcing=SimpleNamespace(temperature=arrays[:,:,0],maximum_temperature=arrays[:,:,1],
        humidity=arrays[:,:,2],rain=arrays[:,:,3],metadata=metadata)
    phenology=json.loads(Path(DEFAULT_PATHS['phenology']).read_text())
    accumulation=np.zeros_like(forcing.temperature)
    for meta in metadata.itertuples():
        i=int(meta.field_index);n=int(meta.forcing_days);dates=pd.date_range(meta.sowing_date,periods=n)
        accumulation[i,:n]=tpv_accumulation(forcing.temperature[i:i+1,:n],forcing.maximum_temperature[i:i+1,:n],
            [meta.latitude],dates.dayofyear,np.ones((1,n),bool),phenology)[0]
        accumulation[i,n:]=accumulation[i,n-1]
    return forcing,accumulation


def standardized_canopy_scenario(forcing,accumulation,thresholds,event_outputs,config):
    """Conditional per-unit LAI transfer using the retained gradual score states.

    This is an explicit score-to-functional-area scenario after predicted
    symptom detection, not measured green area or a locally validated yield.
    """
    path=ROOT/'configurations/wheat_stb/legacy/publication_model.json'
    progression=json.loads(path.read_text())
    valid=np.arange(forcing.temperature.shape[1])[None,:]<forcing.metadata.forcing_days.to_numpy()[:,None]
    functional=np.zeros((*forcing.temperature.shape,8))
    for component in progression['components']:
        fitted=component['fitted']
        active,renewal,area,_=canopy_host(accumulation,forcing.temperature,thresholds,**fitted['host_parameters'])
        active&=valid[:,:,None];renewal[~active]=0.
        preprocessing=fitted.get('weather_preprocessing',{})
        operator=fitted.get('weather_operator',preprocessing.get('weather_operator','daily_or'))
        exposure=None
        if operator=='duration_proxy':
            rate=fitted.get('rain_rate_mm_hour',preprocessing.get('rain_rate_mm_hour'))
            exposure=duration_exposure(forcing.temperature,forcing.maximum_temperature,
                forcing.humidity,forcing.rain,rate)['exposure']
        trajectory=simulate_canopy(np.where(valid,forcing.temperature,0.),forcing.humidity,forcing.rain,
            active,renewal,CanopyParameters(**fitted['parameters']),environmental_response=exposure)
        assert np.max(abs(trajectory.state.sum(axis=-1)-1))<1e-10
        functional+=component['weight']*trajectory.expressed[:,1:]
    _,_,reference_area,_=canopy_host(accumulation,forcing.temperature,thresholds,**config['event_model']['host_parameters'])
    reference_area=reference_area/3
    rows=[]
    for meta in forcing.metadata.itertuples():
        i=int(meta.field_index);n=int(meta.forcing_days);dates=pd.date_range(meta.sowing_date,periods=n)
        event=event_outputs.loc[event_outputs.field_index.eq(i)].iloc[0]
        for leaf in range(3):
            onset=pd.Timestamp(event[f'F{leaf+1}_symptom_date'])
            functional[i,:n,leaf]*=False if pd.isna(onset) else dates>=onset
        found=np.flatnonzero(accumulation[i,:n]>=thresholds[31])
        stem_date=pd.NaT if not len(found) else dates[int(found[0])]
        for scope,start in [('grain_fill_component',event.anthesis_proxy_date),('GS31_to_soft_dough_proxy',stem_date)]:
            had=upper_leaf_had_loss(dates,reference_area[i:i+1,:n],functional[i:i+1,:n],
                [start],[event.soft_dough_date])
            assert np.isnan(had['lost_had3'][0]) or 0<=had['lost_had3'][0]<=had['reference_had3'][0]+1e-10
            loss=transfer_had_yield_loss(had['lost_had3'],config['yield_slopes'])['absolute_loss_t_ha'][0]
            rows.append(dict(field_id=meta.field_id,window=scope,reference_upper3_lai=1.,
                reference_had3_per_unit_lai=had['reference_had3'][0],lost_had3_per_unit_lai=had['lost_had3'][0],
                conditional_loss_low_t_ha_per_unit_lai=loss[0],conditional_loss_central_t_ha_per_unit_lai=loss[1],
                conditional_loss_high_t_ha_per_unit_lai=loss[2],
                area_conversion='retained gradual score as functional fraction after predicted symptoms',
                evidence_role='standardized canopy scenario; not field yield validation'))
    return pd.DataFrame(rows),path


def main():
    if DEST.exists():raise FileExistsError('Use a new immutable priority package.')
    DEST.mkdir(parents=True);(DEST/'figures').mkdir();(DEST/'tables').mkdir()
    config=json.loads(DEFAULT_MODEL.read_text())
    paths=[DEFAULT_MODEL,*DEFAULT_PATHS.values(),ROOT/'configurations/wheat_stb/legacy/publication_model.json',HERE/'event_model/receipt.json',
        HERE/'anthesis_clock/receipt.json',HERE/'literature/yield_loss_primary_evidence.json',
        ROOT/'calibration/seasonal_septoria/infection_priority.py',ROOT/'model/seasonal_septoria/yield_relevance.py',Path(__file__)]
    identity={str(p.relative_to(ROOT)):sha(p) for p in paths}
    (DEST/'configuration_before_replay.json').write_text(json.dumps(dict(registered_utc=datetime.now(timezone.utc).isoformat(),
        experiment='full-season post-refit scenario replay, not outer validation',input_sha256=identity,
        disease_values_used_by_forecast=False,reference_yield_invented=False,
        observed_leaf_area_available=False,conditional_yield_scenario=True),indent=2)+'\n')
    data,weather,_,thresholds=load_inputs(DEFAULT_PATHS)
    forcing,a=extend_forcing(data,weather)
    result=predict_infection_priority(forcing,a,thresholds)
    outputs=result['field_outputs']
    outputs=outputs.merge(data.metadata[['field_id','dataset_id','site_id','season_year','latitude','longitude']],on='field_id',validate='one_to_one')
    outputs['evidence_role']='full-season scenario replay after event refitting'
    outputs.to_csv(DEST/'tables/full_season_top3_timing.csv',index=False)
    complete=outputs.loc[outputs.complete_grain_fill_window]
    standardized,progression_source=standardized_canopy_scenario(forcing,a,thresholds,outputs,config)
    standardized.to_csv(DEST/'tables/standardized_upper3_yield_transfer.csv',index=False)
    timing=[]
    for leaf in range(1,4):
        date=pd.to_datetime(complete[f'F{leaf}_infection_date'])
        relative=(date-complete.anthesis_proxy_date).dt.days
        timing.append(dict(leaf=f'F{leaf}',complete_fields=len(complete),events=int(date.notna().sum()),
            median_infection_days_relative_to_anthesis=float(relative.median()),
            mean_infected_fraction_of_grain_fill=float(complete[f'F{leaf}_infected_grain_fill_fraction'].mean())))
    pd.DataFrame(timing).to_csv(DEST/'tables/upper_leaf_grain_fill_summary.csv',index=False)
    for name in ['assessment_detection_metrics.csv','onset_bracket_metrics.csv','observed_window_occurrence_metrics.csv']:
        shutil.copyfile(HERE/'event_model'/name,DEST/'tables'/name)
    source=json.loads((HERE/'literature/yield_loss_primary_evidence.json').read_text())
    scenarios=[];curves=[];dates=pd.date_range('2020-06-01',periods=81)
    day=np.arange(81)-10;area=np.ones((1,81,3))/3
    for onset in [-10,0,15,30,50]:
        fraction=np.minimum(.8,np.maximum((day-onset)*.04,0))
        damage=np.broadcast_to(fraction[None,:,None],area.shape).copy()
        had=upper_leaf_had_loss(dates,area,damage,[pd.Timestamp('2020-06-11')],[pd.Timestamp('2020-07-25')])
        impact=transfer_had_yield_loss(had['lost_had3'],config['yield_slopes'])
        scenarios.append(dict(symptom_onset_relative_to_anthesis_days=onset,
            reference_upper3_lai=1.,functional_loss_at_final_day=float(fraction[-1]),
            lost_had3_per_unit_reference_lai=float(had['lost_had3'][0]),
            conditional_loss_low_t_ha=float(impact['absolute_loss_t_ha'][0,0]),
            conditional_loss_central_t_ha=float(impact['absolute_loss_t_ha'][0,1]),
            conditional_loss_high_t_ha=float(impact['absolute_loss_t_ha'][0,2]),
            evidence_role='specified canopy/damage illustration; not observed field yield'))
        curves.append(fraction)
    examples=pd.DataFrame(scenarios);examples.to_csv(DEST/'tables/literature_yield_timing_scenarios.csv',index=False)
    assert examples.functional_loss_at_final_day.nunique()==1
    assert examples.conditional_loss_central_t_ha.iloc[-1]==0
    assert examples.conditional_loss_central_t_ha.is_monotonic_decreasing
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'pdf.fonttype':42,
        'axes.spines.top':False,'axes.spines.right':False})
    fig,axes=plt.subplots(1,2,figsize=(8,3.3),layout='constrained')
    for onset,curve in zip(examples.symptom_onset_relative_to_anthesis_days,curves):
        axes[0].plot(day,curve,label=f'{onset:+d}d')
    axes[0].axvspan(0,44,color='#D6E7DE',alpha=.5)
    axes[0].set(xlabel='Days relative to anthesis',ylabel='Specified functional leaf-area loss',
        title='a  Timing relative to grain filling',xlim=(-10,70))
    axes[0].legend(frameon=False,ncol=3,fontsize=8)
    axes[1].errorbar(examples.symptom_onset_relative_to_anthesis_days,examples.conditional_loss_central_t_ha,
        yerr=np.vstack([examples.conditional_loss_central_t_ha-examples.conditional_loss_low_t_ha,
        examples.conditional_loss_high_t_ha-examples.conditional_loss_central_t_ha]),fmt='o-',color='#14696E',capsize=3)
    axes[1].set(xlabel='Symptom onset relative to anthesis (days)',ylabel='Conditional loss per unit upper3 LAI (t/ha)',
        title='b  Transferred top-three HAD response')
    fig.savefig(DEST/'figures/yield_timing.pdf',bbox_inches='tight');fig.savefig(DEST/'figures/yield_timing.png',dpi=300,bbox_inches='tight');plt.close(fig)
    fig,axis=plt.subplots(figsize=(7.3,3.4),layout='constrained')
    relative_all=[]
    for leaf in range(1,4):
        relative_all.extend((pd.to_datetime(complete[f'F{leaf}_infection_date'])-complete.anthesis_proxy_date).dt.days.dropna().tolist())
    bins=np.arange(5*np.floor(min(relative_all)/5),max(5,5*np.ceil(max(relative_all)/5)+5)+1,5)
    for leaf in range(1,4):
        relative=(pd.to_datetime(complete[f'F{leaf}_infection_date'])-complete.anthesis_proxy_date).dt.days.dropna()
        assert np.histogram(relative,bins=bins)[0].sum()==len(relative)
        axis.hist(relative,bins=bins,histtype='step',lw=1.5,label=f'F{leaf}, n={len(relative)}')
    axis.axvline(0,ls='--',color='gray',lw=1)
    axis.set(xlabel='Model-inferred infection date relative to anthesis proxy (days)',ylabel='Field-seasons',title='Full-season replay: upper-leaf timing')
    axis.legend(frameon=False)
    fig.savefig(DEST/'figures/top3_timing_replay.pdf',bbox_inches='tight');fig.savefig(DEST/'figures/top3_timing_replay.png',dpi=300,bbox_inches='tight');plt.close(fig)
    sections=[
        ('title','Wheat infection timing and yield relevance of the top three leaves'),
        ('p','The primary targets are whether effective infection conditions occur, their timing on F1–F3, and the resulting exposure of grain filling to upper-leaf disease. Visible symptoms and model-inferred infection dates remain separate. Event calibration uses zero versus positive detections and assessment brackets, without fitting the numerical magnitudes of infection percentages.'),
        ('h','Infection and symptom events'),
        ('p','The model requires an available leaf and cumulative temperature-weighted humid/rain exposure during a wet spell. A declared dry gap resets insufficient dose. This permits no-event predictions under dry conditions, rather than accumulating a small positive infection pressure indefinitely. Thermal development delays visible symptoms after the inferred infection event. The selected post-evaluation configuration uses flag timing fraction 0.3, leaf interval 120, exposure dose 0.5, dry-gap length three days and delay 20 reference days at 18°C. These are effective field-model parameters conditional on inoculum presence, not laboratory infection requirements or calibrated occurrence probabilities.'),
        ('p','Seventy-two event candidates and twelve independently selected phenology-only candidates retain the five spatial and four forward outer partitions, with three inner location folds. Complete source/location histories share a fold. The primary selection metric is top-three two-sided bracket distance; right and left censoring remain explicit. First-onset interpretation assumes no unobserved earlier transient crossing. Five subsequent positive-to-zero assessments violate persistent symptom detection and remain in scoring. Numerical infection values do not enter event fitting, and changing positive magnitudes leaves signs and forecasts unchanged.'),
        ('p','Retrospective spatial top-three assessment sensitivity is 93.23% and specificity 73.97%; forward sensitivity is 96.53% and specificity 68.96%. Two-sided bracket distance is 2.43 days spatially and 3.52 days forward, with compatibility 64.65% and 57.53%. These distances are deviations outside observed brackets, rather than exact infection-date errors. Independently selected phenology-only distances are 2.37 and 3.61 days. Weather exposure provides little additional timing discrimination in these data. Only three of 217 top-three field windows are zero-only, so whole-season infection-free discrimination remains weakly tested. Actual infection dates are not directly observed.'),
        ('h','Grain-fill timing and upper-leaf exposure'),
        ('p','BBCH65 supplies an anthesis proxy and BBCH85 supplies the soft-dough boundary of the principal grain-fill window. The flowering threshold is calibrated exclusively on original BASF 2017–2018 stage constraints, with the other TPV parameters fixed. The selected C65 is 1165.93 development units; minimum-loss grid values span 1152.46–1177.14 and the within-one-day objective range spans 1069.44–1264.65. Those are finite-grid loss-tolerance ranges, not confidence bounds. Unreached or incomplete windows remain missing.'),
        ('p',f'Extending the archived field forcing through the available harvest-year September window yields complete anthesis-to-soft-dough proxies for {len(complete)} of {len(outputs)} modeled field-seasons. Leaf-specific infection and symptom dates, infected grain-fill days and the overlap of all three infected leaves are retained. This is a full-season scenario replay after event refitting, separate from out-of-fold validation. Binary infection overlap describes yield-relevant exposure and does not represent the fraction of photosynthetic area lost.'),
        ('p','All three modeled leaf infections precede anthesis in this replay, so binary grain-fill exposure is saturated. Median F1/F2/F3 infection dates are 32/44/61 days before the flowering proxy; F1 symptoms begin within grain filling in 15 fields and beforehand in 203. Functional area trajectories provide the additional magnitude information needed for yield effects; infection presence alone is insufficient.'),
        ('h','Literature-based yield linkage'),
        ('p','Functional green-area deficit in the final three leaves is integrated as ΔHAD3 = Σt Σi=1..3 LAI_reference,i(t) × functional_loss_i(t) × Δt. Leaf area is expressed as m² green lamina per m² ground; the resulting deficit has GLAI-day units. Actual layer areas determine contributions, without a universally prescribed F1/F2/F3 percentage weighting. Reference area and functional damage are separate inputs; infection presence is not assigned complete area loss.'),
        ('p','Parker et al. (2004) related wheat yield to the absolute green area duration of the top three leaves. Table5 supports a transferred centre b = 0.0180 t ha⁻¹ per GLAI-day, with cultivar slope predictions 0.0141–0.0207. The conditional difference operator is ΔY = b × ΔHAD3. Relative loss additionally requires an independently supplied reference yield Y0. The cultivar range is sensitivity, not a global confidence interval. Their approximately ten-day assessments began at GS31 and used trapezoidal integration; the present daily rectangle integral and anthesis-to-BBCH85 restriction are adaptations. Equivalence with the original whole observed window requires zero disease-induced green-area loss outside the restricted window.'),
        ('p','Foulkes et al. (2006) support post-anthesis green-duration effects but used the top five leaves and STB/stripe-rust experiments; their coefficients remain separate. Bancal et al. (2015) found that a simple green-duration/yield slope did not transfer uniformly across a broad French genotype/environment dataset. The quantitative relation is therefore a literature-informed conditional transfer, without local matched-yield validation.'),
        ('p',f'The specified timing illustration uses total reference upper-three LAI = 1, equal area allocation solely for the illustration, a common functional-loss progression and a 45-day grain-fill window. All scenarios reach the same final functional loss. A symptom start ten days before anthesis gives {examples.conditional_loss_central_t_ha.iloc[0]:.3f} t ha⁻¹ conditional loss per unit reference LAI; a start 30 days after anthesis gives {examples.conditional_loss_central_t_ha.iloc[3]:.3f}, and onset after the window gives zero contribution within that window. These are scenario values, not measured yields. Numeric field yield loss is undefined until reference leaf area and functional loss are supplied.'),
        ('p','A standardized canopy scenario retains the earlier gradual score trajectories, previously calibrated to numerical severity magnitudes, as a conditional functional-area proxy. Loss before the newly predicted symptom event is suppressed. Reference upper-three area is normalized to one, with equal mature layer areas and the declared unfolding capacities; natural green-area senescence is not measured. Both GS31-to-soft-dough and anthesis-to-soft-dough deficit integrals are provided, distinguishing the broader upper-leaf proxy from the grain-fill component. Transferred tonnes per hectare per unit reference LAI are scenario outputs. Their field interpretation depends on the explicit score-to-functional-area conversion, actual canopy area and source-window assumptions.'),
        ('h','Runtime and interpretation'),
        ('p','The infection-priority runtime returns infection and symptom dates, a top-three grain-fill exposure table, and optional conditional yield transfer. It never reads target disease values. Field-specific forcing horizons prevent padding from advancing either infection or symptom clocks. Unknown crop windows produce missing values, and linear yield transfers exceeding the independent reference yield are flagged without silently clipping them. The existing severity and continental climate archives retain their model identities; new event/yield climate simulations form a separate analysis.'),
        ('h','References'),
        ('p',source['primary_source']['citation']+' https://doi.org/'+source['primary_source']['doi']),
        ('p',source['independent_evidence'][0]['citation']+' https://doi.org/'+source['independent_evidence'][0]['doi']),
        ('p',source['independent_evidence'][1]['citation']+' https://doi.org/'+source['independent_evidence'][1]['doi']),
    ]
    text='\n\n'.join(item[1] for item in sections)+'\n'
    (DEST/'Infection_and_Upper3_Yield.txt').write_text(text)
    doc=Document();doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10)
    pdfmetrics.registerFont(TTFont('PrioritySans',str(Path(matplotlib.get_data_path())/'fonts/ttf/DejaVuSans.ttf')))
    normal=ParagraphStyle('body',fontName='PrioritySans',fontSize=9.5,leading=13.5,spaceAfter=8)
    header=ParagraphStyle('head',fontName='PrioritySans',fontSize=12,leading=15,spaceAfter=8,spaceBefore=8)
    story=[]
    from xml.sax.saxutils import escape
    for kind,content in sections:
        if kind=='p':doc.add_paragraph(content)
        else:doc.add_heading(content,0 if kind=='title' else 1)
        story.append(Paragraph(escape(content),normal if kind=='p' else header))
    for name,caption in [('yield_timing','Figure1. Literature-transferred conditional yield sensitivity to timing. Green shading denotes the specified 45-day grain-fill window. Bars show cultivar slope sensitivity and are not prediction intervals. Total upper-three LAI and the damage curve are explicitly specified illustration inputs.'),
        ('top3_timing_replay','Figure2. Model-inferred upper-leaf infection timing relative to the flowering proxy, in complete full-season historical replay windows. Counts are unweighted field-seasons. These post-refit dates are scenario outputs rather than out-of-fold infection-date validation.')]:
        doc.add_page_break();doc.add_picture(str(DEST/'figures'/f'{name}.png'),width=Inches(6.8));doc.add_paragraph(caption)
        pixels=plt.imread(DEST/'figures'/f'{name}.png')
        story.extend([PageBreak(),Image(str(DEST/'figures'/f'{name}.png'),width=500,height=500*pixels.shape[0]/pixels.shape[1]),Spacer(1,8),Paragraph(escape(caption),normal)])
    doc.save(DEST/'Infection_and_Upper3_Yield.docx')
    SimpleDocTemplate(str(DEST/'Infection_and_Upper3_Yield.pdf'),pagesize=(612,792),leftMargin=56,rightMargin=56,topMargin=42,bottomMargin=42).build(story)
    shutil.copyfile(DEFAULT_MODEL,DEST/'infection_priority_model.json')
    shutil.copyfile(HERE/'literature/yield_loss_primary_evidence.json',DEST/'literature_evidence.json')
    readme='''Infection-first runtime

from calibration.seasonal_septoria.infection_priority import predict_infection_priority
result = predict_infection_priority(data, accumulation, thresholds)
events = result['field_outputs']

data supplies mean/max temperature, mean RH, precipitation and per-field metadata (sowing_date, forcing_days, field_index). The declared TPV index and thresholds must match the crop scenario. Disease magnitudes and labels are not forecast inputs. No-event dates are NaT; incomplete grain-fill windows remain missing.

Numerical conditional yield transfer requires both reference_leaf_area_index (m2 leaf/m2 ground) and functional_loss_fraction (0..1), matching field/day/layer shapes. reference_yield_t_ha is additionally required for relative loss. Binary infection is not functional loss. Coefficient range is cultivar sensitivity, not CI. The main anthesis-to85 window is a conditional adaptation of the published top-three HAD relation.

The event model was selected with zero/positive signs and top3 symptom brackets, rather than percentage RMSE. Detection metrics concern observed dates and do not establish infection-free season specificity. Full-season outputs in this package are post-refit scenarios; untouched data and directly observed infection dates are absent. Original severity/climate outputs remain separate.

The standardized_upper3_yield_transfer.csv table is a conditional per-unit canopy replay: retained score trajectories are converted to functional area after modeled symptom onset. It is not a table of measured or locally calibrated yields. Shared model/phenology code is packaged; raw weather/trial inputs and their original archived sources remain in the existing repository.
'''
    (DEST/'Runtime.txt').write_text(readme)
    sources=[ROOT/'model/seasonal_septoria/infection_events.py',ROOT/'calibration/seasonal_septoria/infection_priority.py',
        ROOT/'model/seasonal_septoria/yield_relevance.py',DEFAULT_MODEL,ROOT/'calibration/seasonal_septoria/docs/INFECTION_PRIORITY_PROTOCOL.txt',
        ROOT/'tests/test_infection_events.py',ROOT/'tests/test_infection_priority.py',ROOT/'tests/test_yield_relevance.py',Path(__file__),
        HERE/'event_model/run.py',HERE/'event_model/receipt.json',HERE/'event_model/verify_outputs.py',
        HERE/'anthesis_clock/threshold_selection.json',HERE/'anthesis_clock/receipt.json',progression_source]
    sources=[p for p in sources if p.exists()]
    sources.extend(sorted((ROOT/'model/seasonal_septoria').glob('*.py')))
    sources.extend(sorted((ROOT/'calibration').rglob('*.py')))
    sources.extend(sorted((ROOT/'calibration').rglob('*.txt')))
    sources.extend(sorted((ROOT/'configurations/wheat_stb/legacy').glob('*.json')))
    sources.extend(sorted((ROOT/'process_model').glob('*.py')))
    sources.extend(sorted((ROOT/'process_model/_support').glob('*.py')))
    sources.extend(sorted((HERE/'event_model').glob('*.py')))
    sources.extend(sorted((ROOT/'analysis/paper_study/structural_evaluation').glob('*.py')))
    sources.extend(ROOT/name for name in ['analysis/paper_study/calibrate_seasonal.py',
        'analysis/paper_study/refine_seasonal_accuracy.py','analysis/paper_study/run_empirical_benchmarks.py',
        'analysis/paper_study/run_structural_canopy.py','process_model/parameters/calibration.json'])
    sources=sorted(set(sources))
    with zipfile.ZipFile(DEST/'Model_and_Evidence.zip','w',zipfile.ZIP_DEFLATED) as archive:
        for p in sources:archive.write(p,str(p.relative_to(ROOT)))
        for p in (DEST/'tables').glob('*'):archive.write(p,'tables/'+p.name)
        for p in [DEST/'Runtime.txt',DEST/'infection_priority_model.json',DEST/'literature_evidence.json']:archive.write(p,p.name)
    assert result['yield_transfer'] is None
    for p,digest in identity.items():assert sha(ROOT/p)==digest
    receipt=dict(status='complete',model_id=result['model_id'],fields=len(outputs),complete_grain_fill_windows=len(complete),
        infection_percentage_fit=False,local_yield_calibration=False,full_replay_is_validation=False,
        secondary_progression_uses_legacy_score_calibration=True,standardized_yield_uses_conditional_area_conversion=True,
        actual_field_yield_predictions_invented=False,scenario_checks_passed=True,
        source_sha256={str(p.relative_to(ROOT)):sha(p) for p in sources},
        output_sha256={str(p.relative_to(DEST)):sha(p) for p in DEST.rglob('*') if p.is_file()})
    (DEST/'receipt.json').write_text(json.dumps(receipt,indent=2)+'\n')
    print(json.dumps(dict(status='complete',fields=len(outputs),complete_windows=len(complete),yield_scenarios=scenarios)))


if __name__=='__main__':main()
