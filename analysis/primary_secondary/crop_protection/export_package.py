"""Append frozen forecasts and new source cells without rewriting observations."""
from datetime import datetime,timezone
from pathlib import Path
import hashlib
import json
import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'data/harmonized'
SCHEMA=pq.read_schema(OUT/'observations.parquet')
RAW_SCHEMA=pq.read_schema(OUT/'raw_table_cells.parquet')


def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()


def scalar(value):
    if value is None or pd.isna(value):return None
    if isinstance(value,datetime):return value.isoformat()
    return str(value)


def predictions():
    parts=[];keys=[]
    inputs=[('primary_secondary_v2_basf_evaluation',ROOT/'analysis/primary_secondary/calibration_v2_age/predictions.csv'),
        ('primary_secondary_v2_french_transfer',ROOT/'analysis/primary_secondary/external_french/transfer_v2/frozen_transfer_predictions.csv')]
    locations=pd.read_csv(ROOT/'data/era5/location_registry.csv').set_index('location_id')
    quantities={'predicted_percent':'predicted_recorded_infection_percent_proxy',
        'unresolved_initial_infected_percent':'unresolved_initial_infected_tissue_percent',
        'external_infected_percent':'external_origin_infected_tissue_percent',
        'secondary_infected_percent':'secondary_origin_infected_tissue_percent',
        'infectious_percent':'modelled_infectious_tissue_percent',
        'pycnidia_proxy_percent':'modelled_pycnidial_area_percent_proxy'}
    for dataset,source in inputs:
        f=pd.read_csv(source)
        if 'fold' not in f:f['fold']='external_'+f.year.astype(str)
        source_rows=pd.Series(f.index+2,index=f.index).astype(str)
        record=f.series_id+'|'+f.start+'|'+f.Date
        base='primary_secondary_v2|'+dataset+'|'+f.fold+'|'+f.model+'|'+record
        french='french' in dataset
        for column,metric in quantities.items():
            if column not in f:continue
            chosen=f.loc[f[column].notna()].copy()
            if chosen.empty:continue
            x=pd.DataFrame(index=chosen.index,columns=SCHEMA.names)
            x['observation_id']=base.loc[chosen.index]+'|'+column
            x['dataset_id']=dataset;x['source_file']=str(source.relative_to(ROOT));x['source_table']='predictions'
            x['source_row']=source_rows.loc[chosen.index];x['source_column']=column
            x['record_id']=record.loc[chosen.index];x['plot_id']=chosen.TrialId.astype(str)
            x['site_id']=chosen.location_id;x['country']=chosen.Country
            x['country']=x.country.str.upper().replace({'SLOWAKIA':'SLOVAKIA','CZECH REPUBLIC':'CZECHIA'})
            x['latitude']=chosen.location_id.map(locations.requested_latitude)
            x['longitude']=chosen.location_id.map(locations.requested_longitude)
            x['season_year']=chosen.year;x['date']=chosen.Date
            x['time_basis']='conditional_hindcast_calendar';x['pathogen']='Zymoseptoria tritici'
            x['organ']='numbered_leaf_rank_'+chosen.leaf_rank.astype(str)
            if french:
                x['cultivar']=chosen.var_origin;x['replicate']=chosen.rep
            x['treatment']=chosen['fold']+'|'+chosen.model
            x['metric']='predicted_pycnidial_area_percent_proxy' if french else metric
            x['unit']='percent';x['value']=chosen[column];x['raw_value']=chosen[column].map(scalar)
            x['measurement_role']='conditional_model_prediction' if column=='predicted_percent' else 'modelled_compartment_state'
            x['quality_status']='external_single_station_transfer' if french else 'development_era_grouped_transfer'
            parts.append(x)
        key=pd.DataFrame({'prediction_key':base,'dataset_id':dataset,'source_file':str(source.relative_to(ROOT)),
            'source_row':source_rows,'fold':f.fold,'model':f.model,'series_id':f.series_id,
            'issue_date':f.start,'assessment_date':f.Date,'forecast_horizon_days':f.day,
            'endpoint':'pycnidial_area_percent' if french else 'recorded_percent_infection',
            'weather_information':'realized post-issue ERA5',
            'source_treatment':f.mixture if french else f.Treatment})
        keys.append(key)
    combined=pd.concat(parts,ignore_index=True)
    for field in SCHEMA:
        combined[field.name]=pd.to_numeric(combined[field.name],errors='coerce') if pa.types.is_floating(field.type) else combined[field.name].map(scalar)
    if not combined.observation_id.is_unique:raise ValueError('Duplicate prediction keys.')
    pq.write_table(pa.Table.from_pandas(combined[SCHEMA.names],schema=SCHEMA,preserve_index=False),
        OUT/'primary_secondary_predictions.parquet',compression='zstd')
    pd.concat(keys,ignore_index=True).to_csv(OUT/'primary_secondary_prediction_key_dictionary.csv',index=False)
    return len(combined),{str(p.relative_to(ROOT)):digest(p) for _,p in inputs}


def raw_source_cells():
    sources=sorted((ROOT/'data/public_septoria/primary_evidence').rglob('*.xlsx'))
    target=OUT/'raw_primary_process_table_cells.parquet'
    rows=0;inventory=[]
    with pq.ParquetWriter(target,RAW_SCHEMA,compression='zstd') as writer:
        for source in sources:
            workbook=load_workbook(source,read_only=True,data_only=False)
            dataset='boixel-moisture-process' if 'boixel' in str(source) else 'orellana-2022-interseason'
            for sheet in workbook:
                records=[]
                for row_number,row in enumerate(sheet.iter_rows(values_only=True),1):
                    for column_number,value in enumerate(row,1):
                        records.append({'dataset_id':dataset,'source_file':str(source.relative_to(ROOT)),
                            'source_table':sheet.title,'source_row':str(row_number),
                            'source_column':str(column_number),'raw_value':scalar(value)})
                    if len(records)>=100000:
                        writer.write_table(pa.Table.from_pylist(records,schema=RAW_SCHEMA));rows+=len(records);records=[]
                if records:writer.write_table(pa.Table.from_pylist(records,schema=RAW_SCHEMA));rows+=len(records)
                inventory.append({'source_file':str(source.relative_to(ROOT)),'sheet':sheet.title,
                    'source_rows':sheet.max_row,'source_columns':sheet.max_column,
                    'rectangular_cells':sheet.max_row*sheet.max_column})
            workbook.close()
    pd.DataFrame(inventory).to_csv(OUT/'raw_primary_process_table_inventory.csv',index=False)
    if rows!=sum(i['rectangular_cells'] for i in inventory):raise ValueError('Source rectangle counts differ.')
    return rows,{str(p.relative_to(ROOT)):digest(p) for p in sources}


def main():
    protected=[OUT/'observations.parquet',OUT/'raw_table_cells.parquet']
    before={str(p.relative_to(ROOT)):digest(p) for p in protected}
    prediction_count,prediction_sources=predictions()
    raw_count,workbooks=raw_source_cells()
    manifest_path=OUT/'manifest.json';manifest=json.loads(manifest_path.read_text())
    name='primary_secondary_predictions.parquet'
    if name not in manifest['common_schema_partitions']:manifest['common_schema_partitions'].append(name)
    definitions=[]
    for name in manifest['common_schema_partitions']:
        if not pq.read_schema(OUT/name).equals(SCHEMA):raise ValueError('Common schema drift.')
        for batch in pq.ParquetFile(OUT/name).iter_batches(batch_size=150000,
                columns=['dataset_id','metric','unit','measurement_role']):
            definitions.append(batch.to_pandas().drop_duplicates())
    pd.concat(definitions).drop_duplicates().sort_values(['dataset_id','metric','unit']).to_csv(
        OUT/'common_metric_dictionary.csv',index=False)
    manifest['common_records']=sum(pq.ParquetFile(OUT/p).metadata.num_rows for p in manifest['common_schema_partitions'])
    manifest['raw_cell_partitions']=['raw_table_cells.parquet','raw_primary_process_table_cells.parquet']
    manifest['raw_cell_records']=sum(pq.ParquetFile(OUT/p).metadata.num_rows for p in manifest['raw_cell_partitions'])
    manifest['updated_utc']=datetime.now(timezone.utc).isoformat()
    manifest['forecast_source_sha256']=prediction_sources
    readme=(OUT/'README.txt').read_text().replace('All39 downloaded public files','The original39 downloaded public files')
    old_raw='raw_table_cells.parquet preserves every cell of all collected spreadsheet/CSV tables, including headers, alternate derived views and model-fit candidates, with original table/row/column keys.'
    old_raw_revision='raw_table_cells.parquet preserves the initial spreadsheet/CSV archive; raw_primary_process_table_cells.parquet preserves all used rectangular cells of the four additional original workbooks. Both retain including headers, alternate derived views and model-fit candidates, with original table/row/column keys.'
    new_raw='raw_table_cells.parquet preserves the initial spreadsheet/CSV archive; raw_primary_process_table_cells.parquet preserves all used rectangular cells of the four additional original workbooks. Both retain original table, row and column keys for headers, measurements and source-derived views.'
    readme=readme.replace(old_raw,new_raw).replace(old_raw_revision,new_raw)
    readme=readme.replace('metric_dictionary.csv defines source study, metric, unit and source column.',
        'metric_dictionary.csv retains source definitions; common_metric_dictionary.csv defines metrics, units and roles across every common partition.')
    paragraph='Primary–secondary v2 forecasts are in primary_secondary_predictions.parquet. Forecasts and source-labelled compartment states are derived model outputs, with no additional field observations. Source-specific endpoints, issue dates, lead times, fold/model identifiers and realized-weather context are retained in primary_secondary_prediction_key_dictionary.csv. The French transfer comprises one independent station across two seasons. Diagnostic severity cutoffs are not validated economic injury or fungicide thresholds.'
    if paragraph not in readme:readme+='\n'+paragraph+'\n'
    origin_definition='Origin-state percentages include latent and visible infection and are not visible-disease source contributions or measured spore-type attribution.'
    if origin_definition not in readme:readme+='\n'+origin_definition+'\n'
    (OUT/'README.txt').write_text(readme)
    manifest['files']=[dict(path=str(p.relative_to(ROOT)),bytes=p.stat().st_size,sha256=digest(p),
        **({'rows':pq.ParquetFile(p).metadata.num_rows} if p.suffix=='.parquet' else {}))
        for p in sorted(OUT.iterdir()) if p.is_file() and p.name!='manifest.json']
    manifest_path.write_text(json.dumps(manifest,indent=2)+'\n')
    receipt={'created_utc':datetime.now(timezone.utc).isoformat(),'prediction_cells':prediction_count,
        'new_raw_workbook_cells':raw_count,'source_workbook_sha256':workbooks,
        'prediction_source_sha256':prediction_sources,'protected_partitions_sha256':before,
        'protected_partitions_unchanged':all(digest(ROOT/p)==h for p,h in before.items()),
        'common_partitions':len(manifest['common_schema_partitions']),'common_cells':manifest['common_records'],
        'raw_cell_partitions':manifest['raw_cell_partitions'],'raw_cells':manifest['raw_cell_records'],
        'prediction_cell_count_is_not_field_sample_count':True}
    (Path(__file__).parent/'package_export_receipt.json').write_text(json.dumps(receipt,indent=2)+'\n')
    print(json.dumps(receipt,indent=2))


if __name__=='__main__':main()
