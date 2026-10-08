"""Create formal crop-protection Methods and Results from verified outputs."""
from datetime import datetime,timezone
import hashlib
import json
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt

ROOT=Path(__file__).resolve().parents[3]
ANALYSIS=Path(__file__).resolve().parent
OUT=ROOT/'publication'
OUT.mkdir(exist_ok=True)


def main():
    severity=pd.read_csv(ANALYSIS/'severity_metrics.csv')
    diagnostic=pd.read_csv(ANALYSIS/'threshold_diagnostics.csv')
    comparisons=pd.read_csv(ANALYSIS/'upper_leaf_paired_comparisons.csv')
    episodes=pd.read_csv(ROOT/'analysis/primary_secondary/calibration_v2_age/episodes.csv')
    fit=json.loads((ROOT/'analysis/primary_secondary/calibration_v2_age/all_development_fit.json').read_text())
    def metric(fold,model,leaf='upper_three'):
        r=severity.loc[severity.fold.eq(fold)&severity.model.eq(model)&severity.leaf_scope.eq(leaf)
            &severity.horizon_scope.eq('all')&severity.stage_scope.eq('all')]
        if len(r)!=1:raise ValueError('Ambiguous publication metric.')
        return r.iloc[0]
    def risk(fold,model):
        r=diagnostic.loc[diagnostic.fold.eq(fold)&diagnostic.model.eq(model)&diagnostic.leaf_scope.eq('upper_three')
            &diagnostic.horizon_scope.eq('all')&diagnostic.stage_scope.eq('all')&diagnostic.cutoff_percent.eq(10)
            &diagnostic.eligibility.eq('initial_below_cutoff')]
        if len(r)!=1:raise ValueError('Ambiguous publication diagnostic.')
        return r.iloc[0]
    def difference(fold):
        r=comparisons.loc[comparisons.fold.eq(fold)&comparisons.baseline.eq('global_logit_trend')].iloc[0]
        return f"{r.rmse_improvement_pp:+.2f} pp (95% location-block interval {r.ci_lower_pp:+.2f} to {r.ci_upper_pp:+.2f} pp)"
    m18=metric('forward_2018','primary_secondary_hidden');m19=metric('forward_2019','primary_secondary_hidden')
    l18=metric('forward_2018','global_logit_trend');l19=metric('forward_2019','global_logit_trend')
    d18=risk('forward_2018','primary_secondary_hidden');d19=risk('forward_2019','primary_secondary_hidden')
    b19=risk('forward_2019','global_logit_trend')
    f18=metric('external_2018','primary_secondary_hidden_BASF_frozen')
    f19=metric('external_2019','primary_secondary_hidden_BASF_frozen')
    sections=[
    ('Crop protection objective',
    'Septoria tritici blotch forecasting concerns the timing and magnitude of disease risk on yield-forming wheat leaves. Leaf emergence, crop development, current disease pressure and weather define the exposure of the upper canopy. The flag leaf and leaf 2 have particular relevance to yield formation, while leaf 3 provides an earlier crop-protection target [1,2]. Primary establishment and subsequent within-canopy spread are components of the epidemic forecast. The principal research outcome is reliable risk prediction across wheat fields and seasons, with sufficient lead time to support scouting and protection decisions. Quantitative attribution to individual spore types is a separate inference requiring source-specific exposure measurements.'),
    ('Field data and observation endpoints',
    f'The BASF modelling cohort contains 46 untreated wheat trials, {episodes.location_id.nunique()} rounded coordinates and {episodes.coordinate_year.nunique()} coordinate-years during 2017–2019. Eligibility requires an explicit numbered leaf, a numeric SEPTTR percentage between 0 and 100, at least three assessment dates per trial–leaf sequence, and no CROP INJURY annotation. The cohort contains 91 sequences and 315 assessments. Each first assessment conditions its forecast, leaving 224 later assessments for scoring. Sixty-six sequences are already positive at initialization. The source P%INF code retains an incompletely specified percentage-infection denominator; the model observation mapping is therefore a recorded-disease proxy rather than a measured tissue compartment [3].\n\n'
    'The independent French source contains adult-leaf pycnidial-area observations on flag leaves, leaf 2 and leaf 3. Means use numerically assessed plants within each plot, cultivar, date and rank. The S sentinel identifies senescence and is excluded from the numeric denominator, with its count retained. Eligibility yields 144 sequences, 527 numeric plot–cultivar–leaf means and 383 subsequent targets. Juvenile leaf numbers are excluded from adult-rank canopy transport. These records represent one station in two seasons, with three experimental blocks per season; plot sequences do not constitute independent stations [4].'),
    ('Conditional crop epidemic model',
    'The leaf-resolved model contains susceptible tissue, three serial latent compartments, visible nonsporulating disease, infectious tissue and removed tissue. Initial unresolved infection, externally acquired infection and secondary within-canopy infection retain separate model labels. The external acquisition hazard is α times a temperature–humidity establishment function. The secondary hazard is β times that establishment function, a rainfall response and the rank-weighted infectious canopy. The canopy kernel declines exponentially with rank distance and is normalized over ranks observed by forecast initialization. No future disease assessment initializes another leaf.\n\n'
    'The working establishment response is exp[−0.5((T−18)/8)²]H/24, where T is daily mean temperature in °C and H is the daily number of hours with relative humidity ≥90%. The rainfall response is 1−exp(−R/2), with R in mm. Latent and nonsporulating progression are scaled by max(T,0)/18. The principal variant fixes the latent mean at 20 days at 18°C, the nonsporulating duration at 3 days and the infectious duration at 21 days. These working response functions and durations are structural assumptions, with comparative process evidence retained separately.\n\n'
    'A fitted initial latent proportion applies to tissue without visible disease and is distributed equally among the three latent stages. Its source remains unresolved. Older visible observations on other leaves retain their observed damage while their infectious share decays exponentially over elapsed time. This decay-only initialization does not reconstruct unobserved infections during the intervening period. For BASF, the observation proxy sums visible nonsporulating, infectious and removed tissue. For French pycnidial area, the proxy sums infectious and removed tissue. Latent tissue is excluded from both observable quantities.'),
    ('Weather, crop stage and forecast information',
    'Daily forcing derives from the archived hourly ERA5 series supplied through Open-Meteo with models=era5, UTC, nearest-grid selection and no elevation correction [5]. Each forecast uses complete daily weather from its initialization date through the day before the target assessment. Relative-humidity exposure is a weather proxy rather than observed canopy leaf wetness. Weather after initialization is realized reanalysis; the evaluation measures retrospective conditional prediction and does not establish accuracy with weather forecasts available at issue time.\n\n'
    'Growth-stage strata use only the BBCH stage recorded at the first assessment: below GS31, GS31–33, GS34–39, GS40–59 and GS60 or later, with missing stages retained as unknown. Subsequent observed stages are excluded from issue-time context. Eighty-nine of the 91 BASF sequences have a recorded initial stage; 12 begin at GS31–33, 31 at GS34–39 and 16 at GS60 or later. The archived T–P–V implementation has independent German station-transfer evidence and explicit sowing dates for five external field-season cases. Missing BASF sowing dates and the absence of a leaf-emergence module prevent field-specific T–P–V coupling in the present disease forecasts.'),
    ('Calibration and transfer evaluation',
    'Parameter estimation minimizes squared disease-proxy error with equal coordinate-year weights, equal leaf-sequence weights within a coordinate-year and equal subsequent-assessment weights within a sequence. The fitted bounds are α∈[0,0.25] day⁻¹, β∈[0,10] day⁻¹ and an initial latent proportion in [0,0.8]. Six deterministic optimization starts assess local convergence. These coefficients scale model hazards and do not estimate viable-spore transmission rates.\n\n'
    'The 2018 forward split trains on 2017 and the 2019 split trains on 2017–2018, with every test coordinate purged from training. Eleven country holdouts retain location separation. Statistical comparators comprise persistence, global logit trend, thermal logit trend and canopy-conditioned logit trend. Mechanism ablations remove external acquisition, secondary acquisition, initial latent infection or all weather effects. Training-only location folds select among latent durations of 10, 20 and 30 days and infectious durations of 14, 21 and 28 days for a sensitivity variant. Earlier benchmark results from these years had already informed development, so the forward results are development-era transfer evidence.\n\n'
    'Paired uncertainty resamples entire locations, retaining all observed years and preserving the coordinate-year scoring denominator. The 5,000-draw intervals describe prediction-error variability conditional on the fitted models. They are not predictive intervals or parameter-confidence intervals. Frozen BASF parameters supply French forecasts without fitting French disease outcomes. French scoring gives equal weight to plot×cultivar×leaf sequences and then assessment dates within each station-year. Mixture plots contribute six such sequences and pure-cultivar plots contribute three; the estimand therefore does not assign equal weight to physical plots.'),
    ('Crop protection diagnostics',
    'Severity errors are stratified by all eligible leaves, the upper three leaves and the flag leaf, and by lead times of 1–7, 8–14, 15–28 and at least 29 days after initialization. The 1%, 5%, 10% and 25% diagnostic cutoffs summarize missed disease and false warnings on each source-specific recorded scale. A second subset includes only sequences below the selected cutoff at initialization. Weighted sensitivity, specificity, warning precision and continuous-severity AUC use the same nested field weights. Undefined class-specific rates remain missing when the required outcome class is absent.\n\n'
    'The cutoffs are descriptive sensitivity analyses specified after development-era severity results were available. They are not validated economic injury levels or fungicide-action thresholds. Deterministic predicted severity is not interpreted as a disease probability. Treatment savings, yield protection and economic benefit require intervention outcomes, harmonized yield measurements and an explicit decision rule; those quantities are not estimated by this untreated-disease evaluation.'),
    ('Field prediction results',
    f'The principal model gives all-leaf forward RMSEs of {metric("forward_2018","primary_secondary_hidden","all_eligible").rmse_pp:.2f} and {metric("forward_2019","primary_secondary_hidden","all_eligible").rmse_pp:.2f} percentage points in 2018 and 2019. On the upper three leaves, RMSE is {m18.rmse_pp:.2f} pp over {int(m18.n_targets)} later assessments at {int(m18.n_coordinate_years)} coordinate-years in 2018, compared with {l18.rmse_pp:.2f} pp for the logit-trend baseline. In 2019, the model RMSE is {m19.rmse_pp:.2f} pp over {int(m19.n_targets)} assessments at {int(m19.n_coordinate_years)} coordinate-years, compared with {l19.rmse_pp:.2f} pp for logit trend. The paired upper-leaf RMSE improvement over logit trend is {difference("forward_2018")} in 2018 and {difference("forward_2019")} in 2019. The comparative benefit is not consistent across years.\n\n'
    f'For upper-leaf sequences initialized below the 10% diagnostic cutoff, the model has weighted sensitivity {100*d18.sensitivity:.1f}% and a false-warning fraction of {100*d18.false_alarm_fraction:.1f}% in 2018. This subset contains {int(d18.n_targets)} targets, including only {int(d18.n_negative_targets)} negative outcomes. In 2019, sensitivity is {100*d19.sensitivity:.1f}% and the false-warning fraction is {100*d19.false_alarm_fraction:.1f}% over {int(d19.n_targets)} targets. The logit trend has sensitivity {100*b19.sensitivity:.1f}% and a false-warning fraction of {100*b19.false_alarm_fraction:.1f}% in that 2019 subset. These differences describe a sensitivity–false-warning tradeoff, with no demonstrated net crop-protection benefit.\n\n'
    f'Frozen BASF transfer to French pycnidial-area scores yields RMSE {f18.rmse_pp:.2f} pp in 2018 and {f19.rmse_pp:.2f} pp in 2019. Persistence yields {metric("external_2018","persistence_BASF_frozen").rmse_pp:.2f} and {metric("external_2019","persistence_BASF_frozen").rmse_pp:.2f} pp, while logit trend yields {metric("external_2018","global_logit_trend_BASF_frozen").rmse_pp:.2f} and {metric("external_2019","global_logit_trend_BASF_frozen").rmse_pp:.2f} pp. The current model does not establish competitive cross-endpoint transfer. The whole-development fit places α at zero and the initial latent proportion at {100*fit["initial_latent"]:.2f}% of unaffected tissue. Initially unresolved infection and secondary acquisition explain the fitted trajectory without a positive external coefficient. This optimum does not establish absence of natural primary infection.'),
    ('Regional interpretation and evidence limits',
    'The field-error maps locate observed trials and their upper-leaf prediction errors. They do not interpolate disease risk across unobserved European land. ERA5 continental layers supply historical climate context, while a separate one-day hourly extraction verifies the forcing pipeline. Full-season disease application additionally depends on wheat-specific area, regional sowing and leaf-emergence inputs, resistance information, initial disease pressure, forcing compatibility and checks on environmental extrapolation. A winter-cereal mask includes several hosts and is not a wheat-area denominator.\n\n'
    'The current evidence supports a reproducible crop-protection model-development benchmark with explicit lead-time and upper-leaf diagnostics. Natural first-infection timing, reliable prospective warnings, calibrated prediction intervals and economic benefit remain unvalidated. Controlled moisture and onset datasets provide supplementary process checks rather than European field-management outcomes. Climate-change attribution requires supported historical disease prediction and a documented climate ensemble; historical ERA5 or arbitrary warming perturbations do not provide future climate projections.')]
    references=[
        '1. AHDB. Introduction to foliar disease management in cereals. https://ahdb.org.uk/knowledge-library/introduction-to-foliar-disease-management-in-cereals',
        '2. AHDB. Fungicide programmes for wheat. https://ahdb.org.uk/knowledge-library/fungicide-programmes-for-wheat',
        '3. European wheat fungicide-trial data archive, 2017–2019. Zenodo. https://doi.org/10.5281/zenodo.6521175',
        '4. Orellana-Torrejon et al. Annual dynamics of Zymoseptoria tritici populations in wheat cultivar mixtures. Recherche Data Gouv data archive. https://doi.org/10.15454/4MAAI0',
        '5. Open-Meteo. Historical Weather API, ERA5 model selection and weather-variable definitions. https://open-meteo.com/en/docs/historical-weather-api']
    title='Septoria forecasting for wheat crop protection: conditional field evaluation'
    text=title+'\n\n'+'\n\n'.join(h+'\n\n'+body for h,body in sections)+'\n\nReferences\n\n'+'\n\n'.join(references)+'\n'
    forbidden=['this report will','this paper will','according to the user','we will','to avoid misunderstanding',
        'writing plan','i have','we have','this report does not']
    if any(s in text.lower() for s in forbidden):raise ValueError('Academic prose contains metawriting.')
    (OUT/'Crop_protection_methods_and_results.txt').write_text(text)
    doc=Document();section=doc.sections[0]
    section.top_margin=section.bottom_margin=Inches(.75)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(6)
    doc.add_heading(title,0)
    for heading,body in sections:
        doc.add_heading(heading,1)
        for paragraph in body.split('\n\n'):doc.add_paragraph(paragraph)
    doc.add_heading('Figures',1)
    doc.add_picture(str(ANALYSIS/'upper_leaf_forecast_horizons.png'),width=Inches(6.4))
    doc.add_paragraph('Figure 1. Upper-leaf forecast error by elapsed days from the first disease assessment. Each horizon is scored with equal coordinate-year, sequence and subsequent-assessment weights. n denotes the number of scored assessments; some bins contain only one or two targets. Curves connect descriptive bin scores and do not represent continuous-horizon validation. Realized future ERA5 forcing supplies the weather-conditioned model.')
    doc.add_picture(str(ANALYSIS/'upper_leaf_field_error_map.png'),width=Inches(6.4))
    doc.add_paragraph('Figure 2. Upper-leaf RMSE at observed BASF coordinates in the 2018 and 2019 forward-location-purged splits. Errors average assessments within sequence and then sequences within coordinate-year. Source coordinates are rounded and represent trial records rather than surveyed field boundaries. Longitude/latitude display; archived Natural Earth country outlines provide geographic context. The common colour scale spans 0–55 percentage points. No spatial disease-risk surface is inferred.')
    doc.add_heading('References',1)
    for reference in references:doc.add_paragraph(reference)
    doc.save(OUT/'Crop_protection_methods_and_results.docx')
    inputs=[ANALYSIS/'severity_metrics.csv',ANALYSIS/'threshold_diagnostics.csv',
        ANALYSIS/'upper_leaf_paired_comparisons.csv',Path(__file__),
        ROOT/'analysis/primary_secondary/calibration_v2_age/episodes.csv',
        ROOT/'analysis/primary_secondary/calibration_v2_age/all_development_fit.json',
        ANALYSIS/'evaluation_contract.json',ANALYSIS/'upper_leaf_forecast_horizons.png',
        ANALYSIS/'upper_leaf_field_error_map.png']
    receipt={'created_utc':datetime.now(timezone.utc).isoformat(),'document_role':'formal methods and actual results excerpt',
        'research_domain':'applied wheat crop protection','complete_submission_manuscript':False,
        'supported_operational_management_utility':False,
        'sources_sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs},
        'outputs_sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest()
            for p in [OUT/'Crop_protection_methods_and_results.txt',OUT/'Crop_protection_methods_and_results.docx']}}
    (ANALYSIS/'publication_excerpt_receipt.json').write_text(json.dumps(receipt,indent=2)+'\n')
    print(json.dumps(receipt,indent=2))


if __name__=='__main__':main()
