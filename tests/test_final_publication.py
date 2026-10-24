"""Final deliverable references, figures and biological typography."""
from pathlib import Path
import re

from docx import Document
from paper.analyses.source_support.verify import verify_documents

PUBLICATION = Path(__file__).resolve().parents[1] / "paper"


def test_final_document_figures_match_exact_extracted_images():
    result = verify_documents()
    assert result["Manuscript.docx"]["embedded_figures"] == 10
    assert result["Supplementary_Information.docx"]["embedded_figures"] == 3


def test_final_manuscript_species_and_citation_typography():
    document = Document(PUBLICATION / "Manuscript.docx")
    found_species = False
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.text in {"Zymoseptoria tritici", "Z. tritici"}:
                found_species = True
                assert run.italic
            if "Septoria tritici blotch" in run.text:
                assert not run.italic
            if re.fullmatch(r"\[\d+(?:,\s*\d+)*\]", run.text):
                assert run.font.superscript
    assert found_species
    text = "\n".join(p.text for p in document.paragraphs)
    assert "https://github.com/SmartAG-Team/paper_cc_impact_wheat_disease" in text
    assert "own-field" in text
    assert "50 km" in text
    assert "References" in text
