"""Render the same formal manuscript sources to editable Word and readable PDF."""
from io import BytesIO
import html
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement, parse_xml
from docx.shared import Cm, Pt
from latex2mathml.converter import convert as latex_to_mathml
from mathml2omml import convert as mathml_to_omml
from lxml import etree
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm, inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer,
    PageBreak, NextPageTemplate, Image, Table, TableStyle, KeepTogether)

from models.monitoring.innovation_common import ROOT
from models.monitoring.inputs import sha
from .figure_registry import MAIN_ROLES,SUPPLEMENT_ROLES,ROLE_NAMES
from .figure_placement import place_figures
from .supplementary_format import table_layout,style_word_table

OUT = ROOT/'paper'


def embedded_figure(block):
    match=re.fullmatch(r'!\[[^\]]*\]\(figures/(figs?\d+)\.png\)',block)
    return match[1] if match else None


def blocks(source):
    source=re.sub(r'\\\[(.*?)\\\]',lambda m:'\n\n@@MATH@@'+m[1].strip()+'@@ENDMATH@@\n\n',source,flags=re.S)
    return [b.strip() for b in re.split(r'\n\s*\n',source) if b.strip()]


def math_lines(latex):
    if '\\begin{aligned}' in latex:
        clean=latex.replace('\\begin{aligned}','').replace('\\end{aligned}','').replace('&','')
        return [s.strip() for s in clean.split('\\\\') if s.strip()]
    if '\\qquad' in latex: return [s.strip().rstrip(',') for s in latex.split('\\qquad')]
    lines=[s.strip().rstrip(',') for s in latex.splitlines() if s.strip()]
    if len(lines)==3 and lines[1].startswith('='): return [lines[0]+lines[1],lines[2]]
    return [' '.join(lines)]


def omml(latex):
    latex=latex.replace('\\widetilde','\\tilde')
    raw=mathml_to_omml(latex_to_mathml(latex))
    if 'xmlns:m=' not in raw: raw=raw.replace('<m:oMath','<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"',1)
    return parse_xml(raw)


def word_runs(paragraph,text):
    for part in re.split(r'(\\\(.*?\\\))',text,flags=re.S):
        if part.startswith('\\('): paragraph._p.append(omml(part[2:-2]))
        else:
            for value in re.split(r'(\*[^*]+\*)',part):
                if value.startswith('*') and value.endswith('*') and len(value)>2:
                    paragraph.add_run(value[1:-1]).italic=True
                else:paragraph.add_run(value)


def inline_pdf(text):
    def mathplain(m):
        root=etree.fromstring(latex_to_mathml(m[1]).encode())
        def convert(node):
            name=etree.QName(node).localname;children=list(node)
            if name=='msub':return convert(children[0])+'<sub>'+convert(children[1])+'</sub>'
            if name=='msup':return convert(children[0])+'<super>'+convert(children[1])+'</super>'
            if name=='msubsup':return convert(children[0])+'<sub>'+convert(children[1])+'</sub><super>'+convert(children[2])+'</super>'
            if name=='mfrac':return '('+convert(children[0])+')/('+convert(children[1])+')'
            if name=='mspace':return ' '
            return html.escape(node.text or '')+''.join(convert(c) for c in children)
        return convert(root)
    parts=[]
    for part in re.split(r'(\\\(.*?\\\))',text,flags=re.S):
        if part.startswith('\\('):parts.append(mathplain(re.match(r'\\\((.*?)\\\)',part,re.S)))
        else:parts.append(re.sub(r'\*([^*]+)\*',r'<i>\1</i>',html.escape(part)))
    return ''.join(parts)


def render_word(source,destination,figures=()):
    supplementary=destination.stem=='Supplementary_Information'
    doc=Document();section=doc.sections[0]
    doc.core_properties.author='Gang Zhao'
    doc.core_properties.title=source.splitlines()[0].lstrip('# ').strip()
    doc.core_properties.subject='Global Change Biology research article and supporting material'
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=Cm(2.2);section.left_margin=section.right_margin=Cm(2.5)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(12)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].font.bold=True
    title_properties=doc.styles['Title']._element.get_or_add_pPr()
    for border in list(title_properties):
        if etree.QName(border).localname=='pBdr':title_properties.remove(border)
    doc.styles['Normal'].paragraph_format.line_spacing=1.25 if supplementary else 1.15 if destination.stem=='Cover_Letter' else 2
    doc.styles['Normal'].paragraph_format.space_after=Pt(4)
    for style in ['Heading 1','Heading 2','Heading 3']:
        doc.styles[style].font.bold=True;doc.styles[style].paragraph_format.keep_with_next=True
    footer=section.footer.paragraphs[0];footer.alignment=1
    field=OxmlElement('w:fldSimple');field.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr','PAGE');footer._p.append(field)
    skip=False;equations=0;inline=[]
    def add_figure(name,caption):
        p=doc.add_paragraph();p.alignment=1;p.paragraph_format.line_spacing=1
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_after=Pt(6)
        p.add_run().add_picture(str(OUT/'figures'/f'{name}.png'),width=Cm(19))
        p=doc.add_paragraph(style='Caption');p.paragraph_format.line_spacing=1.15;p.paragraph_format.keep_together=True;word_runs(p,caption)
    for block in blocks(source):
        if block=='## Figure captions':skip=True;continue
        if skip and block.startswith('## '):skip=False
        if skip:continue
        if embedded_figure(block):
            name=embedded_figure(block);inline.append(name)
            section=doc.add_section(WD_SECTION.NEW_PAGE);section.left_margin=section.right_margin=Cm(1)
            add_figure(name,(OUT/'figures'/f'{name}_caption.txt').read_text().strip())
            section=doc.add_section(WD_SECTION.NEW_PAGE);section.left_margin=section.right_margin=Cm(2.5)
        elif block.startswith('@@MATH@@'):
            latex=block[len('@@MATH@@'):-len('@@ENDMATH@@')]
            for line in math_lines(latex):
                p=doc.add_paragraph();p.alignment=1;p._p.append(omml(line));equations+=1
        elif block.startswith('|'):
            rows=[[c.strip() for c in line.strip().strip('|').split('|')] for line in block.splitlines()]
            rows=[row for row in rows if not all(re.fullmatch(r':?-+:?',c) for c in row)]
            tab=doc.add_table(rows=1,cols=len(rows[0]));tab.style='Table Grid'
            head=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(head)
            for j,value in enumerate(rows[0]):tab.rows[0].cells[j].text=value
            for row in rows[1:]:
                for cell,value in zip(tab.add_row().cells,row):cell.text=value
            for row_number,row in enumerate(tab.rows):
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.line_spacing=1.15
                        if len(tab.rows)<=10:p.paragraph_format.keep_with_next=row_number<len(tab.rows)-1
                        for r in p.runs:r.font.size=Pt(12)
            if supplementary:
                widths,numeric=table_layout(rows[0]);style_word_table(tab,widths,numeric)
        elif block.startswith('#'):
            n=len(block)-len(block.lstrip('#'));text=block[n:].strip()
            p=doc.add_paragraph(style='Title' if n==1 else 'Heading '+str(min(n-1,3)));word_runs(p,text)
            if supplementary and n==2 and re.match(r'S\d+\.',text):p.paragraph_format.page_break_before=True
        else:
            p=doc.add_paragraph(style='Caption' if block.startswith('Table S') else None)
            if block.startswith('Table S'):
                p.paragraph_format.line_spacing=1.15;p.paragraph_format.keep_with_next=True
            word_runs(p,block.replace('\n',' '))
    appended=[(name,caption) for name,caption in figures if name not in inline]
    if appended:
        section=doc.add_section(WD_SECTION.NEW_PAGE);section.left_margin=section.right_margin=Cm(1)
        for i,(name,caption) in enumerate(appended):
            if i:doc.add_page_break()
            add_figure(name,caption)
    if destination.stem=='Manuscript':
        for section in doc.sections:
            node=OxmlElement('w:lnNumType')
            for key,value in [('countBy','1'),('restart','continuous'),('distance','240')]:
                node.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'+key,value)
            section._sectPr.append(node)
    doc.save(destination)
    return dict(path=str(destination.relative_to(OUT)),native_equation_blocks=equations,
        figures=len(figures),inline_figures=inline,appended_figures=[n for n,_ in appended])


def equation_image(latex):
    latex=latex.replace('\\tfrac','\\frac')
    latex=re.sub(r'\\frac\s*1\s*2',r'\\frac{1}{2}',latex)
    fig=plt.figure(figsize=(8,.65),dpi=200);text=fig.text(.01,.5,'$'+latex+'$',fontsize=12,va='center')
    fig.canvas.draw();box=text.get_window_extent(fig.canvas.get_renderer()).expanded(1.025,1.25)
    bio=BytesIO();fig.savefig(bio,format='png',dpi=200,bbox_inches=box.transformed(fig.dpi_scale_trans.inverted()),pad_inches=0)
    plt.close(fig);bio.seek(0);im=PILImage.open(bio);width,height=im.size
    if width/200*inch > A4[0]-5*cm:raise ValueError('Equation needs another line, not a smaller font: '+latex)
    bio.seek(0);return Image(bio,width=width/200*inch,height=height/200*inch)


def render_pdf(source,destination,figures=()):
    supplementary=destination.stem=='Supplementary_Information'
    serif=font_manager.findfont('DejaVu Serif');bold=font_manager.findfont(font_manager.FontProperties(family='DejaVu Serif',weight='bold'))
    italic=font_manager.findfont(font_manager.FontProperties(family='DejaVu Serif',style='italic'))
    pdfmetrics.registerFont(TTFont('ResearchSerif',serif));pdfmetrics.registerFont(TTFont('ResearchSerifBold',bold))
    pdfmetrics.registerFont(TTFont('ResearchSerifItalic',italic))
    pdfmetrics.registerFontFamily('ResearchSerif',normal='ResearchSerif',bold='ResearchSerifBold',italic='ResearchSerifItalic',boldItalic='ResearchSerifBold')
    body=ParagraphStyle('body',fontName='ResearchSerif',fontSize=12,
        leading=15 if supplementary or destination.stem=='Cover_Letter' else 24,spaceAfter=5,splitLongWords=True)
    heading=ParagraphStyle('heading',parent=body,fontName='ResearchSerifBold',leading=17,spaceBefore=12,spaceAfter=6,keepWithNext=True)
    title=ParagraphStyle('title',parent=heading,fontSize=16,leading=21,spaceAfter=14)
    compact=ParagraphStyle('compact',parent=body,leading=15,spaceAfter=6)
    def footer(canvas,doc):
        canvas.setFont('ResearchSerif',12);canvas.drawCentredString(A4[0]/2,1.1*cm,str(doc.page))
    document=BaseDocTemplate(str(destination),pagesize=A4,leftMargin=2.5*cm,rightMargin=2.5*cm,topMargin=2.2*cm,bottomMargin=2.2*cm,
        author='Gang Zhao',title=source.splitlines()[0].lstrip('# ').strip())
    document.addPageTemplates([PageTemplate('text',Frame(2.5*cm,2.2*cm,A4[0]-5*cm,A4[1]-4.4*cm,leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0),onPage=footer),
        PageTemplate('figure',Frame(cm,2*cm,A4[0]-2*cm,A4[1]-4*cm,leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0),onPage=footer)])
    if supplementary:
        def after_flowable(flowable):
            if hasattr(flowable,'supplement_outline'):
                label,level,key=flowable.supplement_outline
                document.canv.bookmarkPage(key);document.canv.addOutlineEntry(label,key,level=level,closed=False)
        document.afterFlowable=after_flowable
    story=[];skip=False;inline=[]
    for block in blocks(source):
        if block=='## Figure captions':skip=True;continue
        if skip and block.startswith('## '):skip=False
        if skip:continue
        if embedded_figure(block):
            name=embedded_figure(block);inline.append(name)
            path=OUT/'figures'/f'{name}.png';im=PILImage.open(path)
            caption=(OUT/'figures'/f'{name}_caption.txt').read_text().strip()
            image_height=19*cm*im.height/im.width
            caption_paragraph=Paragraph(inline_pdf(caption),compact)
            _,caption_height=caption_paragraph.wrap(19*cm,A4[1])
            assert image_height+10+caption_height+compact.spaceAfter<=A4[1]-4*cm, (name,'figure/caption exceed one page; shorten caption or change layout')
            story.extend([NextPageTemplate('figure'),PageBreak(),
                KeepTogether([Image(str(path),width=19*cm,height=image_height),Spacer(1,10),caption_paragraph]),
                NextPageTemplate('text'),PageBreak()])
        elif block.startswith('@@MATH@@'):
            for line in math_lines(block[len('@@MATH@@'):-len('@@ENDMATH@@')]):story.extend([equation_image(line),Spacer(1,8)])
        elif block.startswith('|'):
            rows=[[c.strip() for c in line.strip().strip('|').split('|')] for line in block.splitlines()]
            rows=[row for row in rows if not all(re.fullmatch(r':?-+:?',c) for c in row)]
            n=len(rows[0]);widths=([5.4*cm,2.0*cm,8.6*cm] if n==3 else
                [5.2*cm,3.4*cm,2.2*cm,2.6*cm,2.6*cm] if n==5 else
                [2.3*cm,3.1*cm,4.0*cm,2.1*cm,2.1*cm,2.4*cm])
            if n==3 and rows[0][0]=='Component':widths=[3.1*cm,6.4*cm,6.5*cm]
            numeric=set()
            if supplementary:
                sizes,numeric=table_layout(rows[0]);widths=[v*cm for v in sizes]
            cellstyle=ParagraphStyle('cell',parent=compact,leading=14 if supplementary else 15)
            rightstyle=ParagraphStyle('rightcell',parent=cellstyle,alignment=2)
            headstyle=ParagraphStyle('headercell',parent=cellstyle,fontName='ResearchSerifBold',textColor=colors.HexColor('#254F66'))
            headright=ParagraphStyle('headerright',parent=headstyle,alignment=2)
            data=[[Paragraph(inline_pdf(c),(headright if j in numeric else headstyle) if supplementary and i==0 else rightstyle if j in numeric else cellstyle) for j,c in enumerate(row)] for i,row in enumerate(rows)]
            table=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EEEEEE')),
                ('LINEBELOW',(0,0),(-1,0),.7,colors.grey),('LINEBELOW',(0,-1),(-1,-1),.5,colors.grey),
                ('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3)]))
            if supplementary:
                table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EAF1F5')),
                    ('LINEABOVE',(0,0),(-1,0),.8,colors.HexColor('#70818C')),
                    ('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3),
                    ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F7F9FA')])]))
            if supplementary and len(rows)<=10 and story and isinstance(story[-1],Paragraph) and story[-1].getPlainText().startswith('Table S'):
                caption=story.pop();story.append(KeepTogether([caption,table]))
            else:story.append(KeepTogether([table]) if len(rows)<=10 else table)
            story.append(Spacer(1,8))
        elif block.startswith('#'):
            n=len(block)-len(block.lstrip('#'));label=block[n:].strip()
            if supplementary and n==2 and re.match(r'S\d+\.',label):story.append(PageBreak())
            paragraph=Paragraph(inline_pdf(label),title if n==1 else heading)
            if supplementary:
                level=0 if n==1 else 2 if n==3 and re.match(r'S\d+\.',label) else 1
                paragraph.supplement_outline=(label,level,'supp-'+str(len(story)))
            story.append(paragraph)
        else:
            paragraph=Paragraph(inline_pdf(block.replace('\n',' ')),compact if block.startswith('Table S') else body)
            if block.startswith('Table S'):paragraph.keepWithNext=True
            story.append(paragraph)
    for i,(name,caption) in enumerate(figures):
        if name in inline:continue
        story.extend([NextPageTemplate('figure'),PageBreak()])
        path=OUT/'figures'/f'{name}.png';im=PILImage.open(path)
        story.extend([Image(str(path),width=19*cm,height=19*cm*im.height/im.width),Spacer(1,10),Paragraph(inline_pdf(caption),compact)])
    document.build(story)
    if destination.stem=='Manuscript':
        # Derive review line numbers from final text positions, without rasterizing.
        import pymupdf
        pdf=pymupdf.open(destination);number=0
        for page in pdf:
            positions=[]
            for block in page.get_text('dict')['blocks']:
                if block['type']!=0:continue
                for line in block['lines']:
                    x0,y0,x1,y1=line['bbox']
                    if y0<2*cm or y1>page.rect.height-1.6*cm:continue
                    positions.append((y0,x0,y1))
            seen=[]
            for y0,x0,y1 in sorted(positions):
                if seen and abs(y0-seen[-1])<2:continue
                seen.append(y0);number+=1
                page.insert_textbox(pymupdf.Rect(.08*cm,y0, .85*cm,y1+3),str(number),
                    fontsize=8,fontname='tiro',color=(.45,.45,.45),align=2)
        temporary=destination.with_suffix('.numbered.pdf');pdf.save(temporary);pdf.close()
        temporary.replace(destination)


def build():
    profiles=[]
    for name,figure_names in [('Manuscript',[ROLE_NAMES[r] for r in MAIN_ROLES]),
            ('Supplementary_Information',[ROLE_NAMES[r] for r in SUPPLEMENT_ROLES]),('Cover_Letter',[])]:
        path=OUT/f'{name}.md'
        if not path.exists():continue
        source=place_figures(path.read_text(),figure_names)
        path.write_text(source)
        if name=='Manuscript':(OUT/'Manuscript.txt').write_text(source)
        figures=[(n,(OUT/'figures'/f'{n}_caption.txt').read_text().strip()) for n in figure_names]
        profiles.append(render_word(source,OUT/f'{name}.docx',figures))
        render_pdf(source,OUT/f'{name}.pdf',figures)
    profiles.append(render_word((OUT/'Mathematical_Model_Specification.md').read_text(),OUT/'Mathematical_Model_Specification.docx'))
    (OUT/'document_profiles.json').write_text(json.dumps(dict(documents=profiles,source_renderer_sha256=sha(Path(__file__)),
        pdf_scope='PDF and DOCX share the same source prose, equations and figure files; pagination is renderer-specific',
        native_figure_width_cm=19,manuscript_body_font_pt=12),indent=2)+'\n')
    print('documents',len(profiles),flush=True)
    return OUT


if __name__ == '__main__':build()
